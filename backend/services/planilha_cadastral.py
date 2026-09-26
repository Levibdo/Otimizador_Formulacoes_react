import io
import re
import unicodedata
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


VERSAO_CONTRATO = "1.0"
ABAS = {
    "LEIA_ME": ("CHAVE", "VALOR"),
    "MATERIAS_PRIMAS": ("ACAO", "CODIGO", "NOME", "ATIVA"),
    "NUTRIENTES": ("ACAO", "CODIGO", "NOME", "UNIDADE"),
    "COMPOSICAO_NUTRICIONAL": (
        "ACAO",
        "MATERIA_PRIMA_CODIGO",
        "NUTRIENTE_CODIGO",
        "VALOR",
    ),
    "PRECOS_MP": (
        "ACAO",
        "MATERIA_PRIMA_CODIGO",
        "PRECO_KG",
        "VIGENCIA_INICIO",
        "VIGENCIA_FIM",
    ),
}
ACOES = {"CRIAR", "ATUALIZAR", "DESATIVAR"}
ACOES_POR_ABA = {
    "MATERIAS_PRIMAS": {"CRIAR", "ATUALIZAR", "DESATIVAR"},
    "NUTRIENTES": {"CRIAR", "ATUALIZAR"},
    "COMPOSICAO_NUTRICIONAL": {"CRIAR", "ATUALIZAR"},
    "PRECOS_MP": {"CRIAR"},
}
PADRAO_CODIGO = re.compile(r"^[A-Z][A-Z0-9_]{0,29}$")

MAX_ARQUIVO_BYTES = 5 * 1024 * 1024
MAX_DESCOMPRIMIDO_BYTES = 30 * 1024 * 1024
MAX_ENTRADAS_ZIP = 200
MAX_RAZAO_COMPRESSAO = 100
MAX_ABAS = 5
MAX_LINHAS_POR_ABA = 5_000
MAX_COLUNAS_POR_ABA = 12
MAX_CELULAS_TOTAL = 100_000


def _token(valor) -> str:
    texto = "".join(
        caractere
        for caractere in unicodedata.normalize("NFD", str(valor or "").strip())
        if unicodedata.category(caractere) != "Mn"
    )
    return re.sub(r"[^A-Z0-9]+", "_", texto.upper()).strip("_")


def _vazio(valor) -> bool:
    return valor is None or (isinstance(valor, str) and not valor.strip())


def _texto(valor):
    return None if _vazio(valor) else str(valor).strip()


def _diagnostico(severidade, aba, mensagem, linha=None, coluna=None, codigo=None):
    return {
        "severidade": severidade,
        "aba": aba,
        "linha": linha,
        "coluna": coluna,
        "codigo": codigo,
        "mensagem": mensagem,
    }


def _numero(valor, aba, linha, coluna, diagnosticos, obrigatorio=True):
    if _vazio(valor):
        if obrigatorio:
            diagnosticos.append(_diagnostico("ERRO", aba, "Valor obrigatório ausente.", linha, coluna))
        return None
    if isinstance(valor, bool):
        diagnosticos.append(_diagnostico("ERRO", aba, "Valor numérico inválido.", linha, coluna))
        return None
    try:
        numero = Decimal(str(valor).strip())
    except (InvalidOperation, ValueError):
        diagnosticos.append(_diagnostico("ERRO", aba, "Valor numérico inválido.", linha, coluna))
        return None
    if not numero.is_finite():
        diagnosticos.append(_diagnostico("ERRO", aba, "Valor numérico deve ser finito.", linha, coluna))
        return None
    if numero < 0:
        diagnosticos.append(_diagnostico("ERRO", aba, "Valor negativo não permitido.", linha, coluna))
        return None
    return format(numero, "f")


def _booleano(valor, aba, linha, coluna, diagnosticos, obrigatorio=False):
    if _vazio(valor):
        if obrigatorio:
            diagnosticos.append(_diagnostico("ERRO", aba, "Valor booleano obrigatório ausente.", linha, coluna))
        return None
    if not isinstance(valor, str):
        diagnosticos.append(
            _diagnostico("ERRO", aba, "Booleano inválido; use o texto SIM ou NÃO.", linha, coluna)
        )
        return None
    normalizado = _token(valor)
    if normalizado == "SIM":
        return True
    if normalizado == "NAO":
        return False
    diagnosticos.append(
        _diagnostico("ERRO", aba, "Booleano inválido; use o texto SIM ou NÃO.", linha, coluna)
    )
    return None


def _data(valor, aba, linha, coluna, diagnosticos, obrigatorio=False):
    if _vazio(valor):
        if obrigatorio:
            diagnosticos.append(_diagnostico("ERRO", aba, "Data obrigatória ausente.", linha, coluna))
        return None
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    try:
        return date.fromisoformat(str(valor).strip()).isoformat()
    except ValueError:
        diagnosticos.append(
            _diagnostico("ERRO", aba, "Data inválida; use AAAA-MM-DD.", linha, coluna)
        )
        return None


def _validar_codigo(valor, aba, linha, coluna, diagnosticos):
    if _vazio(valor):
        diagnosticos.append(_diagnostico("ERRO", aba, "Código obrigatório ausente.", linha, coluna))
        return None
    if not isinstance(valor, str):
        diagnosticos.append(
            _diagnostico(
                "ERRO",
                aba,
                "Código deve ser informado como texto para preservar seu conteúdo.",
                linha,
                coluna,
            )
        )
        return None
    codigo = valor.strip().upper()
    if not PADRAO_CODIGO.fullmatch(codigo):
        diagnosticos.append(
            _diagnostico(
                "ERRO",
                aba,
                "Código inválido; use letra inicial, letras maiúsculas, números e sublinhado, com até 30 caracteres.",
                linha,
                coluna,
                codigo,
            )
        )
        return codigo
    return codigo


def _validar_acao(valor, aba, linha, diagnosticos):
    acao = _token(valor)
    if not acao:
        diagnosticos.append(_diagnostico("ERRO", aba, "Ação obrigatória ausente; nenhuma ação é inferida.", linha, "ACAO"))
        return None
    if acao not in ACOES:
        diagnosticos.append(
            _diagnostico("ERRO", aba, "Ação inválida; use CRIAR, ATUALIZAR ou DESATIVAR.", linha, "ACAO")
        )
    elif acao not in ACOES_POR_ABA[aba]:
        permitidas = ", ".join(sorted(ACOES_POR_ABA[aba]))
        diagnosticos.append(
            _diagnostico(
                "ERRO",
                aba,
                f"Ação {acao} não é suportada com segurança nesta aba; use {permitidas}.",
                linha,
                "ACAO",
            )
        )
    return acao


def _inspecionar_zip(conteudo: bytes, diagnosticos) -> bool:
    if len(conteudo) > MAX_ARQUIVO_BYTES:
        diagnosticos.append(_diagnostico("ERRO", None, "Arquivo excede o limite de 5 MiB."))
        return False
    arquivo = io.BytesIO(conteudo)
    if not zipfile.is_zipfile(arquivo):
        diagnosticos.append(_diagnostico("ERRO", None, "Conteúdo inválido ou arquivo XLSX corrompido."))
        return False
    try:
        with zipfile.ZipFile(arquivo) as pacote:
            entradas = pacote.infolist()
            if len(entradas) > MAX_ENTRADAS_ZIP:
                diagnosticos.append(_diagnostico("ERRO", None, "Arquivo XLSX possui entradas internas em excesso."))
                return False
            total = sum(item.file_size for item in entradas)
            if total > MAX_DESCOMPRIMIDO_BYTES:
                diagnosticos.append(_diagnostico("ERRO", None, "Conteúdo descomprimido excede o limite de 30 MiB."))
                return False
            for item in entradas:
                nome = item.filename.replace("\\", "/")
                if item.flag_bits & 1:
                    diagnosticos.append(_diagnostico("ERRO", None, "Arquivos XLSX criptografados não são aceitos."))
                    return False
                if nome.startswith("/") or "../" in nome or nome.startswith("../"):
                    diagnosticos.append(_diagnostico("ERRO", None, "Estrutura interna inválida no arquivo XLSX."))
                    return False
                if "vbaproject" in nome.lower():
                    diagnosticos.append(_diagnostico("ERRO", None, "Macros não são permitidas no template cadastral."))
                    return False
                if item.file_size and item.file_size / max(item.compress_size, 1) > MAX_RAZAO_COMPRESSAO:
                    diagnosticos.append(_diagnostico("ERRO", None, "Arquivo XLSX apresenta taxa de compressão insegura."))
                    return False
    except (OSError, zipfile.BadZipFile):
        diagnosticos.append(_diagnostico("ERRO", None, "Não foi possível inspecionar o arquivo XLSX."))
        return False
    return True


def _mapear_cabecalho(ws, aba, diagnosticos):
    valores = [celula.value for celula in ws[1]] if ws.max_row else []
    normalizados = [_token(valor) for valor in valores]
    while normalizados and not normalizados[-1]:
        normalizados.pop()
    repetidos = sorted({item for item in normalizados if item and normalizados.count(item) > 1})
    for item in repetidos:
        diagnosticos.append(_diagnostico("ERRO", aba, "Cabeçalho repetido.", 1, item))
    esperados = set(ABAS[aba])
    presentes = {item for item in normalizados if item}
    for item in sorted(esperados - presentes):
        diagnosticos.append(_diagnostico("ERRO", aba, "Cabeçalho obrigatório ausente.", 1, item))
    for item in sorted(presentes - esperados):
        diagnosticos.append(_diagnostico("ERRO", aba, "Cabeçalho desconhecido.", 1, item))
    if repetidos or esperados != presentes:
        return None
    return {nome: normalizados.index(nome) + 1 for nome in ABAS[aba]}


def _linhas(ws, mapa):
    for numero in range(2, ws.max_row + 1):
        valores = {coluna: ws.cell(numero, indice).value for coluna, indice in mapa.items()}
        if any(not _vazio(valor) for valor in valores.values()):
            yield numero, valores


def _exigir_texto(valor, aba, linha, coluna, diagnosticos, obrigatorio):
    texto = _texto(valor)
    if obrigatorio and texto is None:
        diagnosticos.append(_diagnostico("ERRO", aba, "Valor obrigatório ausente.", linha, coluna))
    return texto


def _parsear_catalogo(ws, aba, mapa, diagnosticos):
    dados, chaves = [], set()
    for linha, valores in _linhas(ws, mapa):
        acao = _validar_acao(valores["ACAO"], aba, linha, diagnosticos)
        codigo = _validar_codigo(valores["CODIGO"], aba, linha, "CODIGO", diagnosticos)
        if codigo in chaves:
            diagnosticos.append(_diagnostico("ERRO", aba, "Código duplicado nesta aba.", linha, "CODIGO", codigo))
        elif codigo:
            chaves.add(codigo)
        criar = acao == "CRIAR"
        nome = _exigir_texto(valores["NOME"], aba, linha, "NOME", diagnosticos, criar)
        item = {"linha": linha, "acao": acao, "codigo": codigo, "nome": nome}
        if aba == "MATERIAS_PRIMAS":
            item["ativa"] = _booleano(valores["ATIVA"], aba, linha, "ATIVA", diagnosticos)
            if acao == "ATUALIZAR" and nome is None and item["ativa"] is None:
                diagnosticos.append(_diagnostico("ERRO", aba, "ATUALIZAR exige ao menos um campo preenchido.", linha, "ACAO", codigo))
        else:
            item["unidade"] = _exigir_texto(valores["UNIDADE"], aba, linha, "UNIDADE", diagnosticos, criar)
            if acao == "ATUALIZAR" and nome is None and item["unidade"] is None:
                diagnosticos.append(_diagnostico("ERRO", aba, "ATUALIZAR exige ao menos um campo preenchido.", linha, "ACAO", codigo))
        dados.append(item)
    return dados


def _parsear_composicoes(ws, mapa, diagnosticos):
    aba, dados, chaves = "COMPOSICAO_NUTRICIONAL", [], set()
    for linha, valores in _linhas(ws, mapa):
        acao = _validar_acao(valores["ACAO"], aba, linha, diagnosticos)
        mp = _validar_codigo(valores["MATERIA_PRIMA_CODIGO"], aba, linha, "MATERIA_PRIMA_CODIGO", diagnosticos)
        nutriente = _validar_codigo(valores["NUTRIENTE_CODIGO"], aba, linha, "NUTRIENTE_CODIGO", diagnosticos)
        chave = (mp, nutriente)
        if None not in chave and chave in chaves:
            diagnosticos.append(_diagnostico("ERRO", aba, "Composição duplicada nesta aba.", linha, "NUTRIENTE_CODIGO", f"{mp}/{nutriente}"))
        else:
            chaves.add(chave)
        valor = _numero(valores["VALOR"], aba, linha, "VALOR", diagnosticos, acao in {"CRIAR", "ATUALIZAR"})
        dados.append({"linha": linha, "acao": acao, "materia_prima_codigo": mp, "nutriente_codigo": nutriente, "valor": valor})
    return dados


def _parsear_precos(ws, mapa, diagnosticos):
    aba, dados, chaves = "PRECOS_MP", [], set()
    for linha, valores in _linhas(ws, mapa):
        acao = _validar_acao(valores["ACAO"], aba, linha, diagnosticos)
        mp = _validar_codigo(valores["MATERIA_PRIMA_CODIGO"], aba, linha, "MATERIA_PRIMA_CODIGO", diagnosticos)
        preco = _numero(valores["PRECO_KG"], aba, linha, "PRECO_KG", diagnosticos)
        inicio = _data(valores["VIGENCIA_INICIO"], aba, linha, "VIGENCIA_INICIO", diagnosticos, True)
        fim = _data(valores["VIGENCIA_FIM"], aba, linha, "VIGENCIA_FIM", diagnosticos)
        chave = (mp, inicio)
        if None not in chave and chave in chaves:
            diagnosticos.append(_diagnostico("ERRO", aba, "Preço duplicado para a MP e vigência inicial.", linha, "VIGENCIA_INICIO", mp))
        else:
            chaves.add(chave)
        if inicio and fim and fim < inicio:
            diagnosticos.append(_diagnostico("ERRO", aba, "Vigência final anterior à inicial.", linha, "VIGENCIA_FIM", mp))
        dados.append({"linha": linha, "acao": acao, "materia_prima_codigo": mp, "preco_kg": preco, "vigencia_inicio": inicio, "vigencia_fim": fim})
    return dados


def _validar_referencias(dados, diagnosticos):
    mps = {item["codigo"]: item for item in dados["MATERIAS_PRIMAS"] if item["codigo"]}
    nutrientes = {item["codigo"]: item for item in dados["NUTRIENTES"] if item["codigo"]}
    for aba in ("COMPOSICAO_NUTRICIONAL", "PRECOS_MP"):
        for item in dados[aba]:
            indice = item["linha"]
            codigo = item["materia_prima_codigo"]
            if codigo and codigo not in mps:
                diagnosticos.append(_diagnostico("AVISO", aba, "Referência de matéria-prima não declarada no arquivo; deverá existir no banco na pré-validação.", indice, "MATERIA_PRIMA_CODIGO", codigo))
            elif codigo and mps[codigo]["acao"] == "DESATIVAR":
                diagnosticos.append(_diagnostico("ERRO", aba, "Referência aponta para matéria-prima desativada no mesmo lote.", indice, "MATERIA_PRIMA_CODIGO", codigo))
    for item in dados["COMPOSICAO_NUTRICIONAL"]:
        indice = item["linha"]
        codigo = item["nutriente_codigo"]
        if codigo and codigo not in nutrientes:
            diagnosticos.append(_diagnostico("AVISO", "COMPOSICAO_NUTRICIONAL", "Referência de nutriente não declarada no arquivo; deverá existir no banco na pré-validação.", indice, "NUTRIENTE_CODIGO", codigo))
        elif codigo and nutrientes[codigo]["acao"] == "DESATIVAR":
            diagnosticos.append(_diagnostico("ERRO", "COMPOSICAO_NUTRICIONAL", "Referência aponta para nutriente desativado no mesmo lote.", indice, "NUTRIENTE_CODIGO", codigo))


def parsear_planilha_cadastral(conteudo: bytes, nome_arquivo: str) -> dict:
    diagnosticos = []
    dados = {aba: [] for aba in ABAS if aba != "LEIA_ME"}
    if not nome_arquivo.lower().endswith(".xlsx"):
        diagnosticos.append(_diagnostico("ERRO", None, "Formato não suportado; use somente .xlsx."))
        return {"versao": None, "valido": False, "dados": dados, "diagnosticos": diagnosticos}
    if not conteudo:
        diagnosticos.append(_diagnostico("ERRO", None, "Arquivo vazio."))
        return {"versao": None, "valido": False, "dados": dados, "diagnosticos": diagnosticos}
    if not _inspecionar_zip(conteudo, diagnosticos):
        return {"versao": None, "valido": False, "dados": dados, "diagnosticos": diagnosticos}
    try:
        workbook = load_workbook(
            io.BytesIO(conteudo),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception:
        diagnosticos.append(_diagnostico("ERRO", None, "Não foi possível abrir o conteúdo XLSX."))
        return {"versao": None, "valido": False, "dados": dados, "diagnosticos": diagnosticos}

    nomes = workbook.sheetnames
    if len(nomes) > MAX_ABAS:
        diagnosticos.append(_diagnostico("ERRO", None, "Quantidade de abas excede o limite de 5."))
    for aba in sorted(set(ABAS) - set(nomes)):
        diagnosticos.append(_diagnostico("ERRO", aba, "Aba obrigatória ausente."))
    for aba in sorted(set(nomes) - set(ABAS)):
        diagnosticos.append(_diagnostico("ERRO", aba, "Aba desconhecida."))

    total_celulas = 0
    mapas = {}
    for aba in nomes:
        ws = workbook[aba]
        total_celulas += ws.max_row * ws.max_column
        if total_celulas > MAX_CELULAS_TOTAL:
            diagnosticos.append(_diagnostico("ERRO", None, "Workbook excede o limite de 100.000 células."))
            continue
        if ws.max_row > MAX_LINHAS_POR_ABA:
            diagnosticos.append(_diagnostico("ERRO", aba, "Aba excede o limite de 5.000 linhas."))
            continue
        if ws.max_column > MAX_COLUNAS_POR_ABA:
            diagnosticos.append(_diagnostico("ERRO", aba, "Aba excede o limite de 12 colunas."))
            continue
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    diagnosticos.append(_diagnostico("ERRO", aba, "Fórmulas não são permitidas.", cell.row, cell.column_letter))
        if aba in ABAS:
            mapas[aba] = _mapear_cabecalho(ws, aba, diagnosticos)
    versao = None
    if "LEIA_ME" in mapas and mapas["LEIA_ME"]:
        chaves = {}
        for linha, valores in _linhas(workbook["LEIA_ME"], mapas["LEIA_ME"]):
            chave = _token(valores["CHAVE"])
            if not chave:
                diagnosticos.append(_diagnostico("ERRO", "LEIA_ME", "Chave obrigatória ausente.", linha, "CHAVE"))
                continue
            if chave in chaves:
                diagnosticos.append(_diagnostico("ERRO", "LEIA_ME", "Chave duplicada.", linha, "CHAVE", chave))
            chaves[chave] = _texto(valores["VALOR"])
        versao = chaves.get("VERSAO_TEMPLATE")
        if versao != VERSAO_CONTRATO:
            diagnosticos.append(_diagnostico("ERRO", "LEIA_ME", "Versão do template ausente ou incompatível; esperado 1.0.", None, "VALOR", "VERSAO_TEMPLATE"))

    if mapas.get("MATERIAS_PRIMAS"):
        dados["MATERIAS_PRIMAS"] = _parsear_catalogo(workbook["MATERIAS_PRIMAS"], "MATERIAS_PRIMAS", mapas["MATERIAS_PRIMAS"], diagnosticos)
    if mapas.get("NUTRIENTES"):
        dados["NUTRIENTES"] = _parsear_catalogo(workbook["NUTRIENTES"], "NUTRIENTES", mapas["NUTRIENTES"], diagnosticos)
    if mapas.get("COMPOSICAO_NUTRICIONAL"):
        dados["COMPOSICAO_NUTRICIONAL"] = _parsear_composicoes(workbook["COMPOSICAO_NUTRICIONAL"], mapas["COMPOSICAO_NUTRICIONAL"], diagnosticos)
    if mapas.get("PRECOS_MP"):
        dados["PRECOS_MP"] = _parsear_precos(workbook["PRECOS_MP"], mapas["PRECOS_MP"], diagnosticos)
    _validar_referencias(dados, diagnosticos)
    valido = not any(item["severidade"] == "ERRO" for item in diagnosticos)
    return {"versao": versao, "valido": valido, "dados": dados, "diagnosticos": diagnosticos}


def gerar_template_cadastral() -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    exemplos = {
        "LEIA_ME": [
            ("VERSAO_TEMPLATE", VERSAO_CONTRATO),
            ("ESCOPO", "Dados exclusivamente fictícios para demonstrar o contrato cadastral."),
            ("ACOES", "CRIAR, ATUALIZAR ou DESATIVAR; nenhuma ação é inferida."),
            ("IDENTIFICACAO", "Use somente códigos estáveis; códigos não podem ser alterados."),
        ],
        "MATERIAS_PRIMAS": [("CRIAR", "FICT_MP_AVEIA", "Aveia fictícia", "SIM")],
        "NUTRIENTES": [("CRIAR", "FICT_NUT_PROTEINA", "Proteína fictícia", "g/100 g")],
        "COMPOSICAO_NUTRICIONAL": [("CRIAR", "FICT_MP_AVEIA", "FICT_NUT_PROTEINA", 12.5)],
        "PRECOS_MP": [("CRIAR", "FICT_MP_AVEIA", 10.25, date(2026, 1, 1), None)],
    }
    for aba, colunas in ABAS.items():
        ws = workbook.create_sheet(aba)
        ws.append(colunas)
        for linha in exemplos[aba]:
            ws.append(linha)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        ws.freeze_panes = "A2"
        for coluna in ws.columns:
            largura = min(max(len(str(cell.value or "")) for cell in coluna) + 2, 60)
            ws.column_dimensions[coluna[0].column_letter].width = largura
    saida = io.BytesIO()
    workbook.save(saida)
    return saida.getvalue()
