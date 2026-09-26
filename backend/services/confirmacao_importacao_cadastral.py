import io
from datetime import date, timezone
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from models import ComposicaoMateriaPrima, MateriaPrima, Nutriente, PrecoMateriaPrima
from repositories.importacao_cadastral_repository import ImportacaoCadastralRepository
from services.planilha_cadastral import gerar_template_cadastral
from services.pre_validacao_cadastral import pre_validar_planilha_cadastral_completo
from services.staging_importacao_cadastral import serializar_canonico

MAX_CODIGOS_RESULTADO = 500
ORDEM_LOCKS = ("nutrientes", "materias_primas", "composicoes_materias_primas", "precos_materias_primas")


class CredenciaisSessaoInvalidas(Exception):
    pass


class ConflitoSessao(Exception):
    def __init__(self, mensagem):
        self.mensagem = mensagem


class RevalidacaoFalhou(ConflitoSessao):
    pass


def _xlsx_do_payload(payload):
    workbook = load_workbook(io.BytesIO(gerar_template_cadastral()))
    colunas = {
        "MATERIAS_PRIMAS": ("acao", "codigo", "nome", "ativa"),
        "NUTRIENTES": ("acao", "codigo", "nome", "unidade"),
        "COMPOSICAO_NUTRICIONAL": ("acao", "materia_prima_codigo", "nutriente_codigo", "valor"),
        "PRECOS_MP": ("acao", "materia_prima_codigo", "preco_kg", "vigencia_inicio", "vigencia_fim"),
    }
    for aba, campos in colunas.items():
        ws = workbook[aba]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        for item in sorted(payload["dados"][aba], key=lambda valor: valor["linha"]):
            for coluna, campo in enumerate(campos, start=1):
                valor = item.get(campo)
                if campo == "ativa" and isinstance(valor, bool):
                    valor = "SIM" if valor else "NÃO"
                ws.cell(row=item["linha"], column=coluna, value=valor)
    saida = io.BytesIO()
    workbook.save(saida)
    return saida.getvalue()


def _bloquear_cadastros(db):
    if db.bind.dialect.name == "postgresql":
        for tabela in ORDEM_LOCKS:
            db.execute(text(f'LOCK TABLE {tabela} IN SHARE ROW EXCLUSIVE MODE'))


def _por_codigo(db, model, codigos):
    if not codigos:
        return {}
    itens = db.scalars(select(model).where(model.codigo.in_(codigos)).order_by(model.codigo)).all()
    return {item.codigo: item for item in itens}


def _aplicar(db, dados, operacoes):
    resultados = {(item["aba"], item["linha"]): item["resultado"] for item in operacoes}
    for item in sorted(dados["NUTRIENTES"], key=lambda x: (x["codigo"], x["linha"])):
        resultado = resultados.get(("NUTRIENTES", item["linha"]))
        if resultado == "CRIAR":
            db.add(Nutriente(codigo=item["codigo"], nome=item["nome"], unidade=item["unidade"]))
        elif resultado == "ATUALIZAR":
            atual = db.scalar(select(Nutriente).where(Nutriente.codigo == item["codigo"]))
            if item["nome"] is not None:
                atual.nome = item["nome"]
            if item["unidade"] is not None:
                atual.unidade = item["unidade"]
    db.flush()

    for item in sorted(dados["MATERIAS_PRIMAS"], key=lambda x: (x["codigo"], x["linha"])):
        resultado = resultados.get(("MATERIAS_PRIMAS", item["linha"]))
        if resultado == "CRIAR":
            db.add(MateriaPrima(codigo=item["codigo"], nome=item["nome"], ativa=item["ativa"]))
        elif resultado == "ATUALIZAR":
            atual = db.scalar(select(MateriaPrima).where(MateriaPrima.codigo == item["codigo"]))
            if item["nome"] is not None:
                atual.nome = item["nome"]
            if item["ativa"] is not None:
                atual.ativa = item["ativa"]
    db.flush()

    codigos_mp = {item["materia_prima_codigo"] for item in dados["COMPOSICAO_NUTRICIONAL"] + dados["PRECOS_MP"]}
    codigos_mp |= {item["codigo"] for item in dados["MATERIAS_PRIMAS"]}
    codigos_nutriente = {item["nutriente_codigo"] for item in dados["COMPOSICAO_NUTRICIONAL"]}
    codigos_nutriente |= {item["codigo"] for item in dados["NUTRIENTES"]}
    mps = _por_codigo(db, MateriaPrima, codigos_mp)
    nutrientes = _por_codigo(db, Nutriente, codigos_nutriente)

    for item in sorted(dados["COMPOSICAO_NUTRICIONAL"], key=lambda x: (x["materia_prima_codigo"], x["nutriente_codigo"])):
        resultado = resultados.get(("COMPOSICAO_NUTRICIONAL", item["linha"]))
        if resultado == "CRIAR":
            db.add(ComposicaoMateriaPrima(materia_prima_id=mps[item["materia_prima_codigo"]].id,
                                          nutriente_id=nutrientes[item["nutriente_codigo"]].id,
                                          valor=Decimal(item["valor"])))
        elif resultado == "ATUALIZAR":
            atual = db.scalar(select(ComposicaoMateriaPrima).where(
                ComposicaoMateriaPrima.materia_prima_id == mps[item["materia_prima_codigo"]].id,
                ComposicaoMateriaPrima.nutriente_id == nutrientes[item["nutriente_codigo"]].id))
            atual.valor = Decimal(item["valor"])
    db.flush()

    for item in sorted(dados["PRECOS_MP"], key=lambda x: (x["materia_prima_codigo"], x["vigencia_inicio"], x["linha"])):
        if resultados.get(("PRECOS_MP", item["linha"])) == "CRIAR":
            db.add(PrecoMateriaPrima(materia_prima_id=mps[item["materia_prima_codigo"]].id,
                                     preco_kg=Decimal(item["preco_kg"]),
                                     vigencia_inicio=date.fromisoformat(item["vigencia_inicio"]),
                                     vigencia_fim=date.fromisoformat(item["vigencia_fim"]) if item["vigencia_fim"] else None))
    db.flush()

    for item in sorted(dados["MATERIAS_PRIMAS"], key=lambda x: (x["codigo"], x["linha"])):
        if resultados.get(("MATERIAS_PRIMAS", item["linha"])) == "DESATIVAR":
            mps[item["codigo"]].ativa = False
    db.flush()


def _resultado(sessao, validacao, operacoes, confirmado_em):
    afetados = sorted({item["codigo"] for item in operacoes if item["resultado"] != "SEM_ALTERACAO"})
    contagens = {chave: sum(valores[chave] for aba, valores in validacao["resumo"].items() if aba != "GERAL")
                 for chave in ("criar", "atualizar", "desativar", "sem_alteracao")}
    resultado, _ = serializar_canonico({
        "sessao_id": str(sessao.uuid_publico), "arquivo_sha256": sessao.arquivo_sha256,
        "status": "CONFIRMADA", "confirmado_em": confirmado_em,
        "resumo": validacao["resumo"], "totais": contagens,
        "codigos_afetados": afetados[:MAX_CODIGOS_RESULTADO],
        "resultado_truncado": len(afetados) > MAX_CODIGOS_RESULTADO,
        "avisos": [d for d in validacao["diagnosticos"] if d["severidade"] == "AVISO"],
    })
    return resultado


def confirmar_sessao(db: Session, sessao_id, token):
    repo = ImportacaoCadastralRepository(db)
    sessao = repo.obter(sessao_id, bloquear=True)
    if sessao is None or not repo.token_valido(sessao, token):
        raise CredenciaisSessaoInvalidas()
    if sessao.status == "CONFIRMADA":
        return sessao.resultado
    if sessao.status == "EXPIRADA":
        raise ConflitoSessao("A sessão expirou; prepare uma nova importação.")
    if sessao.status == "FALHOU":
        raise ConflitoSessao("A sessão falhou e não pode ser reutilizada; prepare uma nova importação.")
    if repo.marcar_expirada(sessao):
        raise ConflitoSessao("A sessão expirou; prepare uma nova importação.")

    _bloquear_cadastros(db)
    conteudo = _xlsx_do_payload(sessao.payload_normalizado)
    validacao, internos = pre_validar_planilha_cadastral_completo(
        conteudo, "payload-canonico.xlsx", db, sessao.arquivo_sha256
    )
    if not validacao["valido_para_confirmacao"]:
        diagnosticos = validacao["diagnosticos"][:50]
        falha, _ = serializar_canonico({"tipo": "REVALIDACAO", "mensagem": "Os cadastros mudaram desde a preparação. Prepare uma nova sessão.", "diagnosticos": diagnosticos})
        sessao.status = "FALHOU"
        sessao.resultado = falha
        db.flush()
        raise RevalidacaoFalhou("Os cadastros mudaram desde a preparação; prepare uma nova sessão.")

    _aplicar(db, internos["dados"], internos["operacoes"])
    confirmado_em = db.scalar(select(func.clock_timestamp()))
    resultado = _resultado(sessao, validacao, internos["operacoes"], confirmado_em)
    sessao.status = "CONFIRMADA"
    sessao.confirmado_em = confirmado_em
    sessao.resultado = resultado
    db.flush()
    return resultado


def registrar_falha_tecnica(db: Session, sessao_id):
    repo = ImportacaoCadastralRepository(db)
    sessao = repo.obter(sessao_id, bloquear=True)
    if sessao is not None and sessao.status == "PENDENTE":
        sessao.status = "FALHOU"
        sessao.resultado = {"tipo": "ERRO_TECNICO", "mensagem": "A confirmação falhou sem aplicar alterações. Prepare uma nova sessão."}
        db.flush()
