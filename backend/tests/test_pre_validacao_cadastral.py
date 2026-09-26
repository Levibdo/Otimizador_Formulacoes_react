import io
import hashlib
from datetime import date

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from db.base import Base
from db.session import get_db
from models import ComposicaoMateriaPrima, MateriaPrima, Nutriente, PrecoMateriaPrima
from routers.importacoes_cadastrais import router
from services import pre_validacao_cadastral as modulo
from services.planilha_cadastral import gerar_template_cadastral
from services.pre_validacao_cadastral import pre_validar_planilha_cadastral


def planilha(mps=(), nutrientes=(), composicoes=(), precos=()):
    workbook = load_workbook(io.BytesIO(gerar_template_cadastral()))
    for aba in ("MATERIAS_PRIMAS", "NUTRIENTES", "COMPOSICAO_NUTRICIONAL", "PRECOS_MP"):
        ws = workbook[aba]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
    for linha in mps:
        workbook["MATERIAS_PRIMAS"].append(linha)
    for linha in nutrientes:
        workbook["NUTRIENTES"].append(linha)
    for linha in composicoes:
        workbook["COMPOSICAO_NUTRICIONAL"].append(linha)
    for linha in precos:
        workbook["PRECOS_MP"].append(linha)
    saida = io.BytesIO()
    workbook.save(saida)
    return saida.getvalue()


@pytest.fixture
def db():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    with Session(engine, expire_on_commit=False) as session:
        yield session
    engine.dispose()


def cadastrar_base(db, ativa=True, unidade="g/100 g", com_composicao=True, precos=()):
    nutriente = Nutriente(codigo="NUT_EXIST", nome="Nutriente existente", unidade=unidade)
    mp = MateriaPrima(codigo="MP_EXIST", nome="MP existente", ativa=ativa)
    db.add_all([mp, nutriente])
    db.flush()
    if com_composicao:
        db.add(ComposicaoMateriaPrima(materia_prima_id=mp.id, nutriente_id=nutriente.id, valor=10))
    for inicio, fim, valor in precos:
        db.add(PrecoMateriaPrima(materia_prima_id=mp.id, preco_kg=valor,
                                 vigencia_inicio=inicio, vigencia_fim=fim))
    db.commit()
    return mp, nutriente


def diagnosticos(resultado, trecho):
    return [item for item in resultado["diagnosticos"] if trecho in item["mensagem"]]


def test_arquivo_valido_com_criacoes_e_referencias_do_workbook_nao_grava(db):
    conteudo = planilha(
        mps=[("CRIAR", "MP_NOVA", "MP nova", "SIM")],
        nutrientes=[("CRIAR", "NUT_NOVO", "Nutriente novo", "g/100 g")],
        composicoes=[("CRIAR", "MP_NOVA", "NUT_NOVO", "12.345678")],
        precos=[("CRIAR", "MP_NOVA", "9.500000", "2026-01-01", None)],
    )
    resultado = pre_validar_planilha_cadastral(conteudo, "dados.xlsx", db)
    assert resultado["valido_para_confirmacao"] is True
    assert resultado["resumo"]["MATERIAS_PRIMAS"]["criar"] == 1
    assert resultado["resumo"]["NUTRIENTES"]["criar"] == 1
    assert resultado["resumo"]["COMPOSICAO_NUTRICIONAL"]["criar"] == 1
    assert resultado["resumo"]["PRECOS_MP"]["criar"] == 1
    assert len(resultado["sha256"]) == 64
    assert db.scalar(select(func.count()).select_from(MateriaPrima)) == 0
    assert not db.new and not db.dirty and not db.deleted


def test_referencias_do_banco_e_dados_identicos_ficam_sem_alteracao(db):
    cadastrar_base(db)
    conteudo = planilha(
        mps=[("ATUALIZAR", "MP_EXIST", "MP existente", "SIM")],
        nutrientes=[("ATUALIZAR", "NUT_EXIST", "Nutriente existente", "g/100 g")],
        composicoes=[("ATUALIZAR", "MP_EXIST", "NUT_EXIST", 10)],
    )
    resultado = pre_validar_planilha_cadastral(conteudo, "dados.xlsx", db)
    assert resultado["valido_para_confirmacao"] is True
    assert resultado["resumo"]["MATERIAS_PRIMAS"]["sem_alteracao"] == 1
    assert resultado["resumo"]["NUTRIENTES"]["sem_alteracao"] == 1
    assert resultado["resumo"]["COMPOSICAO_NUTRICIONAL"]["sem_alteracao"] == 1
    assert not db.new and not db.dirty and not db.deleted


def test_criacao_existente_atualizacao_inexistente_e_conflito_de_nome(db):
    cadastrar_base(db)
    conteudo = planilha(mps=[
        ("CRIAR", "MP_EXIST", "Outra MP", "SIM"),
        ("ATUALIZAR", "MP_INEXISTENTE", "Nome", "SIM"),
        ("CRIAR", "MP_OUTRA", "MP existente", "SIM"),
    ])
    resultado = pre_validar_planilha_cadastral(conteudo, "dados.xlsx", db)
    assert diagnosticos(resultado, "Código de matéria-prima já cadastrado")
    assert diagnosticos(resultado, "não encontrada para atualização")
    assert diagnosticos(resultado, "nome de matéria-prima já pertence".capitalize())
    assert resultado["valido_para_confirmacao"] is False


def test_desativacao_de_mp_e_mp_ja_inativa(db):
    cadastrar_base(db)
    resultado = pre_validar_planilha_cadastral(
        planilha(mps=[("DESATIVAR", "MP_EXIST", None, None)]), "dados.xlsx", db
    )
    assert resultado["resumo"]["MATERIAS_PRIMAS"]["desativar"] == 1

    db.get(MateriaPrima, 1).ativa = False
    db.commit()
    resultado = pre_validar_planilha_cadastral(
        planilha(mps=[("DESATIVAR", "MP_EXIST", None, None)]), "dados.xlsx", db
    )
    assert resultado["resumo"]["MATERIAS_PRIMAS"]["sem_alteracao"] == 1


def test_mudanca_de_unidade_com_composicao_existente_e_rejeitada(db):
    cadastrar_base(db)
    resultado = pre_validar_planilha_cadastral(
        planilha(nutrientes=[("ATUALIZAR", "NUT_EXIST", None, "mg/100 g")]),
        "dados.xlsx", db,
    )
    assert diagnosticos(resultado, "composições históricas")
    assert resultado["valido_para_confirmacao"] is False


def test_criacao_e_atualizacao_de_composicao_respeitam_existencia(db):
    cadastrar_base(db)
    resultado = pre_validar_planilha_cadastral(
        planilha(composicoes=[("CRIAR", "MP_EXIST", "NUT_EXIST", 11)]),
        "dados.xlsx", db,
    )
    assert diagnosticos(resultado, "Composição já cadastrada")

    resultado = pre_validar_planilha_cadastral(
        planilha(composicoes=[("ATUALIZAR", "MP_EXIST", "NUT_EXIST", 11)]),
        "dados.xlsx", db,
    )
    assert resultado["resumo"]["COMPOSICAO_NUTRICIONAL"]["atualizar"] == 1
    assert resultado["operacoes"][0]["campos_alterados"] == ["valor"]


def test_composicao_vazia_nao_e_convertida_em_zero(db):
    cadastrar_base(db)
    resultado = pre_validar_planilha_cadastral(
        planilha(composicoes=[("ATUALIZAR", "MP_EXIST", "NUT_EXIST", None)]),
        "dados.xlsx", db,
    )
    assert resultado["valido_para_confirmacao"] is False
    assert any(item["coluna"] == "VALOR" and "obrigatório ausente" in item["mensagem"]
               for item in resultado["diagnosticos"])
    assert resultado["resumo"]["COMPOSICAO_NUTRICIONAL"]["sem_alteracao"] == 0


def test_decimal_respeita_numeric_18_6_sem_arredondamento_silencioso(db):
    cadastrar_base(db)
    valido = pre_validar_planilha_cadastral(
        planilha(composicoes=[("ATUALIZAR", "MP_EXIST", "NUT_EXIST", "123456789012.123456")]),
        "dados.xlsx",
        db,
    )
    assert valido["valido_para_confirmacao"] is True

    escala_excedida = pre_validar_planilha_cadastral(
        planilha(composicoes=[("ATUALIZAR", "MP_EXIST", "NUT_EXIST", "1.1234567")]),
        "dados.xlsx",
        db,
    )
    precisao_excedida = pre_validar_planilha_cadastral(
        planilha(composicoes=[("ATUALIZAR", "MP_EXIST", "NUT_EXIST", "1234567890123.123456")]),
        "dados.xlsx",
        db,
    )
    assert diagnosticos(escala_excedida, "precisão suportada")
    assert diagnosticos(precisao_excedida, "precisão suportada")


def test_rejeita_composicao_e_preco_para_mp_inativa(db):
    cadastrar_base(db, ativa=False)
    resultado = pre_validar_planilha_cadastral(
        planilha(
            composicoes=[("ATUALIZAR", "MP_EXIST", "NUT_EXIST", 11)],
            precos=[("CRIAR", "MP_EXIST", 12, "2027-01-01", None)],
        ),
        "dados.xlsx", db,
    )
    assert diagnosticos(resultado, "Composição não pode usar matéria-prima inativa")
    assert diagnosticos(resultado, "Preço não pode usar matéria-prima inativa")
    assert resultado["valido_para_confirmacao"] is False


def test_preco_sem_conflito_e_periodos_contiguos_sao_validos(db):
    cadastrar_base(db, precos=[(date(2026, 1, 1), date(2026, 1, 31), 10)])
    conteudo = planilha(precos=[
        ("CRIAR", "MP_EXIST", 11, "2026-02-01", "2026-02-28"),
        ("CRIAR", "MP_EXIST", 12, "2026-03-01", None),
    ])
    resultado = pre_validar_planilha_cadastral(conteudo, "dados.xlsx", db)
    assert resultado["valido_para_confirmacao"] is True
    assert resultado["resumo"]["PRECOS_MP"]["criar"] == 2


def test_fronteira_inclusiva_na_mesma_data_conflita(db):
    cadastrar_base(db, precos=[(date(2026, 1, 1), date(2026, 1, 31), 10)])
    resultado = pre_validar_planilha_cadastral(
        planilha(precos=[("CRIAR", "MP_EXIST", 11, "2026-01-31", "2026-02-28")]),
        "dados.xlsx",
        db,
    )
    assert resultado["valido_para_confirmacao"] is False
    assert diagnosticos(resultado, "vigência existente iniciada em 2026-01-01")


def test_detecta_sobreposicao_com_banco_no_arquivo_e_periodo_aberto(db):
    cadastrar_base(db, precos=[(date(2026, 1, 1), None, 10)])
    resultado = pre_validar_planilha_cadastral(
        planilha(precos=[
            ("CRIAR", "MP_EXIST", 10, "2026-02-01", "2026-02-28"),
            ("CRIAR", "MP_EXIST", 12, "2026-02-15", "2026-03-01"),
        ]),
        "dados.xlsx", db,
    )
    assert len(diagnosticos(resultado, "vigência existente iniciada em 2026-01-01")) == 2
    assert len(diagnosticos(resultado, "do arquivo")) == 2
    assert resultado["valido_para_confirmacao"] is False


def test_precos_fora_de_ordem_detectam_todas_as_sobreposicoes(db):
    cadastrar_base(db, precos=())
    resultado = pre_validar_planilha_cadastral(
        planilha(precos=[
            ("CRIAR", "MP_EXIST", 13, "2026-03-01", "2026-04-30"),
            ("CRIAR", "MP_EXIST", 11, "2026-01-01", "2026-03-31"),
            ("CRIAR", "MP_EXIST", 12, "2026-02-01", None),
        ]),
        "dados.xlsx",
        db,
    )
    assert resultado["valido_para_confirmacao"] is False
    assert resultado["resumo"]["PRECOS_MP"]["erros"] == 6
    assert {item["linha"] for item in diagnosticos(resultado, "do arquivo")} == {2, 3, 4}


def test_referencia_a_criacao_invalida_no_workbook_e_rejeitada(db):
    cadastrar_base(db)
    resultado = pre_validar_planilha_cadastral(
        planilha(
            mps=[("CRIAR", "MP_EXIST", "Código duplicado", "SIM")],
            composicoes=[("CRIAR", "MP_EXIST", "NUT_EXIST", 12)],
            precos=[("CRIAR", "MP_EXIST", 5, "2027-01-01", None)],
        ),
        "dados.xlsx",
        db,
    )
    assert len(diagnosticos(resultado, "criação inválida no arquivo")) == 2
    assert resultado["resumo"]["COMPOSICAO_NUTRICIONAL"]["criar"] == 0
    assert resultado["resumo"]["PRECOS_MP"]["criar"] == 0


def test_aviso_isolado_mantem_resultado_valido(db, monkeypatch):
    original = modulo.parsear_planilha_cadastral

    def parser_com_aviso(*args, **kwargs):
        resultado = original(*args, **kwargs)
        resultado["diagnosticos"].append({
            "severidade": "AVISO", "aba": "GERAL", "linha": None,
            "coluna": None, "codigo": None, "mensagem": "Aviso cadastral seguro.",
        })
        return resultado

    monkeypatch.setattr(modulo, "parsear_planilha_cadastral", parser_com_aviso)
    resultado = pre_validar_planilha_cadastral(planilha(), "dados.xlsx", db)
    assert resultado["valido_para_confirmacao"] is True
    assert resultado["diagnosticos_total"] == 1
    assert resultado["resumo"]["GERAL"]["avisos"] == 1


def test_erro_estrutural_torna_resultado_invalido(db):
    resultado = pre_validar_planilha_cadastral(b"conteudo-invalido", "dados.xlsx", db)
    assert resultado["valido_para_confirmacao"] is False
    assert resultado["diagnosticos_total"] >= 1


def test_truncamento_de_operacoes_preserva_totais_e_classificacao(db, monkeypatch):
    conteudo = planilha(mps=[
        ("CRIAR", f"MP_{indice:03d}", f"MP fictícia {indice:03d}", "SIM")
        for indice in range(5)
    ])
    monkeypatch.setattr(modulo, "MAX_OPERACOES_RESPOSTA", 2)
    resultado = pre_validar_planilha_cadastral(conteudo, "dados.xlsx", db)
    assert resultado["valido_para_confirmacao"] is True
    assert resultado["operacoes_total"] == 5
    assert len(resultado["operacoes"]) == 2
    assert resultado["resumo"]["MATERIAS_PRIMAS"]["criar"] == 5
    assert resultado["resultado_truncado"] is True


def test_resultado_deterministico_e_diagnosticos_truncados(db, monkeypatch):
    conteudo = planilha(mps=[
        ("CRIAR", f"MP_{indice:03d}", "Mesmo nome", "SIM") for indice in range(8)
    ])
    monkeypatch.setattr(modulo, "MAX_DIAGNOSTICOS_RESPOSTA", 3)
    primeiro = pre_validar_planilha_cadastral(conteudo, "dados.xlsx", db)
    segundo = pre_validar_planilha_cadastral(conteudo, "dados.xlsx", db)
    assert primeiro == segundo
    assert primeiro["resultado_truncado"] is True
    assert len(primeiro["diagnosticos"]) == 3
    assert primeiro["diagnosticos_total"] > 3
    assert primeiro["resumo"]["MATERIAS_PRIMAS"]["erros"] == primeiro["diagnosticos_total"]
    assert primeiro["valido_para_confirmacao"] is False


def test_resultado_independe_da_ordem_retornada_pelo_banco(db, monkeypatch):
    cadastrar_base(db, precos=[
        (date(2026, 1, 1), date(2026, 1, 31), 10),
        (date(2026, 3, 1), date(2026, 3, 31), 12),
    ])
    conteudo = planilha(precos=[("CRIAR", "MP_EXIST", 11, "2026-01-15", "2026-03-15")])
    original = modulo._consultar_estado
    normal = pre_validar_planilha_cadastral(conteudo, "dados.xlsx", db)

    def estado_invertido(*args, **kwargs):
        mps, nutrientes, composicoes, precos, usos = original(*args, **kwargs)
        return list(reversed(mps)), list(reversed(nutrientes)), list(reversed(composicoes)), list(reversed(precos)), usos

    monkeypatch.setattr(modulo, "_consultar_estado", estado_invertido)
    invertido = pre_validar_planilha_cadastral(conteudo, "dados.xlsx", db)
    assert invertido == normal


def test_erro_controlado_nao_inutiliza_sessao(db):
    resultado = pre_validar_planilha_cadastral(b"invalido", "dados.xlsx", db)
    assert resultado["valido_para_confirmacao"] is False
    assert db.scalar(select(func.count()).select_from(MateriaPrima)) == 0


def test_endpoint_invalido_nao_expoe_detalhe_interno(db, monkeypatch):
    app = FastAPI()
    app.include_router(router)

    def override_get_db():
        yield db

    app.dependency_overrides[get_db] = override_get_db
    import routers.importacoes_cadastrais as modulo_router
    monkeypatch.setattr(
        modulo_router,
        "pre_validar_planilha_cadastral",
        lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("senha=segredo SQL interno")),
    )
    with TestClient(app, raise_server_exceptions=False) as client:
        resposta = client.post(
            "/api/v1/importacoes-cadastrais/validar",
            files={"arquivo": ("dados.xlsx", gerar_template_cadastral())},
        )
    assert resposta.status_code == 500
    assert resposta.json() == {"detail": "Não foi possível pré-validar a planilha cadastral."}
    assert "segredo" not in resposta.text


def test_endpoint_calcula_sha256_dos_bytes_exatos_recebidos(db):
    app = FastAPI()
    app.include_router(router)

    def override_get_db():
        yield db

    app.dependency_overrides[get_db] = override_get_db
    conteudo = planilha()
    with TestClient(app) as client:
        resposta = client.post(
            "/api/v1/importacoes-cadastrais/validar",
            files={"arquivo": ("dados.xlsx", conteudo)},
        )
    assert resposta.status_code == 200
    assert resposta.json()["sha256"] == hashlib.sha256(conteudo).hexdigest()
