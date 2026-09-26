import io
import zipfile
from datetime import date

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from routers.importacoes_cadastrais import router
from services import planilha_cadastral as modulo
from services.planilha_cadastral import ABAS, gerar_template_cadastral, parsear_planilha_cadastral


def alterar_template(mutacao):
    workbook = load_workbook(io.BytesIO(gerar_template_cadastral()))
    mutacao(workbook)
    saida = io.BytesIO()
    workbook.save(saida)
    return saida.getvalue()


def mensagens(resultado):
    return [item["mensagem"] for item in resultado["diagnosticos"]]


def test_template_ficticio_obedece_contrato_e_parser_normaliza():
    conteudo = gerar_template_cadastral()
    workbook = load_workbook(io.BytesIO(conteudo), data_only=False)
    assert workbook.sheetnames == list(ABAS)
    assert all(cell.data_type != "f" for ws in workbook for row in ws.iter_rows() for cell in row)

    resultado = parsear_planilha_cadastral(conteudo, "cadastros.xlsx")
    assert resultado["valido"] is True
    assert resultado["versao"] == "1.0"
    assert resultado["diagnosticos"] == []
    assert resultado["dados"]["MATERIAS_PRIMAS"][0] == {
        "linha": 2,
        "acao": "CRIAR",
        "codigo": "FICT_MP_AVEIA",
        "nome": "Aveia fictícia",
        "ativa": True,
    }
    assert resultado["dados"]["PRECOS_MP"][0]["vigencia_inicio"] == "2026-01-01"


def test_endpoint_baixa_template_sem_banco():
    app = FastAPI()
    app.include_router(router)
    resposta = TestClient(app).get("/api/v1/importacoes-cadastrais/template")
    assert resposta.status_code == 200
    assert resposta.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert resposta.headers["x-template-version"] == "1.0"
    assert resposta.headers["content-disposition"] == 'attachment; filename="template-cadastral-v1.0.xlsx"'
    assert "/" not in resposta.headers["content-disposition"]
    assert resposta.headers["x-content-type-options"] == "nosniff"
    rota = next(item for item in router.routes if item.path.endswith("/template"))
    assert rota.dependant.dependencies == []
    assert parsear_planilha_cadastral(resposta.content, "template.xlsx")["valido"]


def test_rejeita_extensao_conteudo_vazio_e_xlsx_corrompido():
    assert "somente .xlsx" in mensagens(parsear_planilha_cadastral(b"x", "dados.csv"))[0]
    assert "Arquivo vazio" in mensagens(parsear_planilha_cadastral(b"", "dados.xlsx"))[0]
    assert "corrompido" in mensagens(parsear_planilha_cadastral("não é zip".encode(), "dados.xlsx"))[0]


def test_rejeita_macro_antes_de_abrir_workbook():
    pacote = io.BytesIO(gerar_template_cadastral())
    with zipfile.ZipFile(pacote, "a") as arquivo:
        arquivo.writestr("xl/vbaProject.bin", "macro fictícia".encode())
    resultado = parsear_planilha_cadastral(pacote.getvalue(), "dados.xlsx")
    assert resultado["valido"] is False
    assert any("Macros não são permitidas" in texto for texto in mensagens(resultado))


def test_rejeita_versao_aba_ausente_e_aba_desconhecida():
    def mutacao(workbook):
        workbook["LEIA_ME"]["B2"] = "2.0"
        del workbook["NUTRIENTES"]
        workbook.create_sheet("INTRUSA")
        workbook.create_sheet("INTRUSA_2")

    resultado = parsear_planilha_cadastral(alterar_template(mutacao), "dados.xlsx")
    assert resultado["valido"] is False
    assert any(item["aba"] == "NUTRIENTES" and "obrigatória ausente" in item["mensagem"] for item in resultado["diagnosticos"])
    assert any(item["aba"] == "INTRUSA" and "desconhecida" in item["mensagem"] for item in resultado["diagnosticos"])
    assert any("Quantidade de abas" in item["mensagem"] for item in resultado["diagnosticos"])
    assert any("Versão do template" in item["mensagem"] for item in resultado["diagnosticos"])


def test_rejeita_cabecalhos_ausentes_repetidos_e_desconhecidos():
    def mutacao(workbook):
        ws = workbook["MATERIAS_PRIMAS"]
        ws["C1"] = "CODIGO"
        ws["E1"] = "CAMPO_INESPERADO"

    resultado = parsear_planilha_cadastral(alterar_template(mutacao), "dados.xlsx")
    assert any("repetido" in texto for texto in mensagens(resultado))
    assert any("obrigatório ausente" in texto for texto in mensagens(resultado))
    assert any("desconhecido" in texto for texto in mensagens(resultado))


def test_rejeita_acao_ausente_invalida_codigo_invalido_e_duplicidade():
    def mutacao(workbook):
        ws = workbook["MATERIAS_PRIMAS"]
        ws.append((None, "CODIGO INVALIDO", "Sem ação", "SIM"))
        ws.append(("REMOVER", "FICT_MP_AVEIA", "Duplicada", "SIM"))

    resultado = parsear_planilha_cadastral(alterar_template(mutacao), "dados.xlsx")
    assert any("Ação obrigatória" in texto for texto in mensagens(resultado))
    assert any("Ação inválida" in texto for texto in mensagens(resultado))
    assert any("Código inválido" in texto for texto in mensagens(resultado))
    assert any("Código duplicado" in texto for texto in mensagens(resultado))


def test_normaliza_codigo_e_acao_sem_gerar_identificador():
    def mutacao(workbook):
        ws = workbook["MATERIAS_PRIMAS"]
        ws["A2"] = " criar "
        ws["B2"] = " fict_mp_aveia "

    resultado = parsear_planilha_cadastral(alterar_template(mutacao), "dados.xlsx")
    item = resultado["dados"]["MATERIAS_PRIMAS"][0]
    assert item["acao"] == "CRIAR"
    assert item["codigo"] == "FICT_MP_AVEIA"


def test_codigo_textual_preserva_zeros_e_codigo_numerico_e_rejeitado():
    def textual(workbook):
        workbook["MATERIAS_PRIMAS"]["B2"] = " mp0001 "

    resultado = parsear_planilha_cadastral(alterar_template(textual), "dados.xlsx")
    assert resultado["dados"]["MATERIAS_PRIMAS"][0]["codigo"] == "MP0001"

    def numerico(workbook):
        workbook["MATERIAS_PRIMAS"]["B2"] = 123

    resultado = parsear_planilha_cadastral(alterar_template(numerico), "dados.xlsx")
    assert any("informado como texto" in texto for texto in mensagens(resultado))


def test_decimal_preserva_precisao_e_celula_vazia_nao_vira_zero():
    preciso = "0.123456789012345678"

    def decimal(workbook):
        workbook["COMPOSICAO_NUTRICIONAL"]["D2"] = preciso
        workbook["PRECOS_MP"]["C2"] = None

    resultado = parsear_planilha_cadastral(alterar_template(decimal), "dados.xlsx")
    assert resultado["dados"]["COMPOSICAO_NUTRICIONAL"][0]["valor"] == preciso
    assert resultado["dados"]["PRECOS_MP"][0]["preco_kg"] is None
    assert any(item["aba"] == "PRECOS_MP" and item["coluna"] == "PRECO_KG" and "obrigatório ausente" in item["mensagem"] for item in resultado["diagnosticos"])


def test_booleanos_aceitam_somente_textos_documentados():
    def mutacao(workbook):
        workbook["MATERIAS_PRIMAS"]["D2"] = True

    resultado = parsear_planilha_cadastral(alterar_template(mutacao), "dados.xlsx")
    assert any("texto SIM ou NÃO" in texto for texto in mensagens(resultado))


def test_acoes_sao_restritas_ao_que_o_schema_atual_suporta():
    def mutacao(workbook):
        workbook["NUTRIENTES"]["A2"] = "DESATIVAR"
        workbook["COMPOSICAO_NUTRICIONAL"]["A2"] = "DESATIVAR"
        workbook["PRECOS_MP"]["A2"] = "ATUALIZAR"

    resultado = parsear_planilha_cadastral(alterar_template(mutacao), "dados.xlsx")
    erros = [item for item in resultado["diagnosticos"] if "não é suportada com segurança" in item["mensagem"]]
    assert {item["aba"] for item in erros} == {"NUTRIENTES", "COMPOSICAO_NUTRICIONAL", "PRECOS_MP"}


def test_valida_referencias_internas_e_referencia_externa_como_aviso():
    def mutacao(workbook):
        workbook["MATERIAS_PRIMAS"]["A2"] = "DESATIVAR"
        workbook["COMPOSICAO_NUTRICIONAL"].append(("CRIAR", "MP_EXISTENTE", "NUT_EXISTENTE", 1))

    resultado = parsear_planilha_cadastral(alterar_template(mutacao), "dados.xlsx")
    assert any(item["severidade"] == "ERRO" and "desativada" in item["mensagem"] for item in resultado["diagnosticos"])
    avisos = [item for item in resultado["diagnosticos"] if item["severidade"] == "AVISO"]
    assert {item["codigo"] for item in avisos} >= {"MP_EXISTENTE", "NUT_EXISTENTE"}


def test_valida_numeros_booleanos_datas_e_periodo():
    def mutacao(workbook):
        workbook["MATERIAS_PRIMAS"]["D2"] = "talvez"
        workbook["COMPOSICAO_NUTRICIONAL"]["D2"] = -1
        ws = workbook["PRECOS_MP"]
        ws["C2"] = "não numérico"
        ws["D2"] = "21/09/2026"
        ws["E2"] = date(2025, 1, 1)

    resultado = parsear_planilha_cadastral(alterar_template(mutacao), "dados.xlsx")
    assert any("Booleano inválido" in texto for texto in mensagens(resultado))
    assert any("negativo" in texto for texto in mensagens(resultado))
    assert any("numérico inválido" in texto for texto in mensagens(resultado))
    assert any("Data inválida" in texto for texto in mensagens(resultado))


def test_rejeita_vigencia_invertida_e_alteracao_de_preco():
    def mutacao(workbook):
        ws = workbook["PRECOS_MP"]
        ws["A2"] = "ATUALIZAR"
        ws["D2"] = date(2026, 2, 1)
        ws["E2"] = date(2026, 1, 1)

    resultado = parsear_planilha_cadastral(alterar_template(mutacao), "dados.xlsx")
    assert any("não é suportada com segurança" in texto for texto in mensagens(resultado))
    assert any("anterior" in texto for texto in mensagens(resultado))


def test_rejeita_formula_com_diagnostico_localizado():
    def mutacao(workbook):
        workbook["COMPOSICAO_NUTRICIONAL"]["D2"] = "=1+1"

    resultado = parsear_planilha_cadastral(alterar_template(mutacao), "dados.xlsx")
    formula = next(item for item in resultado["diagnosticos"] if "Fórmulas" in item["mensagem"])
    assert (formula["aba"], formula["linha"], formula["coluna"]) == ("COMPOSICAO_NUTRICIONAL", 2, "D")


def test_aplica_limites_de_tamanho_dimensoes_e_entradas_zip(monkeypatch):
    conteudo = gerar_template_cadastral()
    monkeypatch.setattr(modulo, "MAX_ARQUIVO_BYTES", len(conteudo) - 1)
    assert "5 MiB" in mensagens(parsear_planilha_cadastral(conteudo, "dados.xlsx"))[0]

    monkeypatch.setattr(modulo, "MAX_ARQUIVO_BYTES", 5 * 1024 * 1024)
    monkeypatch.setattr(modulo, "MAX_LINHAS_POR_ABA", 1)
    resultado = parsear_planilha_cadastral(conteudo, "dados.xlsx")
    assert any("5.000 linhas" in texto for texto in mensagens(resultado))

    monkeypatch.setattr(modulo, "MAX_LINHAS_POR_ABA", 5_000)
    monkeypatch.setattr(modulo, "MAX_COLUNAS_POR_ABA", 1)
    resultado = parsear_planilha_cadastral(conteudo, "dados.xlsx")
    assert any("12 colunas" in texto for texto in mensagens(resultado))

    monkeypatch.setattr(modulo, "MAX_COLUNAS_POR_ABA", 12)
    monkeypatch.setattr(modulo, "MAX_CELULAS_TOTAL", 1)
    resultado = parsear_planilha_cadastral(conteudo, "dados.xlsx")
    assert any("100.000 células" in texto for texto in mensagens(resultado))

    pacote = io.BytesIO()
    with zipfile.ZipFile(pacote, "w") as arquivo:
        for indice in range(201):
            arquivo.writestr(f"entrada-{indice}", b"")
    monkeypatch.setattr(modulo, "MAX_LINHAS_POR_ABA", 5_000)
    assert "entradas internas" in mensagens(parsear_planilha_cadastral(pacote.getvalue(), "dados.xlsx"))[0]


def test_rejeita_taxa_de_compressao_compativel_com_zip_bomb():
    pacote = io.BytesIO()
    with zipfile.ZipFile(pacote, "w", compression=zipfile.ZIP_DEFLATED) as arquivo:
        arquivo.writestr("xl/repetitivo.xml", b"A" * 100_000)
    resultado = parsear_planilha_cadastral(pacote.getvalue(), "dados.xlsx")
    assert any("compressão insegura" in texto for texto in mensagens(resultado))


def test_dados_ausentes_nao_sao_inferidos_nem_removidos():
    def mutacao(workbook):
        ws = workbook["MATERIAS_PRIMAS"]
        ws["A2"] = "ATUALIZAR"
        ws["C2"] = None
        ws["D2"] = None
        workbook["COMPOSICAO_NUTRICIONAL"].delete_rows(2)

    resultado = parsear_planilha_cadastral(alterar_template(mutacao), "dados.xlsx")
    assert resultado["dados"]["MATERIAS_PRIMAS"][0]["nome"] is None
    assert resultado["dados"]["COMPOSICAO_NUTRICIONAL"] == []
    assert any("ATUALIZAR exige" in texto for texto in mensagens(resultado))


def test_resultado_do_parser_e_deterministico():
    conteudo = gerar_template_cadastral()
    primeiro = parsear_planilha_cadastral(conteudo, "dados.xlsx")
    segundo = parsear_planilha_cadastral(conteudo, "dados.xlsx")
    assert primeiro == segundo
