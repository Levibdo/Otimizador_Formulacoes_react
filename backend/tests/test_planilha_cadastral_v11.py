import io
import zipfile

import pytest
from openpyxl import load_workbook

from services.confirmacao_importacao_cadastral import _xlsx_do_payload
from services.planilha_cadastral import (
    ABAS_V1_0,
    ABAS_V1_1,
    gerar_template_cadastral,
    parsear_planilha_cadastral,
)


def alterar(versao="1.1", mutacao=lambda workbook: None):
    workbook = load_workbook(io.BytesIO(gerar_template_cadastral(versao)))
    mutacao(workbook)
    saida = io.BytesIO()
    workbook.save(saida)
    return saida.getvalue()


def erros(resultado):
    return [item["mensagem"] for item in resultado["diagnosticos"] if item["severidade"] == "ERRO"]


def limpar(ws):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)


def test_template_v11_tem_dez_abas_exatas_em_ordem_e_sem_conteudo_ativo():
    conteudo = gerar_template_cadastral()
    workbook = load_workbook(io.BytesIO(conteudo), data_only=False, keep_links=True)
    assert workbook.sheetnames == [
        "LEIA_ME", "MATERIAS_PRIMAS", "NUTRIENTES", "COMPOSICAO_NUTRICIONAL",
        "PRECOS_MP", "CATEGORIAS_PRODUTO", "COMPONENTES_REGULATORIOS",
        "COMPOSICAO_COMPONENTES_MP", "REGRAS_REGULATORIAS_MP",
        "REGRAS_REGULATORIAS_COMPONENTE",
    ]
    assert tuple(cell.value for cell in workbook["CATEGORIAS_PRODUTO"][1]) == (
        "ACAO", "CODIGO", "NOME", "DESCRICAO",
    )
    assert tuple(cell.value for cell in workbook["COMPOSICAO_COMPONENTES_MP"][1]) == (
        "ACAO", "MATERIA_PRIMA_CODIGO", "COMPONENTE_CODIGO", "DATA_REFERENCIA",
        "SITUACAO", "CONCENTRACAO", "FONTE", "OBSERVACAO",
    )
    assert all(cell.data_type != "f" for ws in workbook for row in ws.iter_rows() for cell in row)
    assert workbook._external_links == []
    with zipfile.ZipFile(io.BytesIO(conteudo)) as pacote:
        nomes = [item.filename.lower() for item in pacote.infolist()]
        assert not any("vbaproject" in nome or nome.startswith("xl/externallinks/") for nome in nomes)
    resultado = parsear_planilha_cadastral(conteudo, "template.xlsx")
    assert resultado["valido"] is True
    assert resultado["versao"] == "1.1"


def test_v10_permanece_aceito_e_rejeita_aba_regulatoria():
    assert parsear_planilha_cadastral(gerar_template_cadastral("1.0"), "v10.xlsx")["valido"]
    resultado = parsear_planilha_cadastral(
        alterar("1.0", lambda wb: wb.create_sheet("CATEGORIAS_PRODUTO")), "v10.xlsx"
    )
    assert not resultado["valido"]
    assert any("incompatível com o contrato 1.0" in texto or "desconhecida" in texto for texto in erros(resultado))


def test_v11_incompleto_versao_desconhecida_e_aba_adicional_sao_rejeitados():
    incompleto = parsear_planilha_cadastral(
        alterar(mutacao=lambda wb: wb.remove(wb["COMPONENTES_REGULATORIOS"])), "x.xlsx"
    )
    assert any("obrigatória ausente" in texto for texto in erros(incompleto))
    desconhecida = parsear_planilha_cadastral(
        alterar(mutacao=lambda wb: setattr(wb["LEIA_ME"]["B2"], "value", "9.9")), "x.xlsx"
    )
    assert any("desconhecida" in texto for texto in erros(desconhecida))
    adicional = parsear_planilha_cadastral(
        alterar(mutacao=lambda wb: wb.create_sheet("INTRUSA")), "x.xlsx"
    )
    assert any("Aba desconhecida" in texto for texto in erros(adicional))


def test_versao_ausente_e_cabecalhos_invalidos_sao_rejeitados():
    sem_versao = parsear_planilha_cadastral(
        alterar(mutacao=lambda wb: setattr(wb["LEIA_ME"]["A2"], "value", "OUTRA_CHAVE")), "x.xlsx"
    )
    assert any("VERSAO_TEMPLATE" in texto for texto in erros(sem_versao))

    def cabecalhos(wb):
        ws = wb["CATEGORIAS_PRODUTO"]
        ws["C1"] = "CODIGO"
        ws["E1"] = "INTRUSA"
    resultado = parsear_planilha_cadastral(alterar(mutacao=cabecalhos), "x.xlsx")
    assert any("repetido" in texto for texto in erros(resultado))
    assert any("obrigatório ausente" in texto for texto in erros(resultado))
    assert any("desconhecido" in texto for texto in erros(resultado))


def test_acoes_regulatorias_e_sem_alteracao_declarado():
    def mutar(wb):
        for aba in ("CATEGORIAS_PRODUTO", "COMPONENTES_REGULATORIOS"):
            limpar(wb[aba])
            wb[aba].append(("CRIAR", f"FICT_{aba[:8]}", "Nome", None))
            wb[aba].append(("ATUALIZAR", f"FICT_{aba[:8]}_2", "Novo", None))
            wb[aba].append(("DESATIVAR", f"FICT_{aba[:8]}_3", None, None))
        wb["COMPOSICAO_COMPONENTES_MP"]["A2"] = "ATUALIZAR"
        wb["REGRAS_REGULATORIAS_MP"]["A2"] = "SEM_ALTERACAO"
    resultado = parsear_planilha_cadastral(alterar(mutacao=mutar), "x.xlsx")
    assert any("ATUALIZAR não é suportada" in texto for texto in erros(resultado))
    assert any("Ação inválida" in texto for texto in erros(resultado))


@pytest.mark.parametrize("situacao,concentracao,valido", [
    ("INFORMADO", "0.123456789012345678", True),
    ("INFORMADO", None, False),
    ("AUSENTE_CONFIRMADO", "0", True),
    ("AUSENTE_CONFIRMADO", None, False),
    ("AUSENTE_CONFIRMADO", "1", False),
    ("DESCONHECIDO", None, True),
    ("DESCONHECIDO", "0", False),
])
def test_coerencia_situacao_concentracao_e_decimal_exato(situacao, concentracao, valido):
    def mutar(wb):
        ws = wb["COMPOSICAO_COMPONENTES_MP"]
        ws["E2"], ws["F2"] = situacao, concentracao
    resultado = parsear_planilha_cadastral(alterar(mutacao=mutar), "x.xlsx")
    item = resultado["dados"]["COMPOSICAO_COMPONENTES_MP"][0]
    assert (not any("concentração" in texto.lower() for texto in erros(resultado))) is valido
    if valido and concentracao is not None:
        assert item["concentracao"] == concentracao


@pytest.mark.parametrize("tratamento,minimo,maximo,valido,maximo_esperado", [
    ("PERMITIDA", None, None, True, None),
    ("LIMITADA", None, "25.500000", True, "25.500000"),
    ("LIMITADA", None, None, False, None),
    ("OBRIGATORIA", "1", None, True, None),
    ("OBRIGATORIA", "0", None, False, None),
    ("PROIBIDA", None, None, True, "0"),
    ("PROIBIDA", "0", "0", True, "0"),
    ("PROIBIDA", "1", "0", False, "0"),
])
def test_tratamentos_e_normalizacao(tratamento, minimo, maximo, valido, maximo_esperado):
    def mutar(wb):
        ws = wb["REGRAS_REGULATORIAS_COMPONENTE"]
        ws["D2"], ws["E2"], ws["F2"] = tratamento, minimo, maximo
    resultado = parsear_planilha_cadastral(alterar(mutacao=mutar), "x.xlsx")
    item = resultado["dados"]["REGRAS_REGULATORIAS_COMPONENTE"][0]
    relevantes = [t for t in erros(resultado) if tratamento in t or "mínimo" in t.lower() or "máximo" in t.lower()]
    assert (not relevantes) is valido
    if valido:
        assert item["maximo"] == maximo_esperado
        assert item["unidade"] == "%" and item["base"] == "MASSA_MASSA"


def test_minimo_maior_que_maximo_vigencia_invertida_e_data_nao_textual():
    def mutar(wb):
        ws = wb["REGRAS_REGULATORIAS_MP"]
        ws["E2"], ws["F2"] = "30", "20"
        ws["I2"], ws["J2"] = "2026-02-01", "2026-01-01"
        wb["COMPOSICAO_COMPONENTES_MP"]["D2"] = 46023
    resultado = parsear_planilha_cadastral(alterar(mutacao=mutar), "x.xlsx")
    assert any("mínimo não pode superar" in texto for texto in erros(resultado))
    assert any("Vigência final anterior" in texto for texto in erros(resultado))
    assert any("texto no formato" in texto for texto in erros(resultado))


def test_unicode_e_zeros_iniciais_sao_preservados():
    def mutar(wb):
        wb["CATEGORIAS_PRODUTO"]["B2"] = " fict_0001 "
        wb["CATEGORIAS_PRODUTO"]["C2"] = "Fórmula pediátrica – α"
    resultado = parsear_planilha_cadastral(alterar(mutacao=mutar), "x.xlsx")
    item = resultado["dados"]["CATEGORIAS_PRODUTO"][0]
    assert item["codigo"] == "FICT_0001"
    assert item["nome"] == "Fórmula pediátrica – α"


def test_round_trip_canonico_v10_e_v11():
    for versao in ("1.0", "1.1"):
        primeiro = parsear_planilha_cadastral(gerar_template_cadastral(versao), "x.xlsx")
        payload = {"versao": versao, "dados": primeiro["dados"]}
        segundo = parsear_planilha_cadastral(_xlsx_do_payload(payload), "x.xlsx")
        assert segundo["valido"]
        assert segundo["versao"] == versao
        assert segundo["dados"] == primeiro["dados"]


def test_rejeita_link_externo_no_pacote():
    pacote = io.BytesIO(gerar_template_cadastral())
    with zipfile.ZipFile(pacote, "a") as arquivo:
        arquivo.writestr("xl/externalLinks/externalLink1.xml", b"<externalLink/>")
    resultado = parsear_planilha_cadastral(pacote.getvalue(), "x.xlsx")
    assert any("Links externos" in texto for texto in erros(resultado))


@pytest.mark.parametrize("valor", [1.1, True, "=1.1"])
def test_versao_deve_ser_texto_literal_canonico(valor):
    resultado = parsear_planilha_cadastral(
        alterar(mutacao=lambda wb: setattr(wb["LEIA_ME"]["B2"], "value", valor)), "x.xlsx"
    )
    assert not resultado["valido"]
    assert any(
        "texto literal" in texto or "desconhecida" in texto or "Fórmulas" in texto
        for texto in erros(resultado)
    )


def test_versao_textual_ignora_apenas_espacos_externos():
    resultado = parsear_planilha_cadastral(
        alterar(mutacao=lambda wb: setattr(wb["LEIA_ME"]["B2"], "value", " 1.1 ")), "x.xlsx"
    )
    assert resultado["valido"]
    assert resultado["versao"] == "1.1"


def test_limite_absoluto_e_aba_oculta_sao_rejeitados_antes_do_schema():
    def exceder(wb):
        wb["LEIA_ME"]["B2"] = 9.9
        wb.create_sheet("INTRUSA")
    excesso = parsear_planilha_cadastral(alterar(mutacao=exceder), "x.xlsx")
    assert any("limite absoluto de 10 abas" in texto for texto in erros(excesso))

    oculto = parsear_planilha_cadastral(
        alterar(mutacao=lambda wb: setattr(wb["CATEGORIAS_PRODUTO"], "sheet_state", "hidden")),
        "x.xlsx",
    )
    assert any("ocultas" in texto for texto in erros(oculto))


def test_rejeita_entrada_zip_duplicada():
    pacote = io.BytesIO(gerar_template_cadastral())
    with zipfile.ZipFile(pacote, "a") as arquivo:
        nome = "docProps/app.xml"
        arquivo.writestr(nome, arquivo.read(nome))
    resultado = parsear_planilha_cadastral(pacote.getvalue(), "x.xlsx")
    assert any("duplicadas" in texto for texto in erros(resultado))


def test_xlsx_do_payload_exige_versao_persistida_valida():
    with pytest.raises(ValueError, match="Versão persistida"):
        _xlsx_do_payload({"dados": {}})
    with pytest.raises(ValueError, match="Versão persistida"):
        _xlsx_do_payload({"versao": "9.9", "dados": {}})


def test_round_trip_v11_completo_preserva_semantica():
    def preencher(wb):
        for aba in ABAS_V1_1:
            if aba != "LEIA_ME":
                limpar(wb[aba])
        wb["MATERIAS_PRIMAS"].append(("CRIAR", "MP_0001", "Matéria-prima α", "SIM"))
        wb["NUTRIENTES"].append(("CRIAR", "NUT_0001", "Nutriente ç", "g/100 g"))
        wb["COMPOSICAO_NUTRICIONAL"].append(("CRIAR", "MP_0001", "NUT_0001", "1.230000"))
        wb["PRECOS_MP"].append(("CRIAR", "MP_0001", "0.000000", "2026-01-01", None))
        wb["CATEGORIAS_PRODUTO"].append(("CRIAR", "CAT_0001", "Categoria Ω", None))
        wb["COMPONENTES_REGULATORIOS"].append(("CRIAR", "COMP_0001", "Componente β", None))
        wb["COMPOSICAO_COMPONENTES_MP"].append(
            ("CRIAR", "MP_0001", "COMP_0001", "2026-01-01", "INFORMADO", "1.234567", "Fonte", None)
        )
        wb["COMPOSICAO_COMPONENTES_MP"].append(
            ("CRIAR", "MP_0001", "COMP_0001", "2026-01-02", "DESCONHECIDO", None, None, None)
        )
        wb["COMPOSICAO_COMPONENTES_MP"].append(
            ("CRIAR", "MP_0001", "COMP_0001", "2026-01-03", "AUSENTE_CONFIRMADO", "0", None, None)
        )
        wb["REGRAS_REGULATORIAS_MP"].append(
            ("CRIAR", "CAT_0001", "MP_0001", "LIMITADA", None, "10.000000",
             "Justificativa MP", None, None, None)
        )
        wb["REGRAS_REGULATORIAS_COMPONENTE"].append(
            ("CRIAR", "CAT_0001", "COMP_0001", "PROIBIDA", None, None,
             "Justificativa componente", None, "2026-01-01", None)
        )
    primeiro = parsear_planilha_cadastral(alterar(mutacao=preencher), "x.xlsx")
    assert primeiro["valido"], primeiro["diagnosticos"]
    reconstruido = _xlsx_do_payload({"versao": "1.1", "dados": primeiro["dados"]})
    segundo = parsear_planilha_cadastral(reconstruido, "x.xlsx")
    assert segundo["valido"], segundo["diagnosticos"]
    assert segundo["dados"] == primeiro["dados"]
    composicoes = segundo["dados"]["COMPOSICAO_COMPONENTES_MP"]
    assert [item["concentracao"] for item in composicoes] == ["1.234567", None, "0"]
    assert segundo["dados"]["REGRAS_REGULATORIAS_MP"][0]["vigencia_inicio"] is None


def test_rejeita_relacionamento_externo_com_xml_valido():
    origem = io.BytesIO(gerar_template_cadastral())
    destino = io.BytesIO()
    with zipfile.ZipFile(origem) as entrada, zipfile.ZipFile(destino, "w") as saida:
        for item in entrada.infolist():
            conteudo = entrada.read(item.filename)
            if item.filename == "xl/_rels/workbook.xml.rels":
                conteudo = conteudo.replace(
                    b"</Relationships>",
                    b"<Relationship Id='rIdExternal' Type='urn:test' Target='https://example.invalid' TargetMode='External'/></Relationships>",
                )
            saida.writestr(item, conteudo)
    resultado = parsear_planilha_cadastral(destino.getvalue(), "x.xlsx")
    assert any("Links externos" in texto for texto in erros(resultado))


@pytest.mark.parametrize("nome", ["xl/worksheets/..", "C:/xl/workbook.xml"])
def test_rejeita_caminho_zip_inseguro(nome):
    pacote = io.BytesIO(gerar_template_cadastral())
    with zipfile.ZipFile(pacote, "a") as arquivo:
        arquivo.writestr(nome, b"conteudo")
    resultado = parsear_planilha_cadastral(pacote.getvalue(), "x.xlsx")
    assert any("Estrutura interna inválida" in texto for texto in erros(resultado))
