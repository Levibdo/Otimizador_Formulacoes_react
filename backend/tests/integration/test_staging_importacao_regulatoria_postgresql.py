import io
from hashlib import sha256
from uuid import uuid4

from openpyxl import load_workbook
import pytest
from sqlalchemy import func, select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from models import (
    ComposicaoMateriaPrima, MateriaPrima, Nutriente, PrecoMateriaPrima,
    SessaoImportacaoCadastral,
)
from models.regulatorio import (
    CategoriaProduto, ComponenteRegulatorio, ComposicaoComponenteMP,
    RegraRegulatoria,
)
from repositories.importacao_cadastral_repository import ImportacaoCadastralRepository
from services.confirmacao_importacao_cadastral import _xlsx_do_payload
from services.planilha_cadastral import gerar_template_cadastral, parsear_planilha_cadastral


ABAS = (
    "MATERIAS_PRIMAS", "NUTRIENTES", "COMPOSICAO_NUTRICIONAL", "PRECOS_MP",
    "CATEGORIAS_PRODUTO", "COMPONENTES_REGULATORIOS",
    "COMPOSICAO_COMPONENTES_MP", "REGRAS_REGULATORIAS_MP",
    "REGRAS_REGULATORIAS_COMPONENTE",
)


def planilha_v11_pg():
    workbook = load_workbook(io.BytesIO(gerar_template_cadastral("1.1")))
    for aba in ABAS:
        ws = workbook[aba]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
    workbook["MATERIAS_PRIMAS"].append(("CRIAR", "PG_V11_MP", "MP v1.1", "SIM"))
    workbook["CATEGORIAS_PRODUTO"].append(("CRIAR", "PG_V11_CAT", "Categoria v1.1", None))
    workbook["COMPONENTES_REGULATORIOS"].append(("CRIAR", "PG_V11_COMP", "Componente v1.1", None))
    workbook["COMPOSICAO_COMPONENTES_MP"].append(
        ("CRIAR", "PG_V11_MP", "PG_V11_COMP", "2026-01-01",
         "INFORMADO", "1.000000", None, None)
    )
    workbook["REGRAS_REGULATORIAS_MP"].append(
        ("CRIAR", "PG_V11_CAT", "PG_V11_MP", "LIMITADA", None,
         "10.000000", "Teste PG", None, None, None)
    )
    workbook["REGRAS_REGULATORIAS_COMPONENTE"].append(
        ("CRIAR", "PG_V11_CAT", "PG_V11_COMP", "PROIBIDA", None,
         None, "Teste PG", None, None, None)
    )
    saida = io.BytesIO()
    workbook.save(saida)
    return saida.getvalue()


def snapshot_cadastral(engine):
    tabelas = (
        "materias_primas", "nutrientes", "composicoes_materias_primas",
        "precos_materias_primas", "categorias_produto",
        "componentes_regulatorios", "composicoes_componentes_mp",
        "regras_regulatorias",
    )
    with engine.connect() as conn:
        linhas = {
            tabela: tuple(conn.execute(text(
                f'SELECT id, xmin::text FROM "{tabela}" ORDER BY id'
            )).all())
            for tabela in tabelas
        }
        sequencias = tuple(conn.execute(text(
            "SELECT sequencename, last_value FROM pg_sequences "
            "WHERE schemaname=current_schema() "
            "AND sequencename NOT LIKE 'sessoes_importacao_cadastral%' "
            "ORDER BY sequencename"
        )).all())
    return linhas, sequencias


def preparar(pg_client, conteudo=None):
    conteudo = conteudo or planilha_v11_pg()
    resposta = pg_client.post(
        "/api/v1/importacoes-cadastrais/preparar",
        files={"arquivo": ("regulatorio.xlsx", conteudo)},
    )
    assert resposta.status_code == 201, resposta.text
    return conteudo, resposta.json()


def test_staging_v11_postgresql_modifica_somente_sessao(pg_client, postgres_app):
    _, engine, _ = postgres_app
    antes = snapshot_cadastral(engine)
    conteudo, preparada = preparar(pg_client)
    assert preparada["arquivo_sha256"] == sha256(conteudo).hexdigest()
    assert preparada["operacoes_total"] == 6
    assert snapshot_cadastral(engine) == antes

    with Session(engine) as db:
        sessoes = list(db.scalars(select(SessaoImportacaoCadastral)))
        assert len(sessoes) == 1
        sessao = sessoes[0]
        assert sessao.versao_contrato == "1.1"
        assert sessao.status == "PENDENTE"
        assert sessao.token_hash != preparada["token_confirmacao"]
        assert preparada["token_confirmacao"] not in str(sessao.payload_normalizado)
        payload_jsonb = sessao.payload_normalizado
        assert payload_jsonb["versao"] == sessao.versao_contrato == "1.1"
        assert set(payload_jsonb["dados"]) == set(ABAS)
        assert "LEIA_ME" not in payload_jsonb["dados"]
        assert payload_jsonb["operacoes"][-1]["aba"] == "REGRAS_REGULATORIAS_COMPONENTE"
        reconstruido = _xlsx_do_payload(payload_jsonb)
        assert load_workbook(io.BytesIO(reconstruido)).sheetnames == [
            "LEIA_ME", *ABAS,
        ]
        reparsed = parsear_planilha_cadastral(reconstruido, "jsonb-v11.xlsx")
        assert reparsed["valido"], reparsed["diagnosticos"]
        assert reparsed["dados"] == payload_jsonb["dados"]

    consulta = pg_client.get(
        f'/api/v1/importacoes-cadastrais/{preparada["sessao_id"]}'
    )
    assert consulta.status_code == 200
    assert set(consulta.json()) == {
        "sessao_id", "status", "arquivo_sha256", "versao", "resumo",
        "operacoes_total", "avisos", "criado_em", "expira_em",
    }
    assert "token" not in consulta.text.lower()
    assert "payload" not in consulta.text.lower()


def test_credenciais_invalidas_v11_nao_alteram_estado(pg_client, postgres_app):
    _, engine, _ = postgres_app
    antes = snapshot_cadastral(engine)
    _, preparada = preparar(pg_client)
    caminho = f'/api/v1/importacoes-cadastrais/{preparada["sessao_id"]}/confirmar'

    invalida = pg_client.post(caminho, json={"token": "incorreto"})
    ausente = pg_client.post(
        f"/api/v1/importacoes-cadastrais/{uuid4()}/confirmar",
        json={"token": preparada["token_confirmacao"]},
    )
    assert invalida.status_code == ausente.status_code == 404
    assert invalida.json() == ausente.json()
    assert snapshot_cadastral(engine) == antes
    with Session(engine) as db:
        sessao = db.scalar(select(SessaoImportacaoCadastral))
        assert sessao.status == "PENDENTE"
        assert sessao.resultado is None
        assert sessao.confirmado_em is None


def test_duas_preparacoes_v11_postgresql_sao_independentes(pg_client, postgres_app):
    _, engine, _ = postgres_app
    conteudo = planilha_v11_pg()
    _, primeira = preparar(pg_client, conteudo)
    _, segunda = preparar(pg_client, conteudo)
    assert primeira["sessao_id"] != segunda["sessao_id"]
    assert primeira["token_confirmacao"] != segunda["token_confirmacao"]
    assert primeira["arquivo_sha256"] == segunda["arquivo_sha256"]
    with Session(engine) as db:
        sessoes = list(db.scalars(
            select(SessaoImportacaoCadastral).order_by(
                SessaoImportacaoCadastral.id
            )
        ))
        assert len(sessoes) == 2
        assert sessoes[0].token_hash != sessoes[1].token_hash
        assert sessoes[0].payload_normalizado == sessoes[1].payload_normalizado
        assert sessoes[0].resumo == sessoes[1].resumo
        assert sessoes[0].arquivo_sha256 == sessoes[1].arquivo_sha256


def test_rollback_de_flush_na_preparacao_v11_postgresql(
    pg_client, postgres_app, monkeypatch
):
    _, engine, _ = postgres_app
    original = ImportacaoCadastralRepository.criar

    def falhar(self, **campos):
        original(self, **campos)
        raise RuntimeError("falha controlada após flush")

    monkeypatch.setattr(ImportacaoCadastralRepository, "criar", falhar)
    resposta = pg_client.post(
        "/api/v1/importacoes-cadastrais/preparar",
        files={"arquivo": ("regulatorio.xlsx", planilha_v11_pg())},
    )
    assert resposta.status_code == 500
    assert "segredo" not in resposta.text
    with Session(engine) as db:
        assert db.scalar(
            select(func.count()).select_from(SessaoImportacaoCadastral)
        ) == 0


def test_rollback_de_commit_na_preparacao_v11_postgresql(
    pg_client, postgres_app, monkeypatch
):
    _, engine, _ = postgres_app
    original = Session.commit
    chamadas = 0

    def falhar_primeiro(self):
        nonlocal chamadas
        chamadas += 1
        if chamadas == 1:
            raise RuntimeError("falha controlada antes do commit")
        return original(self)

    monkeypatch.setattr(Session, "commit", falhar_primeiro)
    resposta = pg_client.post(
        "/api/v1/importacoes-cadastrais/preparar",
        files={"arquivo": ("regulatorio.xlsx", planilha_v11_pg())},
    )
    assert resposta.status_code == 500
    assert "segredo" not in resposta.text
    with Session(engine) as db:
        assert db.scalar(
            select(func.count()).select_from(SessaoImportacaoCadastral)
        ) == 0


def test_trigger_protege_conteudo_da_sessao_v11(pg_client, postgres_app):
    _, engine, _ = postgres_app
    _, preparada = preparar(pg_client)
    mutacoes = (
        "uuid_publico = gen_random_uuid()",
        "token_hash = repeat('b', 64)",
        "arquivo_sha256 = repeat('0', 64)",
        "versao_contrato = '1.0'",
        "payload_normalizado = '{}'::jsonb",
        "resumo = '{}'::jsonb",
        "avisos = '[{}]'::jsonb",
        "criado_em = criado_em - interval '1 second'",
        "expira_em = expira_em + interval '1 second'",
    )
    for atribuicao in mutacoes:
        with pytest.raises(IntegrityError, match="imutável"):
            with engine.begin() as conn:
                conn.execute(text(
                    f"UPDATE sessoes_importacao_cadastral SET {atribuicao} "
                    "WHERE uuid_publico=:id"
                ), {"id": preparada["sessao_id"]})
    with pytest.raises(IntegrityError, match="removida"):
        with engine.begin() as conn:
            conn.execute(text(
                "DELETE FROM sessoes_importacao_cadastral "
                "WHERE uuid_publico=:id"
            ), {"id": preparada["sessao_id"]})
