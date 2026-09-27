import io
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier, Event
from datetime import date
from decimal import Decimal

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
from sqlalchemy import event, func, select, text
from sqlalchemy.exc import OperationalError
from sqlalchemy.orm import Session

from models import (
    ComposicaoMateriaPrima, MateriaPrima, Nutriente, PrecoMateriaPrima,
    SessaoImportacaoCadastral,
)
from models.regulatorio import (
    CategoriaProduto, ComponenteRegulatorio, ComposicaoComponenteMP,
    RegraRegulatoria,
)
from services.planilha_cadastral import gerar_template_cadastral


ABAS = (
    "MATERIAS_PRIMAS", "NUTRIENTES", "COMPOSICAO_NUTRICIONAL", "PRECOS_MP",
    "CATEGORIAS_PRODUTO", "COMPONENTES_REGULATORIOS",
    "COMPOSICAO_COMPONENTES_MP", "REGRAS_REGULATORIAS_MP",
    "REGRAS_REGULATORIAS_COMPONENTE",
)
LOCKS_V11 = [
    "nutrientes", "materias_primas", "composicoes_materias_primas",
    "precos_materias_primas", "categorias_produto",
    "componentes_regulatorios", "composicoes_componentes_mp",
    "regras_regulatorias",
]


def montar_planilha(**linhas):
    workbook = load_workbook(io.BytesIO(gerar_template_cadastral("1.1")))
    for aba in ABAS:
        ws = workbook[aba]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        for linha in linhas.get(aba, ()):
            ws.append(linha)
    saida = io.BytesIO()
    workbook.save(saida)
    return saida.getvalue()


def preparar(client, conteudo):
    resposta = client.post(
        "/api/v1/importacoes-cadastrais/preparar",
        files={"arquivo": ("regulatorio.xlsx", conteudo)},
    )
    assert resposta.status_code == 201, resposta.text
    return resposta.json()


def confirmar(client, preparada):
    return client.post(
        f'/api/v1/importacoes-cadastrais/{preparada["sessao_id"]}/confirmar',
        json={"token": preparada["token_confirmacao"]},
    )


def semear_estado(engine):
    with Session(engine) as db, db.begin():
        mp_upd = MateriaPrima(codigo="PG4_MP_UPD", nome="MP antiga", ativa=True)
        mp_des = MateriaPrima(codigo="PG4_MP_DES", nome="MP desativar", ativa=True)
        mp_regra = MateriaPrima(codigo="PG4_MP_REGRA", nome="MP regra", ativa=True)
        cat_upd = CategoriaProduto(codigo="PG4_CAT_UPD", nome="Cat antiga", descricao="Antiga")
        cat_des = CategoriaProduto(codigo="PG4_CAT_DES", nome="Cat desativar")
        cat_regra = CategoriaProduto(codigo="PG4_CAT_REGRA", nome="Cat regra")
        comp_upd = ComponenteRegulatorio(codigo="PG4_COMP_UPD", nome="Comp antigo")
        comp_des = ComponenteRegulatorio(codigo="PG4_COMP_DES", nome="Comp desativar")
        comp_regra = ComponenteRegulatorio(codigo="PG4_COMP_REGRA", nome="Comp regra")
        db.add_all([
            mp_upd, mp_des, mp_regra, cat_upd, cat_des, cat_regra,
            comp_upd, comp_des, comp_regra,
        ])
        db.flush()
        composicao = ComposicaoComponenteMP(
            materia_prima_id=mp_regra.id, componente_id=comp_regra.id,
            data_referencia=date(2025, 1, 1), situacao="INFORMADO",
            concentracao=Decimal("2.000000"), fonte="Histórica",
            observacao="Preservar", ativo=True,
        )
        regra_atualizar = RegraRegulatoria(
            categoria_id=cat_regra.id, tipo_alvo="MATERIA_PRIMA",
            materia_prima_id=mp_regra.id, componente_id=None,
            tratamento="LIMITADA", minimo=None, maximo=Decimal("10"),
            unidade="%", base="MASSA_MASSA", justificativa="Original",
            referencia_normativa="RDC antiga", revisao=1,
            regra_anterior_id=None, ativa=True,
            vigencia_inicio=date(2026, 1, 1),
            vigencia_fim=date(2026, 12, 31),
        )
        regra_desativar = RegraRegulatoria(
            categoria_id=cat_regra.id, tipo_alvo="COMPONENTE",
            materia_prima_id=None, componente_id=comp_regra.id,
            tratamento="PROIBIDA", minimo=None, maximo=Decimal("0"),
            unidade="%", base="MASSA_MASSA", justificativa="Desativar",
            referencia_normativa=None, revisao=1,
            regra_anterior_id=None, ativa=True,
            vigencia_inicio=date(2025, 1, 1),
            vigencia_fim=date(2025, 12, 31),
        )
        db.add_all([composicao, regra_atualizar, regra_desativar])
        db.flush()
        ids = {
            "composicao_id": composicao.id,
            "regra_atualizar_id": regra_atualizar.id,
            "regra_desativar_id": regra_desativar.id,
        }
    return ids


def planilha_completa():
    return montar_planilha(
        MATERIAS_PRIMAS=[
            ("CRIAR", "PG4_MP_NOVA", "MP nova", "SIM"),
            ("ATUALIZAR", "PG4_MP_UPD", "MP atualizada", None),
            ("DESATIVAR", "PG4_MP_DES", None, None),
        ],
        NUTRIENTES=[
            ("CRIAR", "PG4_NUT_NOVO", "Nutriente novo", "g/100 g"),
        ],
        COMPOSICAO_NUTRICIONAL=[
            ("CRIAR", "PG4_MP_NOVA", "PG4_NUT_NOVO", "12.340000"),
        ],
        PRECOS_MP=[
            ("CRIAR", "PG4_MP_NOVA", "9.870000", "2026-01-01", None),
        ],
        CATEGORIAS_PRODUTO=[
            ("CRIAR", "PG4_CAT_NOVA", "Categoria nova", "Unicode çã"),
            ("ATUALIZAR", "PG4_CAT_UPD", None, "Descrição atualizada"),
            ("DESATIVAR", "PG4_CAT_DES", None, None),
        ],
        COMPONENTES_REGULATORIOS=[
            ("CRIAR", "PG4_COMP_NOVO", "Componente novo", None),
            ("ATUALIZAR", "PG4_COMP_UPD", "Componente atualizado", None),
            ("DESATIVAR", "PG4_COMP_DES", None, None),
        ],
        COMPOSICAO_COMPONENTES_MP=[
            ("CRIAR", "PG4_MP_NOVA", "PG4_COMP_NOVO", "2026-01-01",
             "INFORMADO", "1.234567", "Fonte", None),
            ("CRIAR", "PG4_MP_NOVA", "PG4_COMP_NOVO", "2026-01-02",
             "DESCONHECIDO", None, None, None),
            ("CRIAR", "PG4_MP_NOVA", "PG4_COMP_NOVO", "2026-01-03",
             "AUSENTE_CONFIRMADO", "0", None, None),
            ("DESATIVAR", "PG4_MP_REGRA", "PG4_COMP_REGRA", "2025-01-01",
             None, None, None, None),
        ],
        REGRAS_REGULATORIAS_MP=[
            ("CRIAR", "PG4_CAT_NOVA", "PG4_MP_NOVA", "LIMITADA", None,
             "15.000000", "Nova MP", None, None, None),
            ("ATUALIZAR", "PG4_CAT_REGRA", "PG4_MP_REGRA", "LIMITADA", None,
             "20.000000", "Revisada", "RDC nova", "2026-01-01", "2026-12-31"),
        ],
        REGRAS_REGULATORIAS_COMPONENTE=[
            ("CRIAR", "PG4_CAT_NOVA", "PG4_COMP_NOVO", "PROIBIDA", None,
             None, "Novo componente", None, None, None),
            ("DESATIVAR", "PG4_CAT_REGRA", "PG4_COMP_REGRA", None, None,
             None, None, None, "2025-01-01", "2025-12-31"),
        ],
    )


def test_confirmacao_v11_completa_locks_historico_e_idempotencia(
    pg_client, postgres_app, monkeypatch,
):
    _, engine, _ = postgres_app
    ids = semear_estado(engine)
    preparada = preparar(pg_client, planilha_completa())

    comandos = []
    def registrar_sql(conn, cursor, statement, parameters, context, executemany):
        comandos.append(statement)
    event.listen(engine, "before_cursor_execute", registrar_sql)
    commit_original = Session.commit
    commits = 0
    def contar_commit(self):
        nonlocal commits
        commits += 1
        return commit_original(self)
    monkeypatch.setattr(Session, "commit", contar_commit)
    try:
        primeira = confirmar(pg_client, preparada)
    finally:
        event.remove(engine, "before_cursor_execute", registrar_sql)
    assert primeira.status_code == 200, primeira.text
    assert commits == 1

    locks = [
        comando.split()[2]
        for comando in comandos
        if comando.strip().upper().startswith("LOCK TABLE")
    ]
    assert locks == LOCKS_V11
    # Lote com 20 operações: teto fixo prova que não há consulta por linha.
    assert len(comandos) <= 80
    assert sum(
        comando.lstrip().upper().startswith("SELECT") for comando in comandos
    ) <= 20
    resultado = primeira.json()
    assert resultado["status"] == "CONFIRMADA"
    assert resultado["totais"] == {
        "criar": 11, "atualizar": 4, "desativar": 5, "sem_alteracao": 0,
    }
    assert resultado["codigos_afetados"] == sorted(resultado["codigos_afetados"])
    assert all(":" in codigo for codigo in resultado["codigos_afetados"])
    assert "token" not in primeira.text.lower()

    with Session(engine) as db:
        sessao = db.scalar(select(SessaoImportacaoCadastral))
        assert sessao.status == "CONFIRMADA"
        assert sessao.resultado == resultado
        assert db.scalar(select(MateriaPrima).where(
            MateriaPrima.codigo == "PG4_MP_UPD"
        )).nome == "MP atualizada"
        assert db.scalar(select(MateriaPrima).where(
            MateriaPrima.codigo == "PG4_MP_DES"
        )).ativa is False
        categoria = db.scalar(select(CategoriaProduto).where(
            CategoriaProduto.codigo == "PG4_CAT_UPD"
        ))
        assert categoria.nome == "Cat antiga"
        assert categoria.descricao == "Descrição atualizada"
        assert db.scalar(select(CategoriaProduto).where(
            CategoriaProduto.codigo == "PG4_CAT_DES"
        )).ativa is False
        componente = db.scalar(select(ComponenteRegulatorio).where(
            ComponenteRegulatorio.codigo == "PG4_COMP_NOVO"
        ))
        assert (componente.unidade, componente.base) == ("%", "MASSA_MASSA")
        assert db.scalar(select(ComponenteRegulatorio).where(
            ComponenteRegulatorio.codigo == "PG4_COMP_DES"
        )).ativo is False
        composicoes = list(db.scalars(select(ComposicaoComponenteMP).where(
            ComposicaoComponenteMP.materia_prima_id == db.scalar(
                select(MateriaPrima.id).where(MateriaPrima.codigo == "PG4_MP_NOVA")
            )
        ).order_by(ComposicaoComponenteMP.data_referencia)))
        assert [item.concentracao for item in composicoes] == [
            Decimal("1.234567"), None, Decimal("0.000000")
        ]
        historica = db.get(ComposicaoComponenteMP, ids["composicao_id"])
        assert historica.ativo is False
        assert (historica.concentracao, historica.fonte, historica.observacao) == (
            Decimal("2.000000"), "Histórica", "Preservar"
        )
        anterior = db.get(RegraRegulatoria, ids["regra_atualizar_id"])
        assert anterior.ativa is False
        assert anterior.maximo == Decimal("10.000000")
        sucessora = db.scalar(select(RegraRegulatoria).where(
            RegraRegulatoria.regra_anterior_id == anterior.id
        ))
        assert sucessora.revisao == 2
        assert sucessora.maximo == Decimal("20.000000")
        assert sucessora.ativa is True
        assert db.get(RegraRegulatoria, ids["regra_desativar_id"]).ativa is False

    invalida_confirmada = pg_client.post(
        f'/api/v1/importacoes-cadastrais/{preparada["sessao_id"]}/confirmar',
        json={"token": "token-invalido"},
    )
    assert invalida_confirmada.status_code == 404
    assert resultado["sessao_id"] not in invalida_confirmada.text
    with engine.connect() as conn:
        xmin_sessao_antes = conn.execute(text(
            "SELECT xmin::text FROM sessoes_importacao_cadastral "
            "WHERE uuid_publico=:id"
        ), {"id": preparada["sessao_id"]}).scalar_one()

    import services.confirmacao_importacao_cadastral as modulo
    def proibido(*args, **kwargs):
        raise AssertionError("repetição não deve bloquear, revalidar ou aplicar")
    monkeypatch.setattr(modulo, "_bloquear_cadastros", proibido)
    monkeypatch.setattr(modulo, "pre_validar_planilha_cadastral_completo", proibido)
    monkeypatch.setattr(modulo, "aplicar_importacao_regulatoria", proibido)
    segunda = confirmar(pg_client, preparada)
    assert segunda.status_code == 200
    assert segunda.json() == resultado
    with engine.connect() as conn:
        xmin_sessao_depois = conn.execute(text(
            "SELECT xmin::text FROM sessoes_importacao_cadastral "
            "WHERE uuid_publico=:id"
        ), {"id": preparada["sessao_id"]}).scalar_one()
    assert xmin_sessao_depois == xmin_sessao_antes


def test_segunda_sessao_obsoleta_falha_sem_adaptacao(
    pg_client, postgres_app,
):
    _, engine, _ = postgres_app
    conteudo = montar_planilha(
        CATEGORIAS_PRODUTO=[
            ("CRIAR", "PG4_CAT_CONCORRENTE", "Categoria", None),
        ],
    )
    primeira = preparar(pg_client, conteudo)
    segunda = preparar(pg_client, conteudo)
    assert confirmar(pg_client, primeira).status_code == 200
    resposta = confirmar(pg_client, segunda)
    assert resposta.status_code == 409
    assert resposta.json()["detail"]["codigo"] == "REVALIDACAO_DIVERGENTE"
    with Session(engine) as db:
        sessoes = list(db.scalars(
            select(SessaoImportacaoCadastral).order_by(
                SessaoImportacaoCadastral.id
            )
        ))
        assert [item.status for item in sessoes] == ["CONFIRMADA", "FALHOU"]
        assert len(list(db.scalars(select(CategoriaProduto).where(
            CategoriaProduto.codigo == "PG4_CAT_CONCORRENTE"
        )))) == 1


def test_falha_apos_flush_reverte_todo_lote_v11(
    pg_client, postgres_app, monkeypatch,
):
    _, engine, _ = postgres_app
    conteudo = montar_planilha(
        MATERIAS_PRIMAS=[
            ("CRIAR", "PG4_MP_ROLLBACK", "MP rollback", "SIM"),
        ],
        CATEGORIAS_PRODUTO=[
            ("CRIAR", "PG4_CAT_ROLLBACK", "Categoria rollback", None),
        ],
    )
    preparada = preparar(pg_client, conteudo)
    import services.confirmacao_importacao_cadastral as modulo
    original = modulo.aplicar_importacao_regulatoria
    def falhar(db, dados, operacoes):
        original(db, dados, operacoes)
        raise RuntimeError("falha técnica controlada após flush")
    monkeypatch.setattr(modulo, "aplicar_importacao_regulatoria", falhar)
    resposta = confirmar(pg_client, preparada)
    assert resposta.status_code == 500
    assert resposta.json() == {
        "detail": "Não foi possível confirmar a importação cadastral."
    }
    with Session(engine) as db:
        assert db.scalar(select(MateriaPrima).where(
            MateriaPrima.codigo == "PG4_MP_ROLLBACK"
        )) is None
        assert db.scalar(select(CategoriaProduto).where(
            CategoriaProduto.codigo == "PG4_CAT_ROLLBACK"
        )) is None
        sessao = db.scalar(select(SessaoImportacaoCadastral))
        assert sessao.status == "FALHOU"
        assert sessao.resultado["tipo"] == "ERRO_TECNICO"


def test_sem_alteracao_regulatoria_nao_executa_update(
    pg_client, postgres_app,
):
    _, engine, _ = postgres_app
    with Session(engine) as db, db.begin():
        db.add_all([
            CategoriaProduto(
                codigo="PG4_CAT_SEM", nome="Categoria igual",
                descricao="Descrição igual",
            ),
            ComponenteRegulatorio(
                codigo="PG4_COMP_SEM", nome="Componente igual",
                descricao="Descrição igual",
            ),
        ])
    conteudo = montar_planilha(
        CATEGORIAS_PRODUTO=[
            ("ATUALIZAR", "PG4_CAT_SEM", "Categoria igual", "Descrição igual"),
        ],
        COMPONENTES_REGULATORIOS=[
            ("ATUALIZAR", "PG4_COMP_SEM", "Componente igual", "Descrição igual"),
        ],
    )
    preparada = preparar(pg_client, conteudo)
    with engine.connect() as conn:
        antes = tuple(conn.execute(text(
            "SELECT codigo, xmin::text FROM categorias_produto "
            "WHERE codigo='PG4_CAT_SEM' UNION ALL "
            "SELECT codigo, xmin::text FROM componentes_regulatorios "
            "WHERE codigo='PG4_COMP_SEM' ORDER BY codigo"
        )).all())
    resposta = confirmar(pg_client, preparada)
    assert resposta.status_code == 200, resposta.text
    with engine.connect() as conn:
        depois = tuple(conn.execute(text(
            "SELECT codigo, xmin::text FROM categorias_produto "
            "WHERE codigo='PG4_CAT_SEM' UNION ALL "
            "SELECT codigo, xmin::text FROM componentes_regulatorios "
            "WHERE codigo='PG4_COMP_SEM' ORDER BY codigo"
        )).all())
    assert depois == antes
    assert resposta.json()["totais"]["sem_alteracao"] == 2
    assert resposta.json()["codigos_afetados"] == []


def _confirmar_concorrente(app, preparada, barreira):
    with TestClient(app) as client:
        barreira.wait(timeout=10)
        resposta = confirmar(client, preparada)
        return resposta.status_code, resposta.json()


def test_duas_confirmacoes_simultaneas_da_mesma_sessao_sao_idempotentes(
    postgres_app,
):
    app, engine, _ = postgres_app
    with TestClient(app) as client:
        preparada = preparar(client, montar_planilha(
            CATEGORIAS_PRODUTO=[
                ("CRIAR", "PG4_CAT_MESMA", "Categoria mesma sessão", None),
            ],
        ))
    barreira = Barrier(2)
    with ThreadPoolExecutor(max_workers=2) as pool:
        futuros = [
            pool.submit(_confirmar_concorrente, app, preparada, barreira)
            for _ in range(2)
        ]
        respostas = [futuro.result(timeout=30) for futuro in futuros]
    assert respostas[0] == respostas[1]
    assert respostas[0][0] == 200
    with Session(engine) as db:
        assert len(list(db.scalars(select(CategoriaProduto).where(
            CategoriaProduto.codigo == "PG4_CAT_MESMA"
        )))) == 1
        sessao = db.scalar(select(SessaoImportacaoCadastral))
        assert sessao.status == "CONFIRMADA"
        assert sessao.resultado == respostas[0][1]


@pytest.mark.parametrize("caso", ["catalogos", "composicao", "regra", "revisao"])
def test_sessoes_distintas_concorrentes_revalidam_sem_adaptar(
    postgres_app, caso,
):
    app, engine, _ = postgres_app
    with Session(engine) as db, db.begin():
        mp = MateriaPrima(codigo="PG4_CONC_MP", nome="MP concorrente", ativa=True)
        cat = CategoriaProduto(codigo="PG4_CONC_CAT", nome="Categoria concorrente")
        comp = ComponenteRegulatorio(codigo="PG4_CONC_COMP", nome="Componente concorrente")
        db.add_all([mp, cat, comp])
        db.flush()
        if caso == "revisao":
            db.add(RegraRegulatoria(
                categoria_id=cat.id, tipo_alvo="MATERIA_PRIMA",
                materia_prima_id=mp.id, componente_id=None,
                tratamento="LIMITADA", minimo=None, maximo=Decimal("10"),
                unidade="%", base="MASSA_MASSA", justificativa="Original",
                referencia_normativa=None, revisao=1, regra_anterior_id=None,
                ativa=True, vigencia_inicio=date(2026, 1, 1),
                vigencia_fim=date(2026, 12, 31),
            ))
    linhas = {
        "catalogos": dict(
            CATEGORIAS_PRODUTO=[("CRIAR", "PG4_CONC_CAT_N", "Nova categoria", None)],
            COMPONENTES_REGULATORIOS=[("CRIAR", "PG4_CONC_COMP_N", "Novo componente", None)],
        ),
        "composicao": dict(COMPOSICAO_COMPONENTES_MP=[
            ("CRIAR", "PG4_CONC_MP", "PG4_CONC_COMP", "2026-02-01",
             "INFORMADO", "1.000000", None, None),
        ]),
        "regra": dict(REGRAS_REGULATORIAS_COMPONENTE=[
            ("CRIAR", "PG4_CONC_CAT", "PG4_CONC_COMP", "PROIBIDA", None,
             None, "Concorrente", None, None, None),
        ]),
        "revisao": dict(REGRAS_REGULATORIAS_MP=[
            ("ATUALIZAR", "PG4_CONC_CAT", "PG4_CONC_MP", "LIMITADA", None,
             "20.000000", "Revisada", None, "2026-01-01", "2026-12-31"),
        ]),
    }[caso]
    conteudo = montar_planilha(**linhas)
    with TestClient(app) as client:
        preparadas = [preparar(client, conteudo), preparar(client, conteudo)]
    barreira = Barrier(2)
    with ThreadPoolExecutor(max_workers=2) as pool:
        futuros = [
            pool.submit(_confirmar_concorrente, app, item, barreira)
            for item in preparadas
        ]
        respostas = [futuro.result(timeout=30) for futuro in futuros]
    assert sorted(status for status, _ in respostas) == [200, 409]
    conflito = next(corpo for status, corpo in respostas if status == 409)
    assert conflito["detail"] == {
        "codigo": "REVALIDACAO_DIVERGENTE",
        "mensagem": "Os cadastros mudaram desde a preparação; prepare uma nova sessão.",
    }
    with Session(engine) as db:
        assert sorted(db.scalars(select(SessaoImportacaoCadastral.status)).all()) == [
            "CONFIRMADA", "FALHOU",
        ]
        if caso == "catalogos":
            assert db.scalar(select(func.count()).select_from(CategoriaProduto).where(
                CategoriaProduto.codigo == "PG4_CONC_CAT_N"
            )) == 1
            assert db.scalar(select(func.count()).select_from(ComponenteRegulatorio).where(
                ComponenteRegulatorio.codigo == "PG4_CONC_COMP_N"
            )) == 1
        elif caso == "composicao":
            assert db.scalar(select(func.count()).select_from(ComposicaoComponenteMP)) == 1
        elif caso == "regra":
            assert db.scalar(select(func.count()).select_from(RegraRegulatoria)) == 1
        else:
            regras = list(db.scalars(select(RegraRegulatoria).order_by(RegraRegulatoria.revisao)))
            assert [regra.revisao for regra in regras] == [1, 2]
            assert [regra.ativa for regra in regras] == [False, True]


def test_oito_locks_bloqueiam_escritores_e_permitem_leitura(
    pg_client, postgres_app, monkeypatch,
):
    _, engine, _ = postgres_app
    preparada = preparar(pg_client, montar_planilha(
        CATEGORIAS_PRODUTO=[("CRIAR", "PG4_CAT_LOCK", "Categoria lock", None)],
    ))
    import services.confirmacao_importacao_cadastral as modulo
    locks_adquiridos = Event()
    liberar = Event()
    original = modulo._bloquear_cadastros

    def pausar_apos_locks(db, versao):
        original(db, versao)
        locks_adquiridos.set()
        assert liberar.wait(timeout=15)

    monkeypatch.setattr(modulo, "_bloquear_cadastros", pausar_apos_locks)
    with ThreadPoolExecutor(max_workers=1) as pool:
        futuro = pool.submit(confirmar, pg_client, preparada)
        assert locks_adquiridos.wait(timeout=10)
        with engine.connect() as conn:
            assert isinstance(conn.execute(text("SELECT count(*) FROM categorias_produto")).scalar_one(), int)
        for tabela in LOCKS_V11:
            with engine.connect() as conn:
                transacao = conn.begin()
                conn.execute(text("SET LOCAL lock_timeout = '150ms'"))
                with pytest.raises(OperationalError):
                    conn.execute(text(
                        f"LOCK TABLE {tabela} IN ROW EXCLUSIVE MODE"
                    ))
                transacao.rollback()
        liberar.set()
        resposta = futuro.result(timeout=30)
    assert resposta.status_code == 200, resposta.text


@pytest.mark.parametrize("etapa", [
    "_aplicar_catalogos_base",
    "_aplicar_catalogos_regulatorios",
    "_aplicar_composicao_nutricional",
    "_aplicar_precos",
    "_aplicar_composicoes_regulatorias",
    "_aplicar_regras",
])
def test_falhas_injetadas_entre_flushes_revertem_lote_inteiro(
    pg_client, postgres_app, monkeypatch, etapa,
):
    _, engine, _ = postgres_app
    ids = semear_estado(engine)
    preparada = preparar(pg_client, planilha_completa())
    import services.aplicacao_importacao_regulatoria as modulo
    original = getattr(modulo, etapa)

    def falhar(*args, **kwargs):
        original(*args, **kwargs)
        raise RuntimeError(f"falha interna controlada em {etapa}")

    monkeypatch.setattr(modulo, etapa, falhar)
    resposta = confirmar(pg_client, preparada)
    assert resposta.status_code == 500
    assert resposta.json() == {
        "detail": "Não foi possível confirmar a importação cadastral."
    }
    assert etapa not in resposta.text
    with Session(engine) as db:
        assert db.scalar(select(MateriaPrima).where(
            MateriaPrima.codigo == "PG4_MP_NOVA"
        )) is None
        assert db.scalar(select(CategoriaProduto).where(
            CategoriaProduto.codigo == "PG4_CAT_NOVA"
        )) is None
        assert db.get(ComposicaoComponenteMP, ids["composicao_id"]).ativo is True
        anterior = db.get(RegraRegulatoria, ids["regra_atualizar_id"])
        assert anterior.ativa is True
        assert db.scalar(select(RegraRegulatoria).where(
            RegraRegulatoria.regra_anterior_id == anterior.id
        )) is None
        sessao = db.scalar(select(SessaoImportacaoCadastral))
        assert sessao.status == "FALHOU"
        assert sessao.resultado["tipo"] == "ERRO_TECNICO"


def test_falha_entre_desativar_predecessora_e_inserir_sucessora_reverte(
    pg_client, postgres_app,
):
    _, engine, _ = postgres_app
    ids = semear_estado(engine)
    preparada = preparar(pg_client, planilha_completa())
    desativacao_observada = False

    def falhar_antes_da_sucessora(conn, cursor, statement, params, context, many):
        nonlocal desativacao_observada
        sql = statement.lower()
        if sql.startswith("update regras_regulatorias"):
            desativacao_observada = True
        if desativacao_observada and sql.startswith("insert into regras_regulatorias"):
            raise RuntimeError("falha controlada antes da sucessora")

    event.listen(engine, "before_cursor_execute", falhar_antes_da_sucessora)
    try:
        resposta = confirmar(pg_client, preparada)
    finally:
        event.remove(engine, "before_cursor_execute", falhar_antes_da_sucessora)
    assert resposta.status_code == 500
    assert desativacao_observada
    with Session(engine) as db:
        anterior = db.get(RegraRegulatoria, ids["regra_atualizar_id"])
        assert anterior.ativa is True
        assert db.scalar(select(RegraRegulatoria).where(
            RegraRegulatoria.regra_anterior_id == anterior.id
        )) is None
        assert db.scalar(select(SessaoImportacaoCadastral)).status == "FALHOU"


def test_commit_com_conexao_invalidada_permanece_com_resultado_desconhecido(
    pg_client, postgres_app, monkeypatch,
):
    _, engine, _ = postgres_app
    preparada = preparar(pg_client, montar_planilha(
        CATEGORIAS_PRODUTO=[
            ("CRIAR", "PG4_CAT_AMBIGUA", "Categoria ambígua", None),
        ],
    ))

    def commit_ambiguo(self):
        self.connection().invalidate()
        raise RuntimeError("comunicação perdida durante commit")

    monkeypatch.setattr(Session, "commit", commit_ambiguo)
    resposta = confirmar(pg_client, preparada)
    assert resposta.status_code == 500
    assert resposta.json() == {
        "detail": "Não foi possível confirmar a importação cadastral."
    }
    assert "comunicação" not in resposta.text
    with Session(engine) as db:
        sessao = db.scalar(select(SessaoImportacaoCadastral))
        assert sessao.status == "PENDENTE"
        assert sessao.resultado is None
        assert db.scalar(select(CategoriaProduto).where(
            CategoriaProduto.codigo == "PG4_CAT_AMBIGUA"
        )) is None


def test_regra_componente_preserva_cadeia_da_revisao_dois_para_tres(
    pg_client, postgres_app,
):
    _, engine, _ = postgres_app
    with Session(engine) as db, db.begin():
        categoria = CategoriaProduto(codigo="PG4_CAT_R3", nome="Categoria R3")
        componente = ComponenteRegulatorio(codigo="PG4_COMP_R3", nome="Componente R3")
        db.add_all([categoria, componente])
        db.flush()
        revisao_1 = RegraRegulatoria(
            categoria_id=categoria.id, tipo_alvo="COMPONENTE",
            materia_prima_id=None, componente_id=componente.id,
            tratamento="LIMITADA", minimo=Decimal("1"), maximo=Decimal("9"),
            unidade="%", base="MASSA_MASSA", justificativa="Revisão um",
            referencia_normativa="R1", revisao=1, regra_anterior_id=None,
            ativa=False, vigencia_inicio=None, vigencia_fim=None,
        )
        db.add(revisao_1)
        db.flush()
        revisao_2 = RegraRegulatoria(
            categoria_id=categoria.id, tipo_alvo="COMPONENTE",
            materia_prima_id=None, componente_id=componente.id,
            tratamento="LIMITADA", minimo=Decimal("2"), maximo=Decimal("8"),
            unidade="%", base="MASSA_MASSA", justificativa="Revisão dois",
            referencia_normativa="R2", revisao=2,
            regra_anterior_id=revisao_1.id, ativa=True,
            vigencia_inicio=None, vigencia_fim=None,
        )
        db.add(revisao_2)
        db.flush()
        ids = revisao_1.id, revisao_2.id
    preparada = preparar(pg_client, montar_planilha(
        REGRAS_REGULATORIAS_COMPONENTE=[
            ("ATUALIZAR", "PG4_CAT_R3", "PG4_COMP_R3", "LIMITADA",
             "3.000000", "7.000000", "Revisão três", "R3", None, None),
        ],
    ))
    resposta = confirmar(pg_client, preparada)
    assert resposta.status_code == 200, resposta.text
    with Session(engine) as db:
        regras = list(db.scalars(select(RegraRegulatoria).order_by(
            RegraRegulatoria.revisao
        )))
        assert [regra.revisao for regra in regras] == [1, 2, 3]
        assert [regra.ativa for regra in regras] == [False, False, True]
        assert regras[2].regra_anterior_id == ids[1]
        assert regras[1].regra_anterior_id == ids[0]
        assert (regras[0].minimo, regras[0].maximo, regras[0].justificativa) == (
            Decimal("1.000000"), Decimal("9.000000"), "Revisão um"
        )
        assert (regras[1].minimo, regras[1].maximo, regras[1].justificativa) == (
            Decimal("2.000000"), Decimal("8.000000"), "Revisão dois"
        )
        assert (regras[2].minimo, regras[2].maximo, regras[2].justificativa) == (
            Decimal("3.000000"), Decimal("7.000000"), "Revisão três"
        )
        assert all(regra.categoria_id == regras[0].categoria_id for regra in regras)
        assert all(regra.componente_id == regras[0].componente_id for regra in regras)
        assert all(regra.vigencia_inicio is None and regra.vigencia_fim is None for regra in regras)
