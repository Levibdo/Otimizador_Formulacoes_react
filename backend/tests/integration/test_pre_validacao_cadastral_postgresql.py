import io
from datetime import date

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from models import ComposicaoMateriaPrima, MateriaPrima, Nutriente, PrecoMateriaPrima
from services.planilha_cadastral import gerar_template_cadastral


def planilha(mps=(), nutrientes=(), composicoes=(), precos=()):
    workbook = load_workbook(io.BytesIO(gerar_template_cadastral()))
    for aba in ("MATERIAS_PRIMAS", "NUTRIENTES", "COMPOSICAO_NUTRICIONAL", "PRECOS_MP"):
        ws = workbook[aba]
        ws.delete_rows(2, max(ws.max_row - 1, 0))
    for linha in mps:
        workbook["MATERIAS_PRIMAS"].append(linha)
    for linha in nutrientes:
        workbook["NUTRIENTES"].append(linha)
    for linha in composicoes:
        workbook["COMPOSICAO_NUTRICIONAL"].append(linha)
    for linha in precos:
        workbook["PRECOS_MP"].append(linha)
    saida = io.BytesIO()
    workbook.save(saida)
    return saida.getvalue()


def snapshot_semantico(engine):
    with Session(engine) as db:
        mps = tuple(db.execute(select(
            MateriaPrima.id, MateriaPrima.codigo, MateriaPrima.nome, MateriaPrima.ativa,
        ).order_by(MateriaPrima.id)).all())
        nutrientes = tuple(db.execute(select(
            Nutriente.id, Nutriente.codigo, Nutriente.nome, Nutriente.unidade,
        ).order_by(Nutriente.id)).all())
        composicoes = tuple(db.execute(select(
            ComposicaoMateriaPrima.id, ComposicaoMateriaPrima.materia_prima_id,
            ComposicaoMateriaPrima.nutriente_id, ComposicaoMateriaPrima.valor,
        ).order_by(ComposicaoMateriaPrima.id)).all())
        precos = tuple(db.execute(select(
            PrecoMateriaPrima.id, PrecoMateriaPrima.materia_prima_id,
            PrecoMateriaPrima.preco_kg, PrecoMateriaPrima.vigencia_inicio,
            PrecoMateriaPrima.vigencia_fim,
        ).order_by(PrecoMateriaPrima.id)).all())
        return mps, nutrientes, composicoes, precos


def test_pre_validacao_postgresql_resolve_banco_e_nao_escreve(pg_client, postgres_app):
    _, engine, _ = postgres_app
    with Session(engine) as db:
        mp = MateriaPrima(codigo="PG_MP", nome="MP PostgreSQL", ativa=True)
        nutriente = Nutriente(codigo="PG_NUT", nome="Nutriente PostgreSQL", unidade="g/100 g")
        db.add_all([mp, nutriente])
        db.flush()
        db.add_all([
            ComposicaoMateriaPrima(materia_prima_id=mp.id, nutriente_id=nutriente.id, valor=10),
            PrecoMateriaPrima(materia_prima_id=mp.id, preco_kg=5,
                              vigencia_inicio=date(2026, 1, 1), vigencia_fim=date(2026, 1, 31)),
        ])
        db.commit()
    antes = snapshot_semantico(engine)
    conteudo = planilha(
        mps=[("ATUALIZAR", "PG_MP", "MP PostgreSQL", "SIM")],
        nutrientes=[("ATUALIZAR", "PG_NUT", "Nutriente PostgreSQL", "g/100 g")],
        composicoes=[("ATUALIZAR", "PG_MP", "PG_NUT", 11)],
        precos=[("CRIAR", "PG_MP", 6, "2026-02-01", None)],
    )
    resposta = pg_client.post(
        "/api/v1/importacoes-cadastrais/validar",
        files={"arquivo": ("cadastros.xlsx", conteudo,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resposta.status_code == 200
    resultado = resposta.json()
    assert resultado["valido_para_confirmacao"] is True
    assert resultado["resumo"]["MATERIAS_PRIMAS"]["sem_alteracao"] == 1
    assert resultado["resumo"]["COMPOSICAO_NUTRICIONAL"]["atualizar"] == 1
    assert resultado["resumo"]["PRECOS_MP"]["criar"] == 1
    assert snapshot_semantico(engine) == antes


def test_pre_validacao_postgresql_criacoes_internas_e_rollback_total(pg_client, postgres_app):
    _, engine, _ = postgres_app
    antes = snapshot_semantico(engine)
    conteudo = planilha(
        mps=[("CRIAR", "PG_NOVA", "MP nova fictícia", "SIM")],
        nutrientes=[("CRIAR", "PG_NUT_NOVO", "Nutriente novo fictício", "g/100 g")],
        composicoes=[("CRIAR", "PG_NOVA", "PG_NUT_NOVO", 20)],
        precos=[("CRIAR", "PG_NOVA", 8, "2026-01-01", None)],
    )
    primeira = pg_client.post(
        "/api/v1/importacoes-cadastrais/validar",
        files={"arquivo": ("cadastros.xlsx", conteudo)},
    ).json()
    segunda = pg_client.post(
        "/api/v1/importacoes-cadastrais/validar",
        files={"arquivo": ("cadastros.xlsx", conteudo)},
    ).json()
    assert primeira == segunda
    assert primeira["valido_para_confirmacao"] is True
    assert snapshot_semantico(engine) == antes
