import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook
import pytest
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from models import MateriaPrima, SessaoImportacaoCadastral
from models.regulatorio import (
    CategoriaProduto, ComponenteRegulatorio, ComposicaoComponenteMP, RegraRegulatoria,
)
from services.planilha_cadastral import gerar_template_cadastral
from services.pre_validacao_regulatoria import _sobrepoe


pytestmark = pytest.mark.postgresql

ABAS = (
    "MATERIAS_PRIMAS", "NUTRIENTES", "COMPOSICAO_NUTRICIONAL", "PRECOS_MP",
    "CATEGORIAS_PRODUTO", "COMPONENTES_REGULATORIOS", "COMPOSICAO_COMPONENTES_MP",
    "REGRAS_REGULATORIAS_MP", "REGRAS_REGULATORIAS_COMPONENTE",
)


def planilha(**linhas):
    wb = load_workbook(io.BytesIO(gerar_template_cadastral("1.1")))
    for aba in ABAS:
        ws = wb[aba]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        for linha in linhas.get(aba, ()):
            ws.append(linha)
    saida = io.BytesIO(); wb.save(saida); return saida.getvalue()


def validar(client, conteudo):
    resposta = client.post(
        "/api/v1/importacoes-cadastrais/validar",
        files={"arquivo": ("regulatorio.xlsx", conteudo)},
    )
    assert resposta.status_code == 200, resposta.text
    return resposta.json()


def snapshot(engine):
    tabelas = (
        "materias_primas", "categorias_produto", "componentes_regulatorios",
        "composicoes_componentes_mp", "regras_regulatorias", "sessoes_importacao_cadastral",
    )
    with engine.connect() as conn:
        linhas = {
            tabela: tuple(conn.execute(text(
                f'SELECT id, xmin::text FROM "{tabela}" ORDER BY id'
            )).all())
            for tabela in tabelas
        }
        sequencias = tuple(conn.execute(text(
            "SELECT sequencename, last_value FROM pg_sequences "
            "WHERE schemaname = current_schema() ORDER BY sequencename"
        )).all())
        return linhas, sequencias


def test_validar_v11_postgresql_estado_projetado_sem_escrita_e_deterministico(pg_client, postgres_app):
    _, engine, _ = postgres_app
    with Session(engine) as db:
        mp = MateriaPrima(codigo="PG_MP_REG", nome="MP regulatória", ativa=True)
        cat = CategoriaProduto(codigo="PG_CAT_REG", nome="Categoria regulatória", descricao=None)
        comp = ComponenteRegulatorio(codigo="PG_COMP_REG", nome="Componente regulatório", descricao=None)
        db.add_all([mp, cat, comp]); db.flush()
        composicao = ComposicaoComponenteMP(
            materia_prima_id=mp.id, componente_id=comp.id, data_referencia=date(2026, 1, 1),
            situacao="INFORMADO", concentracao=Decimal("1.234567"), fonte=None, observacao=None,
        )
        regra = RegraRegulatoria(
            categoria_id=cat.id, tipo_alvo="MATERIA_PRIMA", materia_prima_id=mp.id,
            componente_id=None, tratamento="LIMITADA", minimo=None, maximo=Decimal("10"),
            unidade="%", base="MASSA_MASSA", justificativa="Regra existente",
            referencia_normativa=None, vigencia_inicio=date(2026, 1, 1),
            vigencia_fim=date(2026, 6, 30), revisao=1, ativa=True,
        )
        db.add_all([composicao, regra]); db.commit()
    antes = snapshot(engine)
    conteudo = planilha(
        CATEGORIAS_PRODUTO=[("ATUALIZAR", "PG_CAT_REG", None, None)],
        COMPONENTES_REGULATORIOS=[("ATUALIZAR", "PG_COMP_REG", None, None)],
        COMPOSICAO_COMPONENTES_MP=[
            ("CRIAR", "PG_MP_REG", "PG_COMP_REG", "2026-01-01", "INFORMADO", "1.234567", None, None),
        ],
        REGRAS_REGULATORIAS_MP=[
            ("DESATIVAR", "PG_CAT_REG", "PG_MP_REG", None, None, None, None, None, "2026-01-01", "2026-06-30"),
            ("CRIAR", "PG_CAT_REG", "PG_MP_REG", "LIMITADA", None, "20", "Nova vigência", None, "2026-07-01", None),
        ],
    )
    primeira = validar(pg_client, conteudo)
    segunda = validar(pg_client, conteudo)
    assert primeira == segunda
    assert primeira["valido_para_confirmacao"], primeira["diagnosticos"]
    assert primeira["resumo"]["COMPOSICAO_COMPONENTES_MP"]["sem_alteracao"] == 1
    assert primeira["resumo"]["REGRAS_REGULATORIAS_MP"]["desativar"] == 1
    assert primeira["resumo"]["REGRAS_REGULATORIAS_MP"]["criar"] == 1
    assert snapshot(engine) == antes


def test_validar_v11_banco_vazio_e_preparar_bloqueado(pg_client, postgres_app):
    _, engine, _ = postgres_app
    conteudo = planilha(
        MATERIAS_PRIMAS=[("CRIAR", "PG_MP_NOVA", "MP nova", "SIM")],
        CATEGORIAS_PRODUTO=[("CRIAR", "PG_CAT_NOVA", "Categoria nova", None)],
        COMPONENTES_REGULATORIOS=[("CRIAR", "PG_COMP_NOVO", "Componente novo", None)],
        COMPOSICAO_COMPONENTES_MP=[
            ("CRIAR", "PG_MP_NOVA", "PG_COMP_NOVO", "2026-01-01", "DESCONHECIDO", None, None, None),
        ],
        REGRAS_REGULATORIAS_COMPONENTE=[
            ("CRIAR", "PG_CAT_NOVA", "PG_COMP_NOVO", "PROIBIDA", None, None, "Teste", None, None, None),
        ],
    )
    validacao = validar(pg_client, conteudo)
    assert validacao["valido_para_confirmacao"], validacao["diagnosticos"]
    resposta = pg_client.post(
        "/api/v1/importacoes-cadastrais/preparar",
        files={"arquivo": ("regulatorio.xlsx", conteudo)},
    )
    assert resposta.status_code == 422
    detalhe = resposta.json()["detail"]
    assert any("preparação do contrato 1.1" in item["mensagem"] for item in detalhe["validacao"]["diagnosticos"])
    with Session(engine) as db:
        assert db.scalar(select(func.count()).select_from(SessaoImportacaoCadastral)) == 0
        assert db.scalar(select(func.count()).select_from(MateriaPrima)) == 0


@pytest.mark.parametrize("inicio_a,fim_a,inicio_b,fim_b", [
    (date(2026, 1, 1), date(2026, 1, 31), date(2026, 2, 1), date(2026, 2, 28)),
    (date(2026, 1, 1), date(2026, 1, 31), date(2026, 1, 31), date(2026, 2, 28)),
    (None, None, date(2026, 1, 1), date(2026, 1, 31)),
    (None, date(2026, 1, 1), date(2026, 1, 2), None),
    (None, date(2026, 1, 1), date(2026, 1, 1), None),
    (date(2026, 1, 1), None, date(2027, 1, 1), None),
    (date.min, date.max, date(2026, 1, 1), date(2026, 1, 1)),
    (date(2026, 1, 1), date(2026, 12, 31), date(2026, 3, 1), date(2026, 4, 1)),
])
def test_sobreposicao_python_equivale_ao_daterange_postgresql(
    postgres_app, inicio_a, fim_a, inicio_b, fim_b
):
    _, engine, _ = postgres_app
    with engine.connect() as conn:
        postgres = conn.scalar(text(
            "SELECT daterange(CAST(:ia AS date), CAST(:fa AS date), '[]') "
            "&& daterange(CAST(:ib AS date), CAST(:fb AS date), '[]')"
        ), {"ia": inicio_a, "fa": fim_a, "ib": inicio_b, "fb": fim_b})
    assert _sobrepoe(inicio_a, fim_a, inicio_b, fim_b) is postgres
