import io
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
import re
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from hashlib import sha256
from uuid import uuid4

from openpyxl import load_workbook
from fastapi.testclient import TestClient
import pytest
from sqlalchemy import inspect, select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from models import (
    ComposicaoMateriaPrima,
    MateriaPrima,
    Nutriente,
    PrecoMateriaPrima,
    SessaoImportacaoCadastral,
)
from services.planilha_cadastral import gerar_template_cadastral, parsear_planilha_cadastral
from services.staging_importacao_cadastral import preparar_sessao
from services.confirmacao_importacao_cadastral import _xlsx_do_payload
from schemas.importacao_cadastral import ConfirmacaoImportacaoRequest


def planilha_valida():
    workbook = load_workbook(io.BytesIO(gerar_template_cadastral()))
    for aba in ("MATERIAS_PRIMAS", "NUTRIENTES", "COMPOSICAO_NUTRICIONAL", "PRECOS_MP"):
        ws = workbook[aba]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
    workbook["MATERIAS_PRIMAS"].append(("CRIAR", "PG_STAGE_MP", "MP staging PG", "SIM"))
    workbook["NUTRIENTES"].append(("CRIAR", "PG_STAGE_NUT", "Nutriente staging PG", "g/100 g"))
    workbook["COMPOSICAO_NUTRICIONAL"].append(("CRIAR", "PG_STAGE_MP", "PG_STAGE_NUT", "20.000000"))
    workbook["PRECOS_MP"].append(("CRIAR", "PG_STAGE_MP", "8.000000", "2026-01-01", None))
    saida = io.BytesIO()
    workbook.save(saida)
    return saida.getvalue()


def snapshot_cadastros(engine):
    with Session(engine) as db:
        return (
            tuple(db.execute(select(MateriaPrima.id, MateriaPrima.codigo, MateriaPrima.nome,
                                    MateriaPrima.ativa).order_by(MateriaPrima.id)).all()),
            tuple(db.execute(select(Nutriente.id, Nutriente.codigo, Nutriente.nome,
                                    Nutriente.unidade).order_by(Nutriente.id)).all()),
            tuple(db.execute(select(ComposicaoMateriaPrima.id,
                                    ComposicaoMateriaPrima.materia_prima_id,
                                    ComposicaoMateriaPrima.nutriente_id,
                                    ComposicaoMateriaPrima.valor)
                             .order_by(ComposicaoMateriaPrima.id)).all()),
            tuple(db.execute(select(PrecoMateriaPrima.id, PrecoMateriaPrima.materia_prima_id,
                                    PrecoMateriaPrima.preco_kg,
                                    PrecoMateriaPrima.vigencia_inicio,
                                    PrecoMateriaPrima.vigencia_fim)
                             .order_by(PrecoMateriaPrima.id)).all()),
        )


def test_preparacao_postgresql_persiste_somente_staging(pg_client, postgres_app):
    _, engine, _ = postgres_app
    antes = snapshot_cadastros(engine)
    conteudo = planilha_valida()
    resposta = pg_client.post(
        "/api/v1/importacoes-cadastrais/preparar",
        files={"arquivo": ("cadastros.xlsx", conteudo)},
    )
    assert resposta.status_code == 201, resposta.text
    corpo = resposta.json()
    assert corpo["arquivo_sha256"] == sha256(conteudo).hexdigest()
    assert corpo["status"] == "PENDENTE"
    assert corpo["operacoes_total"] == 4
    assert snapshot_cadastros(engine) == antes

    with Session(engine) as db:
        sessao = db.scalar(select(SessaoImportacaoCadastral))
        assert sessao is not None
        assert sessao.token_hash != corpo["token_confirmacao"]
        assert corpo["token_confirmacao"] not in str(sessao.payload_normalizado)
        assert sessao.payload_normalizado["dados"]["PRECOS_MP"][0]["preco_kg"] == "8.000000"
        assert sessao.confirmado_em is None
        assert sessao.resultado is None

    consulta = pg_client.get(f'/api/v1/importacoes-cadastrais/{corpo["sessao_id"]}')
    assert consulta.status_code == 200
    assert "token" not in consulta.text.lower()
    assert "payload_normalizado" not in consulta.text


def test_planilha_invalida_nao_cria_staging_postgresql(pg_client, postgres_app):
    _, engine, _ = postgres_app
    resposta = pg_client.post(
        "/api/v1/importacoes-cadastrais/preparar",
        files={"arquivo": ("cadastros.xlsx", "inválido".encode())},
    )
    assert resposta.status_code == 422
    with Session(engine) as db:
        assert db.scalar(select(SessaoImportacaoCadastral)) is None


def test_todos_os_campos_permanentes_sao_imutaveis_postgresql(
    pg_client, postgres_app,
):
    _, engine, _ = postgres_app
    criada = pg_client.post(
        "/api/v1/importacoes-cadastrais/preparar",
        files={"arquivo": ("cadastros.xlsx", planilha_valida())},
    ).json()
    sessao_id = criada["sessao_id"]
    mutacoes = (
        "id = id + 1000",
        "uuid_publico = gen_random_uuid()",
        "token_hash = repeat('b', 64)",
        "arquivo_sha256 = repeat('0', 64)",
        "versao_contrato = '9.9'",
        "payload_normalizado = '{}'::jsonb",
        "resumo = '{}'::jsonb",
        "avisos = '[{}]'::jsonb",
        "criado_em = criado_em - interval '1 second'",
        "expira_em = expira_em + interval '1 second'",
    )
    for atribuicao in mutacoes:
        with pytest.raises(IntegrityError, match="imutável"):
            with engine.begin() as conn:
                conn.execute(text(
                    f"UPDATE sessoes_importacao_cadastral SET {atribuicao} "
                    "WHERE uuid_publico = :id"
                ), {"id": sessao_id})
    with pytest.raises(IntegrityError, match="removida"):
        with engine.begin() as conn:
            conn.execute(text(
                "DELETE FROM sessoes_importacao_cadastral WHERE uuid_publico = :id"
            ), {"id": sessao_id})


def _preparar(pg_client):
    resposta = pg_client.post(
        "/api/v1/importacoes-cadastrais/preparar",
        files={"arquivo": ("cadastros.xlsx", planilha_valida())},
    )
    assert resposta.status_code == 201, resposta.text
    return resposta.json()


def test_transicoes_permitidas_exigem_metadados_coerentes(pg_client, postgres_app):
    _, engine, _ = postgres_app
    confirmada, expirada, falhou = (_preparar(pg_client) for _ in range(3))
    with engine.begin() as conn:
        conn.execute(text("""
            UPDATE sessoes_importacao_cadastral
            SET status = 'CONFIRMADA', confirmado_em = clock_timestamp(),
                resultado = '{"aplicadas": 4}'::jsonb
            WHERE uuid_publico = :id
        """), {"id": confirmada["sessao_id"]})
        conn.execute(text("""
            UPDATE sessoes_importacao_cadastral SET status = 'EXPIRADA'
            WHERE uuid_publico = :id
        """), {"id": expirada["sessao_id"]})
        conn.execute(text("""
            UPDATE sessoes_importacao_cadastral
            SET status = 'FALHOU', resultado = '{"diagnostico": "Falha segura."}'::jsonb
            WHERE uuid_publico = :id
        """), {"id": falhou["sessao_id"]})
    assert pg_client.get(f'/api/v1/importacoes-cadastrais/{confirmada["sessao_id"]}').json()["status"] == "CONFIRMADA"
    assert pg_client.get(f'/api/v1/importacoes-cadastrais/{expirada["sessao_id"]}').json()["status"] == "EXPIRADA"
    assert pg_client.get(f'/api/v1/importacoes-cadastrais/{falhou["sessao_id"]}').json()["status"] == "FALHOU"


def test_transicoes_incompletas_e_estados_finais_sao_rejeitados(pg_client, postgres_app):
    _, engine, _ = postgres_app
    pendente_confirmacao = _preparar(pg_client)
    pendente_falha = _preparar(pg_client)
    pendente_campos = _preparar(pg_client)
    with pytest.raises(IntegrityError):
        with engine.begin() as conn:
            conn.execute(text("UPDATE sessoes_importacao_cadastral SET status = 'CONFIRMADA' WHERE uuid_publico = :id"), {"id": pendente_confirmacao["sessao_id"]})
    with pytest.raises(IntegrityError):
        with engine.begin() as conn:
            conn.execute(text("UPDATE sessoes_importacao_cadastral SET status = 'FALHOU' WHERE uuid_publico = :id"), {"id": pendente_falha["sessao_id"]})
    for atribuicao in ("confirmado_em=clock_timestamp()", "resultado='{}'::jsonb"):
        with pytest.raises(IntegrityError):
            with engine.begin() as conn:
                conn.execute(text(f"UPDATE sessoes_importacao_cadastral SET {atribuicao} WHERE uuid_publico=:id"), {"id": pendente_campos["sessao_id"]})

    finais = {}
    for estado in ("CONFIRMADA", "EXPIRADA", "FALHOU"):
        criada = _preparar(pg_client)
        finais[estado] = criada["sessao_id"]
        with engine.begin() as conn:
            if estado == "CONFIRMADA":
                conn.execute(text("UPDATE sessoes_importacao_cadastral SET status='CONFIRMADA', confirmado_em=clock_timestamp(), resultado='{}'::jsonb WHERE uuid_publico=:id"), {"id": criada["sessao_id"]})
            elif estado == "FALHOU":
                conn.execute(text("UPDATE sessoes_importacao_cadastral SET status='FALHOU', resultado='{}'::jsonb WHERE uuid_publico=:id"), {"id": criada["sessao_id"]})
            else:
                conn.execute(text("UPDATE sessoes_importacao_cadastral SET status='EXPIRADA' WHERE uuid_publico=:id"), {"id": criada["sessao_id"]})

    tentativas = (
        (finais["CONFIRMADA"], "status='EXPIRADA'"),
        (finais["CONFIRMADA"], "status='FALHOU'"),
        (finais["CONFIRMADA"], "resultado=jsonb_build_object('alterado', true)"),
        (finais["EXPIRADA"], "status='PENDENTE'"),
        (finais["EXPIRADA"], "status='CONFIRMADA', confirmado_em=clock_timestamp(), resultado='{}'::jsonb"),
        (finais["FALHOU"], "status='PENDENTE'"),
        (finais["FALHOU"], "status='CONFIRMADA', confirmado_em=clock_timestamp()"),
        (finais["FALHOU"], "resultado=jsonb_build_object('alterado', true)"),
    )
    for identificador, atribuicao in tentativas:
        with pytest.raises(IntegrityError):
            with engine.begin() as conn:
                conn.execute(text(f"UPDATE sessoes_importacao_cadastral SET {atribuicao} WHERE uuid_publico=:id"), {"id": identificador})


def test_token_original_ausente_de_todas_as_colunas_e_digests_validos(pg_client, postgres_app):
    _, engine, _ = postgres_app
    criada = _preparar(pg_client)
    with engine.connect() as conn:
        contem_token = conn.scalar(text("""
            SELECT row_to_json(s)::text LIKE :padrao
            FROM sessoes_importacao_cadastral s WHERE uuid_publico = :id
        """), {"id": criada["sessao_id"], "padrao": f'%{criada["token_confirmacao"]}%'})
        hashes = conn.execute(text("SELECT token_hash, arquivo_sha256 FROM sessoes_importacao_cadastral WHERE uuid_publico=:id"), {"id": criada["sessao_id"]}).one()
    assert contem_token is False
    assert re.fullmatch(r"[0-9a-f]{64}", hashes.token_hash)
    assert re.fullmatch(r"[0-9a-f]{64}", hashes.arquivo_sha256)


def test_expiracao_persistida_nao_altera_cadastros_postgresql(pg_client, postgres_app):
    _, engine, _ = postgres_app
    antes = snapshot_cadastros(engine)
    with Session(engine) as db:
        sessao, _, _ = preparar_sessao(
            planilha_valida(), "dados.xlsx", db,
            agora=datetime.now(timezone.utc) - timedelta(hours=25),
        )
        sessao_id = sessao.uuid_publico
        registro_id = sessao.id
        db.commit()
    with Session(engine) as db:
        original = db.get(SessaoImportacaoCadastral, registro_id)
        staging_antes = (original.payload_normalizado, original.resumo, original.avisos)
    resposta = pg_client.get(f"/api/v1/importacoes-cadastrais/{sessao_id}")
    repetida = pg_client.get(f"/api/v1/importacoes-cadastrais/{sessao_id}")
    assert resposta.status_code == repetida.status_code == 200
    assert resposta.json() == repetida.json()
    assert resposta.json()["status"] == "EXPIRADA"
    with Session(engine) as db:
        atual = db.scalar(select(SessaoImportacaoCadastral).where(SessaoImportacaoCadastral.uuid_publico == sessao_id))
        assert (atual.payload_normalizado, atual.resumo, atual.avisos) == staging_antes
    assert snapshot_cadastros(engine) == antes


def test_uuid_inexistente_retorna_mensagem_segura_postgresql(pg_client):
    resposta = pg_client.get(f"/api/v1/importacoes-cadastrais/{uuid4()}")
    assert resposta.status_code == 404
    assert resposta.json() == {"detail": "Sessão de importação não encontrada."}


def test_migracao_08_downgrade_e_upgrade(postgres_app, alembic_runner):
    _, engine, _ = postgres_app
    alembic_runner(engine.url, "downgrade", "20260919_07")
    assert "sessoes_importacao_cadastral" not in inspect(engine).get_table_names()
    alembic_runner(engine.url, "upgrade", "head")
    assert "sessoes_importacao_cadastral" in inspect(engine).get_table_names()


def test_confirmacao_aplica_lote_completo_e_e_idempotente(pg_client, postgres_app, monkeypatch):
    _, engine, _ = postgres_app
    preparada = _preparar(pg_client)
    caminho = f'/api/v1/importacoes-cadastrais/{preparada["sessao_id"]}/confirmar'
    primeira = pg_client.post(caminho, json={"token": preparada["token_confirmacao"]})

    import services.confirmacao_importacao_cadastral as modulo
    def caminho_proibido(*args, **kwargs):
        raise AssertionError("confirmação idempotente não deve bloquear ou revalidar cadastros")
    monkeypatch.setattr(modulo, "_bloquear_cadastros", caminho_proibido)
    monkeypatch.setattr(modulo, "pre_validar_planilha_cadastral_completo", caminho_proibido)
    segunda = pg_client.post(caminho, json={"token": preparada["token_confirmacao"]})
    assert primeira.status_code == segunda.status_code == 200
    assert primeira.json() == segunda.json()
    assert primeira.json()["status"] == "CONFIRMADA"
    assert "token" not in primeira.text.lower()
    with Session(engine) as db:
        mp = db.scalar(select(MateriaPrima).where(MateriaPrima.codigo == "PG_STAGE_MP"))
        nutriente = db.scalar(select(Nutriente).where(Nutriente.codigo == "PG_STAGE_NUT"))
        assert mp and nutriente
        assert db.scalar(select(ComposicaoMateriaPrima).where(
            ComposicaoMateriaPrima.materia_prima_id == mp.id,
            ComposicaoMateriaPrima.nutriente_id == nutriente.id)).valor == 20
        assert len(db.scalars(select(PrecoMateriaPrima).where(
            PrecoMateriaPrima.materia_prima_id == mp.id)).all()) == 1
        sessao = db.scalar(select(SessaoImportacaoCadastral))
        assert sessao.resultado == primeira.json()


def test_payload_canonico_faz_round_trip_sem_perda_semantica():
    dados = {
        "MATERIAS_PRIMAS": [
            {"linha": 2, "acao": "CRIAR", "codigo": "MP_001", "nome": "Proteína çã", "ativa": True},
            {"linha": 4, "acao": "ATUALIZAR", "codigo": "MP_002", "nome": None, "ativa": False},
            {"linha": 5, "acao": "DESATIVAR", "codigo": "MP_003", "nome": None, "ativa": None},
        ],
        "NUTRIENTES": [
            {"linha": 2, "acao": "CRIAR", "codigo": "NUT_001", "nome": "Ácido fólico", "unidade": "µg/100 g"},
        ],
        "COMPOSICAO_NUTRICIONAL": [
            {"linha": 2, "acao": "CRIAR", "materia_prima_codigo": "MP_001", "nutriente_codigo": "NUT_001", "valor": "999999999999.999999"},
            {"linha": 3, "acao": "ATUALIZAR", "materia_prima_codigo": "MP_002", "nutriente_codigo": "NUT_001", "valor": "0.000000"},
        ],
        "PRECOS_MP": [
            {"linha": 2, "acao": "CRIAR", "materia_prima_codigo": "MP_001", "preco_kg": "0.000000", "vigencia_inicio": "2026-09-25", "vigencia_fim": None},
        ],
    }
    reconstruido = _xlsx_do_payload({"dados": dados})
    resultado = parsear_planilha_cadastral(reconstruido, "payload-canonico.xlsx")
    assert resultado["dados"] == dados


def test_token_do_schema_tem_representacao_protegida():
    segredo = "A" * 43
    dados = ConfirmacaoImportacaoRequest(token=segredo)
    assert segredo not in repr(dados)
    assert dados.token.get_secret_value() == segredo


def test_token_incorreto_e_uuid_inexistente_tem_resposta_generica(pg_client, postgres_app):
    _, engine, _ = postgres_app
    preparada = _preparar(pg_client)
    respostas = (
        pg_client.post(f'/api/v1/importacoes-cadastrais/{preparada["sessao_id"]}/confirmar', json={"token": "incorreto"}),
        pg_client.post(f'/api/v1/importacoes-cadastrais/{uuid4()}/confirmar', json={"token": preparada["token_confirmacao"]}),
        pg_client.post('/api/v1/importacoes-cadastrais/uuid-invalido/confirmar', json={"token": preparada["token_confirmacao"]}),
    )
    assert [item.status_code for item in respostas] == [404, 404, 404]
    assert respostas[0].json() == respostas[1].json() == respostas[2].json()
    assert preparada["token_confirmacao"] not in respostas[0].text + respostas[1].text
    with Session(engine) as db:
        sessao = db.scalar(select(SessaoImportacaoCadastral))
        assert sessao.status == "PENDENTE"


def test_revalidacao_detecta_conflito_sem_aplicacao_parcial(pg_client, postgres_app):
    _, engine, _ = postgres_app
    preparada = _preparar(pg_client)
    with Session(engine) as db, db.begin():
        db.add(MateriaPrima(codigo="PG_STAGE_MP", nome="Conflito posterior", ativa=True))
    resposta = pg_client.post(
        f'/api/v1/importacoes-cadastrais/{preparada["sessao_id"]}/confirmar',
        json={"token": preparada["token_confirmacao"]},
    )
    assert resposta.status_code == 409
    with Session(engine) as db:
        assert db.scalar(select(SessaoImportacaoCadastral)).status == "FALHOU"
        assert db.scalar(select(Nutriente).where(Nutriente.codigo == "PG_STAGE_NUT")) is None
        assert len(db.scalars(select(MateriaPrima).where(MateriaPrima.codigo == "PG_STAGE_MP")).all()) == 1


def test_falha_tecnica_apos_flush_faz_rollback_integral(pg_client, postgres_app, monkeypatch):
    _, engine, _ = postgres_app
    preparada = _preparar(pg_client)
    antes = snapshot_cadastros(engine)
    import services.confirmacao_importacao_cadastral as modulo

    def falhar(db, dados, operacoes):
        db.add(Nutriente(codigo="PG_PARCIAL", nome="Não deve persistir", unidade="g"))
        db.flush()
        raise RuntimeError("senha=segredo SQL interno")

    monkeypatch.setattr(modulo, "_aplicar", falhar)
    resposta = pg_client.post(
        f'/api/v1/importacoes-cadastrais/{preparada["sessao_id"]}/confirmar',
        json={"token": preparada["token_confirmacao"]},
    )
    assert resposta.status_code == 500
    assert "segredo" not in resposta.text
    with Session(engine) as db:
        assert db.scalar(select(Nutriente).where(Nutriente.codigo == "PG_PARCIAL")) is None
        assert snapshot_cadastros(engine) == antes
        sessao = db.scalar(select(SessaoImportacaoCadastral))
        assert sessao.status == "FALHOU"
        assert sessao.resultado["tipo"] == "ERRO_TECNICO"


def test_falha_no_commit_reverte_lote_e_registra_falha(pg_client, postgres_app, monkeypatch):
    _, engine, _ = postgres_app
    preparada = _preparar(pg_client)
    antes = snapshot_cadastros(engine)
    commit_original = Session.commit
    chamadas = 0

    def falhar_primeiro_commit(self):
        nonlocal chamadas
        chamadas += 1
        if chamadas == 1:
            raise RuntimeError("falha simulada antes do commit")
        return commit_original(self)

    monkeypatch.setattr(Session, "commit", falhar_primeiro_commit)
    resposta = pg_client.post(
        f'/api/v1/importacoes-cadastrais/{preparada["sessao_id"]}/confirmar',
        json={"token": preparada["token_confirmacao"]},
    )
    assert resposta.status_code == 500
    assert resposta.json() == {"detail": "Não foi possível confirmar a importação cadastral."}
    assert snapshot_cadastros(engine) == antes
    with Session(engine) as db:
        assert db.scalar(select(SessaoImportacaoCadastral)).status == "FALHOU"


def test_falha_ao_registrar_falhou_mantem_resposta_segura_e_rollback(
    pg_client, postgres_app, monkeypatch, caplog,
):
    _, engine, _ = postgres_app
    preparada = _preparar(pg_client)
    antes = snapshot_cadastros(engine)
    import services.confirmacao_importacao_cadastral as servico
    import routers.importacoes_cadastrais as rota

    def falhar_aplicacao(*args, **kwargs):
        raise RuntimeError("falha controlada")

    def falhar_registro(*args, **kwargs):
        raise RuntimeError("falha controlada no registro")

    monkeypatch.setattr(servico, "_aplicar", falhar_aplicacao)
    monkeypatch.setattr(rota, "registrar_falha_tecnica", falhar_registro)
    resposta = pg_client.post(
        f'/api/v1/importacoes-cadastrais/{preparada["sessao_id"]}/confirmar',
        json={"token": preparada["token_confirmacao"]},
    )
    assert resposta.status_code == 500
    assert resposta.json() == {"detail": "Não foi possível confirmar a importação cadastral."}
    assert preparada["token_confirmacao"] not in caplog.text
    assert snapshot_cadastros(engine) == antes
    with Session(engine) as db:
        assert db.scalar(select(SessaoImportacaoCadastral)).status == "PENDENTE"


def test_resultado_truncado_persistido_e_retornado_sao_identicos(
    pg_client, postgres_app, monkeypatch,
):
    _, engine, _ = postgres_app
    import services.confirmacao_importacao_cadastral as modulo
    monkeypatch.setattr(modulo, "MAX_CODIGOS_RESULTADO", 1)
    preparada = _preparar(pg_client)
    resposta = pg_client.post(
        f'/api/v1/importacoes-cadastrais/{preparada["sessao_id"]}/confirmar',
        json={"token": preparada["token_confirmacao"]},
    )
    assert resposta.status_code == 200
    assert resposta.json()["resultado_truncado"] is True
    assert len(resposta.json()["codigos_afetados"]) == 1
    with Session(engine) as db:
        sessao = db.scalar(select(SessaoImportacaoCadastral))
        assert sessao.resultado == resposta.json()


def planilha_operacoes(mps=(), nutrientes=(), composicoes=(), precos=()):
    workbook = load_workbook(io.BytesIO(gerar_template_cadastral()))
    for aba in ("MATERIAS_PRIMAS", "NUTRIENTES", "COMPOSICAO_NUTRICIONAL", "PRECOS_MP"):
        ws = workbook[aba]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
    for linha in mps: workbook["MATERIAS_PRIMAS"].append(linha)
    for linha in nutrientes: workbook["NUTRIENTES"].append(linha)
    for linha in composicoes: workbook["COMPOSICAO_NUTRICIONAL"].append(linha)
    for linha in precos: workbook["PRECOS_MP"].append(linha)
    saida = io.BytesIO(); workbook.save(saida); return saida.getvalue()


def test_confirmacao_atualiza_desativa_zero_e_preserva_historico(pg_client, postgres_app):
    _, engine, _ = postgres_app
    with Session(engine) as db, db.begin():
        mp = MateriaPrima(codigo="PG_ATUAL", nome="MP antiga", ativa=True)
        desativar = MateriaPrima(codigo="PG_DESAT", nome="MP desativar", ativa=True)
        nut = Nutriente(codigo="PG_NUT_AT", nome="Nutriente antigo", unidade="g")
        db.add_all([mp, desativar, nut]); db.flush()
        db.add(ComposicaoMateriaPrima(materia_prima_id=mp.id, nutriente_id=nut.id, valor=5))
        db.add(PrecoMateriaPrima(materia_prima_id=mp.id, preco_kg=10,
                                 vigencia_inicio=date(2026, 1, 1), vigencia_fim=date(2026, 1, 31)))
    arquivo = planilha_operacoes(
        mps=[("ATUALIZAR", "PG_ATUAL", "MP atualizada", None), ("DESATIVAR", "PG_DESAT", None, None)],
        nutrientes=[("ATUALIZAR", "PG_NUT_AT", "Nutriente atualizado", None)],
        composicoes=[("ATUALIZAR", "PG_ATUAL", "PG_NUT_AT", "0.000000")],
        precos=[("CRIAR", "PG_ATUAL", "11.000000", "2026-02-01", None)],
    )
    preparada = pg_client.post("/api/v1/importacoes-cadastrais/preparar", files={"arquivo": ("dados.xlsx", arquivo)}).json()
    resposta = pg_client.post(f'/api/v1/importacoes-cadastrais/{preparada["sessao_id"]}/confirmar', json={"token": preparada["token_confirmacao"]})
    assert resposta.status_code == 200, resposta.text
    with Session(engine) as db:
        mp = db.scalar(select(MateriaPrima).where(MateriaPrima.codigo == "PG_ATUAL"))
        nut = db.scalar(select(Nutriente).where(Nutriente.codigo == "PG_NUT_AT"))
        assert mp.nome == "MP atualizada" and nut.nome == "Nutriente atualizado"
        assert db.scalar(select(ComposicaoMateriaPrima).where(ComposicaoMateriaPrima.materia_prima_id == mp.id)).valor == 0
        precos = list(db.scalars(select(PrecoMateriaPrima).where(PrecoMateriaPrima.materia_prima_id == mp.id).order_by(PrecoMateriaPrima.vigencia_inicio)))
        assert [(p.vigencia_inicio.isoformat(), p.vigencia_fim.isoformat() if p.vigencia_fim else None) for p in precos] == [("2026-01-01", "2026-01-31"), ("2026-02-01", None)]
        assert db.scalar(select(MateriaPrima).where(MateriaPrima.codigo == "PG_DESAT")).ativa is False


def test_sem_alteracao_nao_executa_update(pg_client, postgres_app):
    _, engine, _ = postgres_app
    with Session(engine) as db, db.begin():
        mp = MateriaPrima(codigo="PG_SEM_ALT", nome="Mesmo nome", ativa=True)
        nutriente = Nutriente(codigo="PG_SEM_NUT", nome="Mesmo nutriente", unidade="g")
        db.add_all([mp, nutriente])
        db.flush()
        db.add(ComposicaoMateriaPrima(
            materia_prima_id=mp.id, nutriente_id=nutriente.id, valor=Decimal("1.250000")
        ))
    arquivo = planilha_operacoes(
        mps=[("ATUALIZAR", "PG_SEM_ALT", "Mesmo nome", "SIM")],
        nutrientes=[("ATUALIZAR", "PG_SEM_NUT", "Mesmo nutriente", "g")],
        composicoes=[("ATUALIZAR", "PG_SEM_ALT", "PG_SEM_NUT", "1.250000")],
    )
    preparada = pg_client.post(
        "/api/v1/importacoes-cadastrais/preparar", files={"arquivo": ("dados.xlsx", arquivo)}
    ).json()
    with engine.connect() as conn:
        antes = conn.execute(text(
            "SELECT xmin::text FROM materias_primas WHERE codigo='PG_SEM_ALT'"
        )).scalar_one()
    resposta = pg_client.post(
        f'/api/v1/importacoes-cadastrais/{preparada["sessao_id"]}/confirmar',
        json={"token": preparada["token_confirmacao"]},
    )
    assert resposta.status_code == 200, resposta.text
    with engine.connect() as conn:
        depois = conn.execute(text(
            "SELECT xmin::text FROM materias_primas WHERE codigo='PG_SEM_ALT'"
        )).scalar_one()
    assert depois == antes
    assert resposta.json()["totais"]["sem_alteracao"] == 3


def test_confirmacoes_concorrentes_mesma_sessao_sao_idempotentes(postgres_app):
    app, engine, _ = postgres_app
    with TestClient(app) as client:
        preparada = _preparar(client)
    barreira = Barrier(2)
    def confirmar():
        with TestClient(app) as client:
            barreira.wait(timeout=10)
            resposta = client.post(f'/api/v1/importacoes-cadastrais/{preparada["sessao_id"]}/confirmar', json={"token": preparada["token_confirmacao"]})
            return resposta.status_code, resposta.json()
    with ThreadPoolExecutor(max_workers=2) as pool:
        resultados = [f.result(timeout=30) for f in [pool.submit(confirmar), pool.submit(confirmar)]]
    assert resultados[0] == resultados[1]
    assert resultados[0][0] == 200
    with Session(engine) as db:
        assert len(db.scalars(select(MateriaPrima).where(MateriaPrima.codigo == "PG_STAGE_MP")).all()) == 1
        assert len(db.scalars(select(PrecoMateriaPrima)).all()) == 1


def test_sessoes_concorrentes_conflitantes_nao_aplicam_parcialmente(postgres_app):
    app, engine, _ = postgres_app
    with TestClient(app) as client:
        preparadas = [_preparar(client), _preparar(client)]
    barreira = Barrier(2)
    def confirmar(item):
        with TestClient(app) as client:
            barreira.wait(timeout=10)
            return client.post(f'/api/v1/importacoes-cadastrais/{item["sessao_id"]}/confirmar', json={"token": item["token_confirmacao"]}).status_code
    with ThreadPoolExecutor(max_workers=2) as pool:
        codigos = sorted(f.result(timeout=30) for f in [pool.submit(confirmar, item) for item in preparadas])
    assert codigos == [200, 409]
    with Session(engine) as db:
        assert len(db.scalars(select(MateriaPrima).where(MateriaPrima.codigo == "PG_STAGE_MP")).all()) == 1
        assert len(db.scalars(select(ComposicaoMateriaPrima)).all()) == 1
        assert len(db.scalars(select(PrecoMateriaPrima)).all()) == 1
        assert sorted(db.scalars(select(SessaoImportacaoCadastral.status)).all()) == ["CONFIRMADA", "FALHOU"]


def test_sessao_expirada_nao_confirma(pg_client, postgres_app):
    _, engine, _ = postgres_app
    with Session(engine) as db:
        sessao, token, _ = preparar_sessao(planilha_valida(), "dados.xlsx", db,
                                            agora=datetime.now(timezone.utc) - timedelta(hours=25))
        identificador = sessao.uuid_publico; db.commit()
    resposta = pg_client.post(f"/api/v1/importacoes-cadastrais/{identificador}/confirmar", json={"token": token})
    assert resposta.status_code == 409
    with Session(engine) as db:
        assert db.scalar(select(SessaoImportacaoCadastral)).status == "EXPIRADA"
        assert snapshot_cadastros(engine) == ((), (), (), ())
