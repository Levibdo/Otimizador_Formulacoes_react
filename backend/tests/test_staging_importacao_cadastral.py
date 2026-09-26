import io
import base64
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from enum import Enum
from hashlib import sha256
from uuid import uuid4

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from db.base import Base
from db.session import get_db
from models import MateriaPrima, SessaoImportacaoCadastral
from repositories.importacao_cadastral_repository import ImportacaoCadastralRepository
from routers.importacoes_cadastrais import router
from services.planilha_cadastral import gerar_template_cadastral
from services.staging_importacao_cadastral import preparar_sessao, serializar_canonico


def planilha_valida():
    workbook = load_workbook(io.BytesIO(gerar_template_cadastral()))
    for aba in ("MATERIAS_PRIMAS", "NUTRIENTES", "COMPOSICAO_NUTRICIONAL", "PRECOS_MP"):
        ws = workbook[aba]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
    workbook["MATERIAS_PRIMAS"].append(("CRIAR", "STAGE_MP", "MP staging", "SIM"))
    workbook["NUTRIENTES"].append(("CRIAR", "STAGE_NUT", "Nutriente staging", "g/100 g"))
    workbook["COMPOSICAO_NUTRICIONAL"].append(("CRIAR", "STAGE_MP", "STAGE_NUT", "12.345600"))
    workbook["PRECOS_MP"].append(("CRIAR", "STAGE_MP", "9.500000", "2026-01-01", None))
    saida = io.BytesIO()
    workbook.save(saida)
    return saida.getvalue()


@pytest.fixture
def db():
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    with Session(engine, expire_on_commit=False) as session:
        yield session
    engine.dispose()


@pytest.fixture
def client(db):
    app = FastAPI()
    app.include_router(router)

    def override_get_db():
        yield db

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app, raise_server_exceptions=False) as test_client:
        yield test_client


def test_preparacao_valida_hash_token_payload_e_repeticao(client, db):
    conteudo = planilha_valida()
    respostas = [client.post(
        "/api/v1/importacoes-cadastrais/preparar",
        files={"arquivo": ("cadastros.xlsx", conteudo)},
    ) for _ in range(2)]
    assert [item.status_code for item in respostas] == [201, 201]
    primeira, segunda = [item.json() for item in respostas]
    assert primeira["status"] == "PENDENTE"
    assert primeira["arquivo_sha256"] == sha256(conteudo).hexdigest()
    assert primeira["operacoes_total"] == 4
    assert primeira["token_confirmacao"] != segunda["token_confirmacao"]
    assert primeira["sessao_id"] != segunda["sessao_id"]
    token_bytes = base64.urlsafe_b64decode(primeira["token_confirmacao"] + "==")
    assert len(token_bytes) == 32

    salvas = list(db.scalars(select(SessaoImportacaoCadastral).order_by(
        SessaoImportacaoCadastral.id
    )).all())
    assert len(salvas) == 2
    assert all(item.token_hash not in (primeira["token_confirmacao"], segunda["token_confirmacao"])
               for item in salvas)
    assert ImportacaoCadastralRepository.token_valido(salvas[0], primeira["token_confirmacao"])
    assert not ImportacaoCadastralRepository.token_valido(salvas[0], segunda["token_confirmacao"])
    assert salvas[0].payload_normalizado["dados"]["MATERIAS_PRIMAS"][0]["codigo"] == "STAGE_MP"
    assert salvas[0].payload_normalizado["dados"]["COMPOSICAO_NUTRICIONAL"][0]["valor"] == "12.345600"
    assert db.scalar(select(func.count()).select_from(MateriaPrima)) == 0


def test_endpoint_validar_permanece_sem_staging(client, db):
    resposta = client.post(
        "/api/v1/importacoes-cadastrais/validar",
        files={"arquivo": ("cadastros.xlsx", planilha_valida())},
    )
    assert resposta.status_code == 200
    assert resposta.json()["valido_para_confirmacao"] is True
    assert db.scalar(select(func.count()).select_from(SessaoImportacaoCadastral)) == 0


def test_planilha_invalida_nao_cria_sessao(client, db):
    resposta = client.post(
        "/api/v1/importacoes-cadastrais/preparar",
        files={"arquivo": ("cadastros.xlsx", "arquivo inválido".encode())},
    )
    assert resposta.status_code == 422
    assert db.scalar(select(func.count()).select_from(SessaoImportacaoCadastral)) == 0


def test_consulta_nao_expoe_segredos_payload_ou_digest(client, db):
    preparada = client.post(
        "/api/v1/importacoes-cadastrais/preparar",
        files={"arquivo": ("cadastros.xlsx", planilha_valida())},
    ).json()
    resposta = client.get(f'/api/v1/importacoes-cadastrais/{preparada["sessao_id"]}')
    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["status"] == "PENDENTE"
    assert not ({"token_confirmacao", "token_hash", "payload_normalizado"} & corpo.keys())
    assert preparada["token_confirmacao"] not in resposta.text


def test_expiracao_em_24_horas_e_uuid_inexistente(client, db):
    agora = datetime.now(timezone.utc) - timedelta(hours=25)
    sessao, _, _ = preparar_sessao(planilha_valida(), "dados.xlsx", db, agora=agora)
    db.commit()
    resposta = client.get(f"/api/v1/importacoes-cadastrais/{sessao.uuid_publico}")
    assert resposta.status_code == 200
    assert resposta.json()["status"] == "EXPIRADA"
    repetida = client.get(f"/api/v1/importacoes-cadastrais/{sessao.uuid_publico}")
    assert repetida.status_code == 200
    assert repetida.json() == resposta.json()
    assert sessao.expira_em - sessao.criado_em == timedelta(hours=24)

    ausente = client.get(f"/api/v1/importacoes-cadastrais/{uuid4()}")
    assert ausente.status_code == 404
    assert ausente.json() == {"detail": "Sessão de importação não encontrada."}


def test_sessao_confirmada_nunca_expira(client, db):
    sessao, _, _ = preparar_sessao(
        planilha_valida(), "dados.xlsx", db,
        agora=datetime.now(timezone.utc) - timedelta(hours=25),
    )
    sessao.status = "CONFIRMADA"
    sessao.confirmado_em = datetime.now(timezone.utc) - timedelta(hours=1)
    sessao.resultado = {"aplicadas": 4}
    db.commit()
    primeira = client.get(f"/api/v1/importacoes-cadastrais/{sessao.uuid_publico}")
    segunda = client.get(f"/api/v1/importacoes-cadastrais/{sessao.uuid_publico}")
    assert primeira.status_code == segunda.status_code == 200
    assert primeira.json()["status"] == segunda.json()["status"] == "CONFIRMADA"


def test_rollback_e_mensagem_segura_quando_criacao_falha(client, db, monkeypatch):
    original = ImportacaoCadastralRepository.criar

    def falhar(self, **campos):
        original(self, **campos)
        raise RuntimeError("senha=segredo SQL /caminho/interno")

    monkeypatch.setattr(ImportacaoCadastralRepository, "criar", falhar)
    resposta = client.post(
        "/api/v1/importacoes-cadastrais/preparar",
        files={"arquivo": ("cadastros.xlsx", planilha_valida())},
    )
    assert resposta.status_code == 500
    assert resposta.json() == {"detail": "Não foi possível preparar a importação cadastral."}
    assert "segredo" not in resposta.text
    assert db.scalar(select(func.count()).select_from(SessaoImportacaoCadastral)) == 0


def test_serializacao_json_canonica_de_decimal_data_datetime_e_enum():
    class Situacao(str, Enum):
        OK = "OK"

    valor = {
        "z": Decimal("1.230000"),
        "a": [date(2026, 1, 2), datetime(2026, 1, 2, 3, 4, tzinfo=timezone.utc), Situacao.OK],
    }
    normalizado, texto = serializar_canonico(valor)
    repetido, texto_repetido = serializar_canonico(valor)
    assert normalizado == repetido == {"a": ["2026-01-02", "2026-01-02T03:04:00Z", "OK"], "z": "1.230000"}
    assert texto == texto_repetido == '{"a":["2026-01-02","2026-01-02T03:04:00Z","OK"],"z":"1.230000"}'
    _, reordenado = serializar_canonico({"a": valor["a"], "z": valor["z"]})
    assert reordenado == texto


@pytest.mark.parametrize("valor", [float("nan"), float("inf"), float("-inf")])
def test_serializacao_rejeita_numeros_nao_finitos(valor):
    with pytest.raises(ValueError, match="não finito"):
        serializar_canonico({"valor": valor})


def test_serializacao_rejeita_tipos_imprevistos_e_datetime_sem_timezone():
    with pytest.raises(TypeError, match="não serializável"):
        serializar_canonico({"valor": object()})
    with pytest.raises(TypeError, match="sem timezone"):
        serializar_canonico({"valor": datetime(2026, 1, 1)})


def test_rotas_estaticas_nao_sao_capturadas_pela_rota_uuid(client):
    assert client.get("/api/v1/importacoes-cadastrais/template").status_code == 200
    assert client.get("/api/v1/importacoes-cadastrais/validar").status_code == 405
    assert client.get("/api/v1/importacoes-cadastrais/preparar").status_code == 405
    invalido = client.get("/api/v1/importacoes-cadastrais/nao-e-uuid")
    assert invalido.status_code == 404
    assert invalido.json() == {"detail": "Not Found"}


def test_token_nao_aparece_na_representacao_do_model(client, db):
    preparada = client.post(
        "/api/v1/importacoes-cadastrais/preparar",
        files={"arquivo": ("cadastros.xlsx", planilha_valida())},
    ).json()
    sessao = db.scalar(select(SessaoImportacaoCadastral))
    assert preparada["token_confirmacao"] not in repr(sessao)
