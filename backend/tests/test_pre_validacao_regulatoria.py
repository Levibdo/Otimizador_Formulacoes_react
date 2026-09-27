import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from db.base import Base
from models import MateriaPrima, SessaoImportacaoCadastral
from models.regulatorio import (
    CategoriaProduto, ComponenteRegulatorio, ComposicaoComponenteMP, RegraRegulatoria,
)
from services.planilha_cadastral import gerar_template_cadastral
from services.pre_validacao_cadastral import (
    pre_validar_planilha_cadastral, pre_validar_planilha_cadastral_completo,
)
from services.staging_importacao_cadastral import preparar_sessao


ABAS_DADOS = (
    "MATERIAS_PRIMAS", "NUTRIENTES", "COMPOSICAO_NUTRICIONAL", "PRECOS_MP",
    "CATEGORIAS_PRODUTO", "COMPONENTES_REGULATORIOS", "COMPOSICAO_COMPONENTES_MP",
    "REGRAS_REGULATORIAS_MP", "REGRAS_REGULATORIAS_COMPONENTE",
)


@pytest.fixture
def db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine, expire_on_commit=False) as session:
        yield session
    engine.dispose()


def planilha(**linhas):
    wb = load_workbook(io.BytesIO(gerar_template_cadastral("1.1")))
    for aba in ABAS_DADOS:
        ws = wb[aba]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        for linha in linhas.get(aba, ()):
            ws.append(linha)
    saida = io.BytesIO(); wb.save(saida); return saida.getvalue()


def validar(db, **linhas):
    return pre_validar_planilha_cadastral(planilha(**linhas), "regulatorio.xlsx", db)


def codigos(resultado):
    return {item["codigo"] for item in resultado["diagnosticos"]}


def cadastrar_base(db, *, ativos=True):
    mp = MateriaPrima(codigo="MP_EXIST", nome="MP existente", ativa=ativos)
    cat = CategoriaProduto(codigo="CAT_EXIST", nome="Categoria existente", descricao="Descrição", ativa=ativos)
    comp = ComponenteRegulatorio(codigo="COMP_EXIST", nome="Componente existente", descricao="Descrição", ativo=ativos)
    db.add_all([mp, cat, comp]); db.commit()
    return mp, cat, comp


def test_referencias_criadas_no_lote_e_ordem_do_resumo(db):
    resultado = validar(db,
        MATERIAS_PRIMAS=[("CRIAR", "MP_0001", "MP nova", "SIM")],
        CATEGORIAS_PRODUTO=[("CRIAR", "CAT_0001", "Categoria nova", None)],
        COMPONENTES_REGULATORIOS=[("CRIAR", "COMP_0001", "Componente novo", None)],
        COMPOSICAO_COMPONENTES_MP=[
            ("CRIAR", "MP_0001", "COMP_0001", "2026-01-01", "INFORMADO", "1.234567", None, None),
            ("CRIAR", "MP_0001", "COMP_0001", "2026-01-02", "DESCONHECIDO", None, None, None),
            ("CRIAR", "MP_0001", "COMP_0001", "2026-01-03", "AUSENTE_CONFIRMADO", "0", None, None),
        ],
        REGRAS_REGULATORIAS_MP=[
            ("CRIAR", "CAT_0001", "MP_0001", "LIMITADA", None, "10", "Justificativa", None, None, None),
        ],
        REGRAS_REGULATORIAS_COMPONENTE=[
            ("CRIAR", "CAT_0001", "COMP_0001", "PROIBIDA", None, None, "Justificativa", None, None, None),
        ],
    )
    assert resultado["valido_para_confirmacao"], resultado["diagnosticos"]
    assert list(resultado["resumo"])[:10] == [
        "LEIA_ME", "MATERIAS_PRIMAS", "NUTRIENTES", "CATEGORIAS_PRODUTO",
        "COMPONENTES_REGULATORIOS", "COMPOSICAO_NUTRICIONAL", "PRECOS_MP",
        "COMPOSICAO_COMPONENTES_MP", "REGRAS_REGULATORIAS_MP",
        "REGRAS_REGULATORIAS_COMPONENTE",
    ]
    assert resultado["resumo"]["COMPOSICAO_COMPONENTES_MP"]["criar"] == 3
    assert not db.new and not db.dirty and not db.deleted


def test_catalogos_criar_identico_atualizar_desativar_e_inativos(db):
    _, cat, comp = cadastrar_base(db)
    identico = validar(db,
        CATEGORIAS_PRODUTO=[("CRIAR", "CAT_EXIST", "Categoria existente", "Descrição")],
        COMPONENTES_REGULATORIOS=[("CRIAR", "COMP_EXIST", "Componente existente", "Descrição")],
    )
    assert identico["resumo"]["CATEGORIAS_PRODUTO"]["sem_alteracao"] == 1
    assert identico["resumo"]["COMPONENTES_REGULATORIOS"]["sem_alteracao"] == 1

    alterado = validar(db,
        CATEGORIAS_PRODUTO=[("ATUALIZAR", "CAT_EXIST", None, "Nova descrição")],
        COMPONENTES_REGULATORIOS=[("DESATIVAR", "COMP_EXIST", None, None)],
    )
    assert alterado["resumo"]["CATEGORIAS_PRODUTO"]["atualizar"] == 1
    assert alterado["resumo"]["COMPONENTES_REGULATORIOS"]["desativar"] == 1
    assert cat.descricao == "Descrição" and comp.ativo is True

    cat.ativa = False; comp.ativo = False; db.commit()
    inativo = validar(db,
        CATEGORIAS_PRODUTO=[("CRIAR", "CAT_EXIST", "Categoria existente", "Descrição")],
        COMPONENTES_REGULATORIOS=[("ATUALIZAR", "COMP_EXIST", "Outro", None)],
    )
    assert {"REATIVACAO_NAO_SUPORTADA", "REGISTRO_INATIVO"} <= codigos(inativo)


def test_referencia_desativada_no_lote_e_dependencia_invalida(db):
    cadastrar_base(db)
    resultado = validar(db,
        COMPONENTES_REGULATORIOS=[("DESATIVAR", "COMP_EXIST", None, None)],
        COMPOSICAO_COMPONENTES_MP=[
            ("CRIAR", "MP_EXIST", "COMP_EXIST", "2026-01-01", "INFORMADO", "1", None, None),
        ],
    )
    assert "REFERENCIA_DESATIVADA_NO_LOTE" in codigos(resultado)
    assert resultado["resumo"]["COMPONENTES_REGULATORIOS"]["desativar"] == 1
    assert resultado["resumo"]["COMPOSICAO_COMPONENTES_MP"]["criar"] == 0


def test_composicao_identica_conflitante_inativa_e_desativacao(db):
    mp, _, comp = cadastrar_base(db)
    registro = ComposicaoComponenteMP(
        materia_prima_id=mp.id, componente_id=comp.id, data_referencia=date(2026, 1, 1),
        situacao="INFORMADO", concentracao=Decimal("1.234567"), fonte="Fonte", observacao=None,
    )
    db.add(registro); db.commit()
    igual = validar(db, COMPOSICAO_COMPONENTES_MP=[
        ("CRIAR", "MP_EXIST", "COMP_EXIST", "2026-01-01", "INFORMADO", "1.234567", "Fonte", None),
    ])
    assert igual["resumo"]["COMPOSICAO_COMPONENTES_MP"]["sem_alteracao"] == 1
    conflito = validar(db, COMPOSICAO_COMPONENTES_MP=[
        ("CRIAR", "MP_EXIST", "COMP_EXIST", "2026-01-01", "INFORMADO", "2", "Fonte", None),
    ])
    assert "ATUALIZACAO_IMUTAVEL" in codigos(conflito)
    registro.ativo = False; db.commit()
    inativa = validar(db, COMPOSICAO_COMPONENTES_MP=[
        ("CRIAR", "MP_EXIST", "COMP_EXIST", "2026-01-01", "INFORMADO", "1.234567", "Fonte", None),
    ])
    assert "REATIVACAO_NAO_SUPORTADA" in codigos(inativa)
    desativar = validar(db, COMPOSICAO_COMPONENTES_MP=[
        ("DESATIVAR", "MP_EXIST", "COMP_EXIST", "2026-01-01", None, None, None, None),
    ])
    assert desativar["resumo"]["COMPOSICAO_COMPONENTES_MP"]["sem_alteracao"] == 1


def criar_regra(db, cat, *, mp=None, comp=None, ativa=True, inicio=None, fim=None, maximo="10", revisao=1):
    regra = RegraRegulatoria(
        categoria_id=cat.id, tipo_alvo="MATERIA_PRIMA" if mp else "COMPONENTE",
        materia_prima_id=mp.id if mp else None, componente_id=comp.id if comp else None,
        tratamento="LIMITADA", minimo=None, maximo=Decimal(maximo), unidade="%",
        base="MASSA_MASSA", justificativa="Justificativa", referencia_normativa=None,
        vigencia_inicio=inicio, vigencia_fim=fim, ativa=ativa, revisao=revisao,
    )
    db.add(regra); db.commit(); return regra


def test_regra_identica_revisao_prevista_sobreposicao_e_historico(db):
    mp, cat, comp = cadastrar_base(db)
    regra = criar_regra(db, cat, mp=mp)
    igual = validar(db, REGRAS_REGULATORIAS_MP=[
        ("CRIAR", "CAT_EXIST", "MP_EXIST", "LIMITADA", None, "10", "Justificativa", None, None, None),
    ])
    assert igual["resumo"]["REGRAS_REGULATORIAS_MP"]["sem_alteracao"] == 1
    revisao = validar(db, REGRAS_REGULATORIAS_MP=[
        ("ATUALIZAR", "CAT_EXIST", "MP_EXIST", "LIMITADA", None, "20", "Nova justificativa", None, None, None),
    ])
    assert revisao["resumo"]["REGRAS_REGULATORIAS_MP"]["atualizar"] == 1
    assert regra.maximo == Decimal("10")

    sobreposta = validar(db, REGRAS_REGULATORIAS_MP=[
        ("CRIAR", "CAT_EXIST", "MP_EXIST", "LIMITADA", None, "5", "Outra", None, "2026-01-01", None),
    ])
    assert "REGRA_SOBREPOSTA" in codigos(sobreposta)
    regra.ativa = False; db.commit()
    historico = validar(db, REGRAS_REGULATORIAS_MP=[
        ("DESATIVAR", "CAT_EXIST", "MP_EXIST", None, None, None, None, None, None, None),
    ])
    assert historico["resumo"]["REGRAS_REGULATORIAS_MP"]["sem_alteracao"] == 1


def test_troca_de_periodo_desativar_e_criar_e_regra_por_componente(db):
    mp, cat, comp = cadastrar_base(db)
    criar_regra(db, cat, mp=mp, inicio=date(2026, 1, 1), fim=date(2026, 6, 30))
    resultado = validar(db,
        REGRAS_REGULATORIAS_MP=[
            ("DESATIVAR", "CAT_EXIST", "MP_EXIST", None, None, None, None, None, "2026-01-01", "2026-06-30"),
            ("CRIAR", "CAT_EXIST", "MP_EXIST", "LIMITADA", None, "15", "Nova vigência", None, "2026-07-01", None),
        ],
        REGRAS_REGULATORIAS_COMPONENTE=[
            ("CRIAR", "CAT_EXIST", "COMP_EXIST", "PROIBIDA", None, None, "Proibição", None, None, None),
        ],
    )
    assert resultado["valido_para_confirmacao"], resultado["diagnosticos"]
    assert resultado["resumo"]["REGRAS_REGULATORIAS_MP"]["desativar"] == 1
    assert resultado["resumo"]["REGRAS_REGULATORIAS_MP"]["criar"] == 1
    assert resultado["resumo"]["REGRAS_REGULATORIAS_COMPONENTE"]["criar"] == 1


def test_preparacao_v11_reutiliza_pre_validacao_sem_aplicar_cadastros(db):
    conteudo = planilha(
        CATEGORIAS_PRODUTO=[("CRIAR", "CAT_0001", "Categoria", None)],
    )
    resposta = pre_validar_planilha_cadastral(conteudo, "x.xlsx", db)
    assert resposta["valido_para_confirmacao"]
    completo, internos = pre_validar_planilha_cadastral_completo(conteudo, "x.xlsx", db)
    assert completo == resposta
    assert internos["dados"]["CATEGORIAS_PRODUTO"][0]["codigo"] == "CAT_0001"
    sessao, token, validacao = preparar_sessao(conteudo, "x.xlsx", db)
    assert sessao.versao_contrato == "1.1"
    assert token and validacao["valido_para_confirmacao"]
    assert db.scalar(select(SessaoImportacaoCadastral)) is sessao


def test_chave_duplicada_tem_codigo_estavel(db):
    resultado = validar(db, CATEGORIAS_PRODUTO=[
        ("CRIAR", "CAT_DUP", "Uma", None),
        ("DESATIVAR", "CAT_DUP", None, None),
    ])
    assert "CHAVE_DUPLICADA" in codigos(resultado)
    assert "ACAO_CONFLITANTE" in codigos(resultado)


def test_truncamento_preserva_totais_v11(db):
    linhas = [
        ("CRIAR", f"CAT_{indice:04d}", f"Categoria {indice}", None)
        for indice in range(510)
    ]
    resultado = validar(db, CATEGORIAS_PRODUTO=linhas)
    assert resultado["valido_para_confirmacao"]
    assert resultado["operacoes_total"] == 510
    assert len(resultado["operacoes"]) == 500
    assert resultado["resultado_truncado"] is True
    assert resultado["resumo"]["CATEGORIAS_PRODUTO"]["criar"] == 510


def test_operacoes_invalidas_nao_contaminam_estado_projetado(db):
    mp, cat, comp = cadastrar_base(db)
    criar_regra(db, cat, mp=mp, inicio=date(2026, 1, 1), fim=date(2026, 12, 31))

    categoria_invalida = validar(db,
        CATEGORIAS_PRODUTO=[
            ("CRIAR", "CAT_INVALIDA", None, None),
            ("CRIAR", "CAT_INDEPENDENTE", "Categoria independente", None),
        ],
        REGRAS_REGULATORIAS_COMPONENTE=[
            ("CRIAR", "CAT_INVALIDA", "COMP_EXIST", "PROIBIDA", None, None,
             "Referência inválida", None, None, None),
        ],
    )
    assert "DEPENDENCIA_INVALIDA" in codigos(categoria_invalida)
    assert categoria_invalida["resumo"]["CATEGORIAS_PRODUTO"]["criar"] == 1
    assert categoria_invalida["resumo"]["REGRAS_REGULATORIAS_COMPONENTE"]["criar"] == 0

    componente_conflitante = validar(db,
        COMPONENTES_REGULATORIOS=[
            ("CRIAR", "COMP_EXIST", "Conteúdo conflitante", None),
        ],
        COMPOSICAO_COMPONENTES_MP=[
            ("CRIAR", "MP_EXIST", "COMP_EXIST", "2026-02-01", "INFORMADO", "1", None, None),
        ],
    )
    assert {"CRIAR_CONFLITA_EXISTENTE", "DEPENDENCIA_INVALIDA"} <= codigos(componente_conflitante)

    regra = validar(db, REGRAS_REGULATORIAS_MP=[
        ("DESATIVAR", "CAT_EXIST", "MP_EXIST", "LIMITADA", None, None,
         None, None, "2026-01-01", "2026-12-31"),
        ("CRIAR", "CAT_EXIST", "MP_EXIST", "LIMITADA", None, "20",
         "Nova sobreposta", None, "2026-06-01", None),
    ])
    assert "REGRA_SOBREPOSTA" in codigos(regra)
    assert regra["resumo"]["REGRAS_REGULATORIAS_MP"]["desativar"] == 0


def test_mp_invalida_nao_fica_disponivel_para_regra(db):
    _, cat, _ = cadastrar_base(db)
    resultado = validar(db,
        MATERIAS_PRIMAS=[("CRIAR", "MP_INVALIDA", None, "SIM")],
        REGRAS_REGULATORIAS_MP=[
            ("CRIAR", "CAT_EXIST", "MP_INVALIDA", "LIMITADA", None, "10",
             "Referência", None, None, None),
        ],
    )
    assert "DEPENDENCIA_INVALIDA" in codigos(resultado)
    assert resultado["resumo"]["REGRAS_REGULATORIAS_MP"]["criar"] == 0


def test_atualizacao_invalida_nao_altera_estado_e_dependencia_falha(db):
    _, _, comp = cadastrar_base(db)
    resultado = validar(db,
        COMPONENTES_REGULATORIOS=[
            ("ATUALIZAR", "COMP_EXIST", "X" * 201, None),
        ],
        COMPOSICAO_COMPONENTES_MP=[
            ("CRIAR", "MP_EXIST", "COMP_EXIST", "2026-02-01", "INFORMADO", "1", None, None),
        ],
    )
    assert {"VALOR_EXCEDE_LIMITE", "DEPENDENCIA_INVALIDA"} <= codigos(resultado)
    assert comp.nome == "Componente existente"


def test_linhas_duplicadas_independem_da_ordem(db):
    linhas_a = [
        ("CRIAR", "CAT_DUP_2", "Primeira", None),
        ("CRIAR", "CAT_DUP_2", "Segunda", None),
    ]
    linhas_b = list(reversed(linhas_a))
    resultados = [validar(db, CATEGORIAS_PRODUTO=linhas) for linhas in (linhas_a, linhas_b)]
    for resultado in resultados:
        assert "CHAVE_DUPLICADA" in codigos(resultado)
        assert resultado["operacoes_total"] == 0
        assert resultado["resumo"]["CATEGORIAS_PRODUTO"]["erros"] == 1


def test_teto_de_consultas_e_ausencia_de_n_mais_um(db):
    consultas = 0
    def contar(conn, cursor, statement, parameters, context, executemany):
        nonlocal consultas
        consultas += 1
    event.listen(db.bind, "before_cursor_execute", contar)
    try:
        resultado = validar(db,
            MATERIAS_PRIMAS=[
                ("CRIAR", f"MP_Q_{i:03d}", f"MP consulta {i}", "SIM") for i in range(40)
            ],
            CATEGORIAS_PRODUTO=[
                ("CRIAR", f"CAT_Q_{i:03d}", f"Categoria consulta {i}", None) for i in range(40)
            ],
            COMPONENTES_REGULATORIOS=[
                ("CRIAR", f"COMP_Q_{i:03d}", f"Componente consulta {i}", None) for i in range(40)
            ],
        )
    finally:
        event.remove(db.bind, "before_cursor_execute", contar)
    assert resultado["valido_para_confirmacao"]
    assert consultas <= 10


def test_fluxo_validar_nao_chama_metodos_de_escrita():
    class SessaoSomenteLeitura(Session):
        def add(self, *args, **kwargs):
            raise AssertionError("add não permitido")
        def add_all(self, *args, **kwargs):
            raise AssertionError("add_all não permitido")
        def delete(self, *args, **kwargs):
            raise AssertionError("delete não permitido")
        def flush(self, *args, **kwargs):
            raise AssertionError("flush não permitido")
        def commit(self, *args, **kwargs):
            raise AssertionError("commit não permitido")

    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with SessaoSomenteLeitura(engine) as somente_leitura:
        resultado = validar(
            somente_leitura,
            CATEGORIAS_PRODUTO=[("CRIAR", "CAT_READ_ONLY", "Categoria", None)],
        )
        assert resultado["valido_para_confirmacao"]
        assert not somente_leitura.new and not somente_leitura.dirty and not somente_leitura.deleted
    engine.dispose()


def test_precisao_numeric_18_6_rejeitada_antes_do_banco(db):
    cadastrar_base(db)
    resultado = validar(db,
        COMPOSICAO_COMPONENTES_MP=[
            ("CRIAR", "MP_EXIST", "COMP_EXIST", "2026-01-01", "INFORMADO",
             "1.1234567", None, None),
        ],
        REGRAS_REGULATORIAS_COMPONENTE=[
            ("CRIAR", "CAT_EXIST", "COMP_EXIST", "LIMITADA", None,
             "1.1234567", "Teste", None, None, None),
        ],
    )
    assert [item["codigo"] for item in resultado["diagnosticos"]].count(
        "VALOR_EXCEDE_PRECISAO"
    ) == 2
