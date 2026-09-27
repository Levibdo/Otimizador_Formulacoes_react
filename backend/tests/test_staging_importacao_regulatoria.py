import io
from datetime import timedelta
from hashlib import sha256

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from db.base import Base
from db.session import get_db
from models import MateriaPrima, SessaoImportacaoCadastral
from routers.importacoes_cadastrais import router
from services.confirmacao_importacao_cadastral import _xlsx_do_payload
from services.planilha_cadastral import gerar_template_cadastral, parsear_planilha_cadastral

@pytest.fixture
def db():
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    with Session(engine, expire_on_commit=False) as session:
        yield session
    engine.dispose()


@pytest.fixture
def client(db):
    app = FastAPI()
    app.include_router(router)
    def override_get_db():
        yield db
    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app, raise_server_exceptions=False) as test_client:
        yield test_client


ABAS_DADOS = (
    "MATERIAS_PRIMAS", "NUTRIENTES", "COMPOSICAO_NUTRICIONAL", "PRECOS_MP",
    "CATEGORIAS_PRODUTO", "COMPONENTES_REGULATORIOS",
    "COMPOSICAO_COMPONENTES_MP", "REGRAS_REGULATORIAS_MP",
    "REGRAS_REGULATORIAS_COMPONENTE",
)

def planilha_v11():
    workbook = load_workbook(io.BytesIO(gerar_template_cadastral("1.1")))
    for aba in ABAS_DADOS:
        ws = workbook[aba]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
    workbook["MATERIAS_PRIMAS"].append(("CRIAR", "MP_0001", "Proteína çã", "SIM"))
    workbook["CATEGORIAS_PRODUTO"].append(("CRIAR", "CAT_0001", "Categoria çã", "Descrição Unicode"))
    workbook["COMPONENTES_REGULATORIOS"].append(("CRIAR", "COMP_0001", "Componente µ", None))
    workbook["COMPOSICAO_COMPONENTES_MP"].append(("CRIAR", "MP_0001", "COMP_0001", "2026-01-01", "INFORMADO", "1.230000", "Fonte", None))
    workbook["COMPOSICAO_COMPONENTES_MP"].append(("CRIAR", "MP_0001", "COMP_0001", "2026-01-02", "DESCONHECIDO", None, None, ""))
    workbook["COMPOSICAO_COMPONENTES_MP"].append(("CRIAR", "MP_0001", "COMP_0001", "2026-01-03", "AUSENTE_CONFIRMADO", "0", None, None))
    workbook["REGRAS_REGULATORIAS_MP"].append(("CRIAR", "CAT_0001", "MP_0001", "LIMITADA", "0.000000", "10.000000", "Regra MP", None, None, None))
    workbook["REGRAS_REGULATORIAS_COMPONENTE"].append(("CRIAR", "CAT_0001", "COMP_0001", "PROIBIDA", None, None, "Regra componente", "RDC çã", "2026-01-01", None))
    saida = io.BytesIO()
    workbook.save(saida)
    return saida.getvalue()

def test_preparacao_v11_persiste_payload_completo_e_round_trip(client, db, caplog):
    conteudo = planilha_v11()
    resposta = client.post("/api/v1/importacoes-cadastrais/preparar", files={"arquivo": ("regulatorio.xlsx", conteudo)})
    assert resposta.status_code == 201, resposta.text
    corpo = resposta.json()
    assert corpo["versao"] == "1.1"
    assert corpo["arquivo_sha256"] == sha256(conteudo).hexdigest()
    assert corpo["status"] == "PENDENTE"
    assert corpo["operacoes_total"] == 8
    sessao = db.scalar(select(SessaoImportacaoCadastral))
    assert sessao.versao_contrato == "1.1"
    assert sessao.status == "PENDENTE"
    assert sessao.expira_em - sessao.criado_em == timedelta(hours=24)
    payload = sessao.payload_normalizado
    assert payload["versao"] == "1.1"
    assert set(payload["dados"]) == set(ABAS_DADOS)
    assert payload["dados"]["MATERIAS_PRIMAS"][0]["codigo"] == "MP_0001"
    composicoes = payload["dados"]["COMPOSICAO_COMPONENTES_MP"]
    assert [item["concentracao"] for item in composicoes] == ["1.230000", None, "0"]
    assert payload["dados"]["REGRAS_REGULATORIAS_MP"][0]["vigencia_inicio"] is None
    assert "token" not in str(payload).lower()
    assert corpo["token_confirmacao"] not in str(payload)
    assert corpo["token_confirmacao"] not in caplog.text
    reconstruido = _xlsx_do_payload(payload)
    reparsed = parsear_planilha_cadastral(reconstruido, "canonico-v11.xlsx")
    assert reparsed["valido"], reparsed["diagnosticos"]
    assert reparsed["dados"] == payload["dados"]

def test_duas_preparacoes_v11_criam_sessoes_e_credenciais_distintas(client, db):
    conteudo = planilha_v11()
    respostas = [client.post("/api/v1/importacoes-cadastrais/preparar", files={"arquivo": ("regulatorio.xlsx", conteudo)}).json() for _ in range(2)]
    assert respostas[0]["sessao_id"] != respostas[1]["sessao_id"]
    assert respostas[0]["token_confirmacao"] != respostas[1]["token_confirmacao"]
    assert respostas[0]["arquivo_sha256"] == respostas[1]["arquivo_sha256"]
    sessoes = list(db.scalars(select(SessaoImportacaoCadastral).order_by(SessaoImportacaoCadastral.id)))
    assert len(sessoes) == 2
    assert sessoes[0].payload_normalizado == sessoes[1].payload_normalizado
    assert sessoes[0].token_hash != sessoes[1].token_hash

def test_confirmacao_v11_autentica_antes_de_bloquear_fluxo(client, db, monkeypatch):
    preparada = client.post("/api/v1/importacoes-cadastrais/preparar", files={"arquivo": ("regulatorio.xlsx", planilha_v11())}).json()
    import services.confirmacao_importacao_cadastral as modulo
    def proibido(*args, **kwargs):
        raise AssertionError("não deve bloquear ou revalidar cadastros v1.1")
    monkeypatch.setattr(modulo, "_bloquear_cadastros", proibido)
    monkeypatch.setattr(modulo, "pre_validar_planilha_cadastral_completo", proibido)
    caminho = f'/api/v1/importacoes-cadastrais/{preparada["sessao_id"]}/confirmar'
    invalida = client.post(caminho, json={"token": "token-invalido"})
    assert invalida.status_code == 404
    assert invalida.json() == {"detail": "Sessão ou credencial de confirmação inválida."}
    respostas = [client.post(caminho, json={"token": preparada["token_confirmacao"]}) for _ in range(2)]
    assert [item.status_code for item in respostas] == [409, 409]
    assert respostas[0].json() == respostas[1].json() == {"detail": {"codigo": "CONFIRMACAO_V11_NAO_DISPONIVEL", "mensagem": "A confirmação do contrato 1.1 ainda não está disponível."}}
    sessao = db.scalar(select(SessaoImportacaoCadastral))
    assert sessao.status == "PENDENTE"
    assert sessao.resultado is None
    assert db.scalar(select(func.count()).select_from(MateriaPrima)) == 0

def test_falha_de_serializacao_v11_nao_cria_sessao(client, db, monkeypatch):
    import services.staging_importacao_cadastral as modulo
    def falhar(*args, **kwargs):
        raise TypeError("conteúdo secreto não serializável")
    monkeypatch.setattr(modulo, "serializar_canonico", falhar)
    resposta = client.post("/api/v1/importacoes-cadastrais/preparar", files={"arquivo": ("regulatorio.xlsx", planilha_v11())})
    assert resposta.status_code == 500
    assert resposta.json() == {"detail": "Não foi possível preparar a importação cadastral."}
    assert "secreto" not in resposta.text
    assert db.scalar(select(SessaoImportacaoCadastral)) is None

def test_falha_de_commit_v11_remove_sessao_parcial(client, db, monkeypatch):
    def falhar_commit():
        raise RuntimeError("falha controlada antes do commit")
    monkeypatch.setattr(db, "commit", falhar_commit)
    resposta = client.post("/api/v1/importacoes-cadastrais/preparar", files={"arquivo": ("regulatorio.xlsx", planilha_v11())})
    assert resposta.status_code == 500
    assert resposta.json() == {"detail": "Não foi possível preparar a importação cadastral."}
    assert "segredo" not in resposta.text
    assert db.scalar(select(SessaoImportacaoCadastral)) is None

def test_preparacao_revalida_conteudo_e_estado_atual(client, db):
    conteudo = planilha_v11()
    validada = client.post("/api/v1/importacoes-cadastrais/validar", files={"arquivo": ("regulatorio.xlsx", conteudo)})
    assert validada.status_code == 200
    assert validada.json()["valido_para_confirmacao"]
    db.add(MateriaPrima(codigo="MP_0001", nome="Conflito posterior", ativa=True))
    db.commit()
    preparada = client.post("/api/v1/importacoes-cadastrais/preparar", files={"arquivo": ("regulatorio.xlsx", conteudo)})
    assert preparada.status_code == 422
    assert db.scalar(select(SessaoImportacaoCadastral)) is None


def test_preparacao_v11_persiste_aviso_seguro(client, db, monkeypatch):
    import services.staging_importacao_cadastral as modulo
    original = modulo.pre_validar_planilha_cadastral_completo

    def com_aviso(*args, **kwargs):
        validacao, internos = original(*args, **kwargs)
        aviso = {
            "severidade": "AVISO", "aba": "GERAL", "linha": None,
            "coluna": None, "codigo": "AVISO_TESTE",
            "mensagem": "Aviso seguro para revisão.",
        }
        validacao = {**validacao, "diagnosticos": [*validacao["diagnosticos"], aviso]}
        internos = {**internos, "diagnosticos": [*internos["diagnosticos"], aviso]}
        return validacao, internos

    monkeypatch.setattr(modulo, "pre_validar_planilha_cadastral_completo", com_aviso)
    resposta = client.post(
        "/api/v1/importacoes-cadastrais/preparar",
        files={"arquivo": ("regulatorio.xlsx", planilha_v11())},
    )
    assert resposta.status_code == 201
    assert resposta.json()["avisos"] == [{
        "aba": "GERAL", "codigo": "AVISO_TESTE", "coluna": None,
        "linha": None, "mensagem": "Aviso seguro para revisão.",
        "severidade": "AVISO",
    }]
    assert db.scalar(select(SessaoImportacaoCadastral)).avisos == resposta.json()["avisos"]


def test_conteudos_v11_diferentes_nao_compartilham_hash_ou_payload(client, db):
    primeiro = planilha_v11()
    workbook = load_workbook(io.BytesIO(primeiro))
    workbook["MATERIAS_PRIMAS"]["C2"] = "Proteína alterada"
    saida = io.BytesIO()
    workbook.save(saida)
    segundo = saida.getvalue()

    respostas = [
        client.post(
            "/api/v1/importacoes-cadastrais/preparar",
            files={"arquivo": ("regulatorio.xlsx", conteudo)},
        ).json()
        for conteudo in (primeiro, segundo)
    ]
    assert respostas[0]["arquivo_sha256"] == sha256(primeiro).hexdigest()
    assert respostas[1]["arquivo_sha256"] == sha256(segundo).hexdigest()
    assert respostas[0]["arquivo_sha256"] != respostas[1]["arquivo_sha256"]
    sessoes = list(db.scalars(
        select(SessaoImportacaoCadastral).order_by(SessaoImportacaoCadastral.id)
    ))
    assert sessoes[0].payload_normalizado != sessoes[1].payload_normalizado


def test_payload_rejeita_inconsistencia_antes_de_criar_sessao(client, db, monkeypatch):
    import services.staging_importacao_cadastral as modulo
    original = modulo.pre_validar_planilha_cadastral_completo

    def inconsistente(*args, **kwargs):
        validacao, internos = original(*args, **kwargs)
        return {**validacao, "operacoes_total": validacao["operacoes_total"] + 1}, internos

    monkeypatch.setattr(
        modulo, "pre_validar_planilha_cadastral_completo", inconsistente
    )
    resposta = client.post(
        "/api/v1/importacoes-cadastrais/preparar",
        files={"arquivo": ("regulatorio.xlsx", planilha_v11())},
    )
    assert resposta.status_code == 500
    assert resposta.json() == {
        "detail": "Não foi possível preparar a importação cadastral."
    }
    assert db.scalar(select(SessaoImportacaoCadastral)) is None


def test_payload_v11_com_muitas_operacoes_nao_e_truncado(client, db):
    workbook = load_workbook(io.BytesIO(gerar_template_cadastral("1.1")))
    for aba in ABAS_DADOS:
        ws = workbook[aba]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
    for indice in range(510):
        workbook["CATEGORIAS_PRODUTO"].append(
            ("CRIAR", f"CAT_LIM_{indice:04d}", f"Categoria {indice}", None)
        )
    saida = io.BytesIO()
    workbook.save(saida)
    resposta = client.post(
        "/api/v1/importacoes-cadastrais/preparar",
        files={"arquivo": ("limite.xlsx", saida.getvalue())},
    )
    assert resposta.status_code == 201, resposta.text
    assert resposta.json()["operacoes_total"] == 510
    sessao = db.scalar(select(SessaoImportacaoCadastral))
    assert len(sessao.payload_normalizado["operacoes"]) == 510
    assert len(sessao.payload_normalizado["dados"]["CATEGORIAS_PRODUTO"]) == 510
    consulta = client.get(
        f"/api/v1/importacoes-cadastrais/{sessao.uuid_publico}"
    )
    assert consulta.status_code == 200
    assert "payload_normalizado" not in consulta.text
    assert "operacoes" not in consulta.json()


def test_reconstrucao_v10_preserva_cinco_abas_e_booleano():
    payload = {
        "versao": "1.0",
        "dados": {
            "MATERIAS_PRIMAS": [{
                "linha": 2, "acao": "CRIAR", "codigo": "MP_0007",
                "nome": "Matéria-prima", "ativa": False,
            }],
            "NUTRIENTES": [],
            "COMPOSICAO_NUTRICIONAL": [],
            "PRECOS_MP": [],
        },
    }
    conteudo = _xlsx_do_payload(payload)
    workbook = load_workbook(io.BytesIO(conteudo))
    assert workbook.sheetnames == [
        "LEIA_ME", "MATERIAS_PRIMAS", "NUTRIENTES",
        "COMPOSICAO_NUTRICIONAL", "PRECOS_MP",
    ]
    reparsed = parsear_planilha_cadastral(conteudo, "canonico-v10.xlsx")
    assert reparsed["dados"] == payload["dados"]
    assert reparsed["dados"]["MATERIAS_PRIMAS"][0]["ativa"] is False
